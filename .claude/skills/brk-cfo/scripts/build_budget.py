#!/usr/bin/env python3
"""يبني موازنة سنة كاملة (افتراضياً 2027) بمعادلات حية من بنود سنة الأساس ومحركاتها:
حجم المبيعات، وزيادة الأسعار، وسعر الصرف (رسمي وسوق ومرجح)، والتضخم، والرواتب، مع قائمة دخل شهرية،
وتدفق نقدي شهري بأيام التحصيل والدفع، وسيناريوهات، وحساسية سعر الصرف.

الاستخدام:
  python build_budget.py --base base.csv --fx-base 1450 --fx-market 1668 --fx-market-basis "بورصة الكفاح 2026-10-17" \
      --market-share 0.3 [--year 2027] [--volume-growth 0.05] [--price-increase 0.12] [--inflation 0.03] \
      [--wage-increase 0.1] [--headcount-change 0] [--season "1,1,1.2,..."] [--khazen <مجلد الخازن>] \
      [--opening-cash 3000000000] [--opening-receivables 11509329026] [--dso 91] [--dpo 30] \
      --out "الموازنة_التقديرية_2027.xlsx" --check --pdf

ملف الأساس (عناوين عربية أو إنكليزية):
  line/البند، group/المجموعة، driver/المحرك، amount/المبلغ (لسنة الأساس)، اختياري: months/أشهر_الأساس (افتراضياً 12)،
  usd_share/نسبة_الدولار (0 إلى 1)، source/المصدر
  المجموعة "إيرادات" للمبيعات الإجمالية قبل خصم الوكلاء. المجموعتان "اندثار" و"تمويل" تحت الربح التشغيلي.
  المحركات: مبيعات، حجم، تضخم، رواتب، ثابت، نسبة من المبيعات، يدوي (تنكتب أشهره يدوياً بالملف).

كل الافتراضات خلايا زرقاء تنعدل يدوياً. التضخم على الجزء المحلي فقط، والجزء الدولاري يتحرك بسعر الصرف الفعلي
مقابل سعر صرف الأساس، حتى لا يُحسب الأثر مرتين.
"""
import argparse
import csv
import datetime as dt
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_cost_workbook import (CENTER, F_BODY, F_BOLD, F_INPUT, F_NOTE, F_SECTION, F_SUB, F_TITLE,  # noqa: E402
                                 FMT_PCT, RIGHT, export_pdf, house_style_issues, page_setup, recalc_check,
                                 set_widths, style_body, style_header, style_total)
from validate_costs import parse_number  # noqa: E402

COMPANY = 'شركة البركة للاستثمارات الصناعية'
FMT_IQD = '#,##0;[Red](#,##0);"-"'
FMT_M = '#,##0,,;[Red](#,##0,,);"-"'          # بالملايين
FMT_PCT_SIGNED = '0.0%;[Red]-0.0%;"-"'
WRAP = Alignment(horizontal='right', vertical='center', wrap_text=True)
MONTHS = ['كانون الثاني', 'شباط', 'آذار', 'نيسان', 'أيار', 'حزيران', 'تموز', 'آب', 'أيلول', 'تشرين الأول', 'تشرين الثاني', 'كانون الأول']
S_ASSUME, S_SUM, S_LINES, S_PL, S_CASH, S_SCEN, S_METHOD = (
    'الافتراضات', 'الملخص', 'بنود الموازنة', 'قائمة الدخل الشهرية', 'التدفق النقدي', 'السيناريوهات', 'المنهجية والاعتماد')
REV, DEP, FIN = 'إيرادات', 'اندثار', 'تمويل'
DRIVERS = ['مبيعات', 'حجم', 'تضخم', 'رواتب', 'ثابت', 'نسبة من المبيعات', 'يدوي']
SEASONAL = {'مبيعات', 'حجم', 'نسبة من المبيعات'}
ALIASES = {
    'line': ('line', 'البند'), 'group': ('group', 'المجموعة'), 'driver': ('driver', 'المحرك'),
    'amount': ('amount', 'المبلغ', 'amount_base'), 'months': ('months', 'أشهر_الأساس'),
    'usd_share': ('usd_share', 'نسبة_الدولار'), 'source': ('source', 'المصدر'),
}


def read_base(path):
    with open(path, encoding='utf-8-sig', newline='') as f:
        rows = list(csv.DictReader(f))
    keys = {k.strip().lower(): k for k in rows[0]} if rows else {}
    col = {f: next((keys[a.lower()] for a in al if a.lower() in keys), None) for f, al in ALIASES.items()}
    for need in ('line', 'group', 'driver', 'amount'):
        if not col[need]:
            sys.exit(f'ملف الأساس: العمود {need} مطلوب')
    out, problems = [], []
    for r in rows:
        g = lambda f: (r.get(col[f]) or '').strip() if col[f] else ''
        line = {'line': g('line'), 'group': g('group'), 'driver': g('driver'), 'amount': parse_number(g('amount')),
                'months': parse_number(g('months')) or 12, 'usd': parse_number(g('usd_share')) or 0, 'source': g('source')}
        if line['driver'] not in DRIVERS:
            problems.append(f'{line["line"]}: محرك غير معروف "{line["driver"]}" (المسموح: {"، ".join(DRIVERS)})')
        if line['amount'] is None:
            problems.append(f'{line["line"]}: بدون مبلغ أساس')
        if line['group'] == REV and line['usd']:
            problems.append(f'{line["line"]}: الإيرادات بالدينار، نسبة الدولار لازم صفر')
        out.append(line)
    if not any(l['group'] == REV for l in out):
        problems.append('ماكو بند بمجموعة إيرادات')
    if problems:
        print('مشاكل بملف الأساس:')
        for p in problems:
            print('  -', p)
        sys.exit(1)
    return out


def season_from_khazen(folder):
    """أوزان موسمية من صافي مبيعات الخازن الشهرية، والأشهر الناقصة بمتوسط الموجود."""
    path = Path(folder) / '16_فواتير_البيع_تفاصيل.csv'
    sums = defaultdict(float)
    with open(path, encoding='utf-8-sig', newline='') as f:
        for r in csv.DictReader(f):
            sign = -1 if r['BillKind'] == '021101' else 1
            sums[int(r['BillDate'][5:7])] += sign * float(r['NetTotal'] or 0)
    avg = sum(sums.values()) / len(sums)
    return [round(sums.get(m, avg)) for m in range(1, 13)], f'صافي مبيعات الخازن الشهرية، والأشهر الناقصة ({12 - len(sums)}) بالمتوسط'


def build(args, lines):
    wb = Workbook()
    wb.remove(wb.active)
    A = {}
    yr = args.year

    # ---------- الافتراضات ----------
    ws = wb.create_sheet(S_ASSUME)
    ws.merge_cells('A1:C1')
    ws['A1'], ws['A1'].font, ws['A1'].alignment = COMPANY, F_TITLE, CENTER
    ws.merge_cells('A2:C2')
    ws['A2'], ws['A2'].font, ws['A2'].alignment = f'افتراضات الموازنة التقديرية لسنة {yr}', F_SUB, CENTER
    ws.merge_cells('A3:C3')
    ws['A3'] = 'الخلايا الزرقاء إدخال يدوي، وكل الأوراق تتحدث منها.'
    ws['A3'].font, ws['A3'].alignment = F_NOTE, CENTER
    for j, h in enumerate(['البند', 'القيمة', 'الأساس والمصدر'], start=1):
        ws.cell(row=5, column=j, value=h)
    style_header(ws, 5, 1, 3)
    need = 'يلزم تحديده من الإدارة'
    items = [
        ('h_fx', 'سعر الصرف', None, None, None),
        ('fx_base', 'سعر الصرف بأرقام سنة الأساس', args.fx_base, '#,##0', args.fx_base_basis),
        ('fx_off', 'سعر الصرف الرسمي', args.fx_official, '#,##0', 'قرار مجلس الوزراء 544، نافذ من 2026-10-07'),
        ('fx_mkt', 'سعر صرف السوق الموازي المتوقع', args.fx_market, '#,##0', args.fx_market_basis),
        ('mkt_share', 'نسبة الدولار المشترى من السوق', args.market_share, FMT_PCT, args.market_share_basis),
        ('fx_eff', 'سعر الصرف الفعلي المرجح', None, '#,##0', 'تُحسب: الرسمي × (1 − نسبة السوق) + السوق × نسبة السوق'),
        ('h_grow', 'النمو والأسعار', None, None, None),
        ('vol', 'نمو حجم المبيعات', args.volume_growth, FMT_PCT, args.volume_basis or need),
        ('price', 'زيادة أسعار البيع (متوسط السنة)', args.price_increase, FMT_PCT, args.price_basis or need),
        ('disc', 'خصم الوكلاء', args.discount, FMT_PCT, 'سياسة الخصم الحالية'),
        ('h_cost', 'الكلف', None, None, None),
        ('infl', 'التضخم السنوي للكلف المحلية', args.inflation, FMT_PCT, args.inflation_basis),
        ('wage', 'زيادة الرواتب', args.wage_increase, FMT_PCT, args.wage_basis or need),
        ('hc', 'تغير عدد الموظفين', args.headcount_change, FMT_PCT, args.headcount_basis or need),
        ('tax', 'ضريبة الدخل على الشركات', args.income_tax, FMT_PCT, 'القانون 113 لسنة 1982 المعدل: 15%'),
        ('h_cash', 'النقد', None, None, None),
        ('cash0', 'الرصيد النقدي الافتتاحي', args.opening_cash, FMT_IQD, args.opening_cash_basis or 'يُدخل من الصناديق والمصارف بتاريخ بداية السنة'),
        ('ar0', 'ذمم الوكلاء الافتتاحية', args.opening_receivables, FMT_IQD, args.opening_receivables_basis or 'يُدخل من كشف الوكلاء'),
        ('ap0', 'ذمم الموردين الافتتاحية', args.opening_payables, FMT_IQD, 'يُدخل من كشف الموردين'),
        ('dso', 'أيام التحصيل', args.dso, '0', 'متوسط أيام تحصيل الوكلاء، خط الأساس 2026-09-30'),
        ('dpo', 'أيام الدفع للموردين', args.dpo, '0', args.dpo_basis or 'تقدير يلزم تأكيده'),
        ('lag_in', 'تأخير التحصيل بالأشهر', None, '0', 'تُحسب: أيام التحصيل ÷ 30 مقربة'),
        ('lag_out', 'تأخير الدفع بالأشهر', None, '0', 'تُحسب: أيام الدفع ÷ 30 مقربة'),
    ]
    R = {key: i for i, (key, *_) in enumerate(items, start=6)}
    for key, i in R.items():
        A[key] = f"'{S_ASSUME}'!$B${i}"
    formulas = {
        'fx_eff': f'=B{R["fx_off"]}*(1-B{R["mkt_share"]})+B{R["fx_mkt"]}*B{R["mkt_share"]}',
        'lag_in': f'=ROUND(B{R["dso"]}/30,0)', 'lag_out': f'=ROUND(B{R["dpo"]}/30,0)',
    }
    for key, i in R.items():
        _, label, val, fmt, basis = items[i - 6]
        if key.startswith('h_'):
            ws.cell(row=i, column=1, value=label).font = F_SECTION
            continue
        style_body(ws.cell(row=i, column=1, value=label), font=F_BOLD, align=RIGHT)
        style_body(ws.cell(row=i, column=2, value=formulas.get(key, val)), fmt,
                   font=F_BODY if key in formulas else F_INPUT, align=CENTER)
        style_body(ws.cell(row=i, column=3, value=basis), align=WRAP)
    # الموسمية
    r0 = len(items) + 8
    ws.cell(row=r0, column=1, value='الموسمية الشهرية للمبيعات والكلف المتغيرة').font = F_SECTION
    ws.cell(row=r0 + 1, column=1, value=args.season_basis).font = F_NOTE
    for j, h in enumerate(['الشهر', 'الوزن (إدخال)', 'النسبة'], start=1):
        ws.cell(row=r0 + 2, column=j, value=h)
    style_header(ws, r0 + 2, 1, 3)
    s_first, s_last = r0 + 3, r0 + 14
    for m, w in enumerate(args.season, start=0):
        rr = s_first + m
        style_body(ws.cell(row=rr, column=1, value=MONTHS[m]), font=F_BOLD, align=RIGHT)
        style_body(ws.cell(row=rr, column=2, value=w), '#,##0', font=F_INPUT, align=CENTER)
        style_body(ws.cell(row=rr, column=3, value=f'=B{rr}/SUM($B${s_first}:$B${s_last})'), FMT_PCT, align=CENTER)
    SEASON = lambda m: f"'{S_ASSUME}'!$C${s_first + m}"
    set_widths(ws, [40, 18, 80])

    # ---------- بنود الموازنة ----------
    ln = wb.create_sheet(S_LINES)
    cols = [('line', 'البند'), ('group', 'المجموعة'), ('driver', 'المحرك'), ('usd', 'نسبة الدولار'), ('amount', 'مبلغ الأساس'),
            ('months', 'أشهر الأساس'), ('annual0', 'الأساس السنوي'), ('factor', f'معامل {yr}'), ('budget', f'موازنة {yr}'),
            ('chg', 'التغير'), ('usd_part', 'الجزء الدولاري'), ('source', 'المصدر')] + [(f'm{m}', MONTHS[m]) for m in range(12)]
    C = {k: get_column_letter(j) for j, (k, _) in enumerate(cols, start=1)}
    for j, (_, h) in enumerate(cols, start=1):
        ln.cell(row=1, column=j, value=h)
    style_header(ln, 1, 1, len(cols))
    ln.row_dimensions[1].height = 32
    v, p, inf, fe, fb = A['vol'], A['price'], A['infl'], A['fx_eff'], A['fx_base']
    for i, l in enumerate(lines, start=2):
        x = {k: f'{C[k]}{i}' for k in C}
        mix = f'((1-{x["usd"]})*(1+{inf})+{x["usd"]}*{fe}/{fb})'
        factor = (f'=IF({x["driver"]}="مبيعات",(1+{v})*(1+{p}),IF({x["driver"]}="حجم",(1+{v})*{mix},'
                  f'IF({x["driver"]}="تضخم",{mix},IF({x["driver"]}="رواتب",(1+{A["wage"]})*(1+{A["hc"]}),'
                  f'IF({x["driver"]}="نسبة من المبيعات",(1+{v})*(1+{p}),1)))))')
        manual = l['driver'] == 'يدوي'
        months_range = f'{C["m0"]}{i}:{C["m11"]}{i}'
        row = {
            'line': (l['line'], F_BODY, None), 'group': (l['group'], F_BODY, None), 'driver': (l['driver'], F_INPUT, None),
            'usd': (l['usd'], F_INPUT, FMT_PCT), 'amount': (l['amount'], F_INPUT, FMT_IQD), 'months': (l['months'], F_INPUT, '0'),
            'annual0': (f'={x["amount"]}/{x["months"]}*12', F_BODY, FMT_IQD),
            'factor': ('' if manual else factor, F_BODY, '0.000'),
            'budget': (f'=SUM({months_range})' if manual else f'={x["annual0"]}*{x["factor"]}', F_BOLD, FMT_IQD),
            'chg': (f'=IF({x["annual0"]}=0,"",{x["budget"]}/{x["annual0"]}-1)', F_BODY, FMT_PCT_SIGNED),
            'usd_part': (f'=IF(OR({x["driver"]}="حجم",{x["driver"]}="تضخم"),{x["annual0"]}*IF({x["driver"]}="حجم",1+{v},1)'
                         f'*{x["usd"]}*{fe}/{fb},0)', F_BODY, FMT_IQD),
            'source': (l['source'], F_BODY, None),
        }
        for m in range(12):
            if manual:
                row[f'm{m}'] = (None, F_INPUT, FMT_IQD)
            else:
                w = SEASON(m) if l['driver'] in SEASONAL else '1/12'
                row[f'm{m}'] = (f'={x["budget"]}*{w}', F_BODY, FMT_IQD)
        for j, (k, _) in enumerate(cols, start=1):
            val, font, fmt = row[k]
            style_body(ln.cell(row=i, column=j, value=val), fmt, font=font, align=CENTER if k in ('group', 'driver', 'months') else RIGHT)
    last = len(lines) + 1
    dv = DataValidation(type='list', formula1='"' + ','.join(d for d in DRIVERS if d != 'يدوي') + '"', allow_blank=False)
    ln.add_data_validation(dv)
    for i, l in enumerate(lines, start=2):
        if l['driver'] != 'يدوي':
            dv.add(f'{C["driver"]}{i}')
    tot = last + 1
    ln.cell(row=tot, column=1, value='صافي (الإيرادات ناقص الكلف)')
    for k in ['annual0', 'budget', 'usd_part'] + [f'm{m}' for m in range(12)]:
        col = C[k]
        if k == 'usd_part':
            ln[f'{col}{tot}'] = f'=SUM({col}2:{col}{last})'
        else:
            ln[f'{col}{tot}'] = f'=SUMIF($B$2:$B${last},"{REV}",{col}2:{col}{last})-SUMIF($B$2:$B${last},"<>{REV}",{col}2:{col}{last})'
    style_total(ln, tot, 1, len(cols))
    for k in ['annual0', 'budget', 'usd_part'] + [f'm{m}' for m in range(12)]:
        ln[f'{C[k]}{tot}'].number_format = FMT_IQD
    set_widths(ln, [30, 14, 14, 9, 16, 8, 16, 9, 16, 9, 15, 30] + [14] * 12)
    ln.freeze_panes = 'B2'
    LN = lambda k: f"'{S_LINES}'!${C[k]}$2:${C[k]}${last}"

    # ---------- قائمة الدخل الشهرية ----------
    pl = wb.create_sheet(S_PL)
    groups = []
    for l in lines:
        if l['group'] not in groups and l['group'] not in (REV, DEP, FIN):
            groups.append(l['group'])
    heads = ['البند'] + MONTHS + [f'سنة {yr}', 'الأساس السنوي', 'التغير']
    for j, h in enumerate(heads, start=1):
        pl.cell(row=1, column=j, value=h)
    style_header(pl, 1, 1, len(heads))
    YC, BC, CC = get_column_letter(14), get_column_letter(15), get_column_letter(16)
    rows, P = [], {}

    def add(key, label, per_month, base=None, bold=False):
        r = len(rows) + 2
        P[key] = r
        rows.append((key, label, per_month, base, bold))

    sumif = lambda grp, colk: f"SUMIF({LN('group')},\"{grp}\",{LN(colk)})"
    add('gross', 'المبيعات الإجمالية', lambda m: f'={sumif(REV, f"m{m}")}', f'={sumif(REV, "annual0")}')
    add('disc', 'خصم الوكلاء', lambda m: f'=-{get_column_letter(m + 2)}{P["gross"]}*{A["disc"]}',
        f'=-{BC}{P["gross"]}*{A["disc"]}')
    add('net', 'صافي المبيعات', lambda m: f'={get_column_letter(m + 2)}{P["gross"]}+{get_column_letter(m + 2)}{P["disc"]}',
        f'={BC}{P["gross"]}+{BC}{P["disc"]}', True)
    for g in groups:
        add(f'g_{g}', g, lambda m, g=g: f'=-{sumif(g, f"m{m}")}', f'=-{sumif(g, "annual0")}')
    first_g, last_g = P['net'] + 1, P['net'] + len(groups)
    add('ebitda', 'الربح قبل الاندثار والتمويل والضريبة',
        lambda m: f'=SUM({get_column_letter(m + 2)}{P["net"]}:{get_column_letter(m + 2)}{last_g})',
        f'=SUM({BC}{P["net"]}:{BC}{last_g})', True)
    add('dep', 'الاندثار', lambda m: f'=-{sumif(DEP, f"m{m}")}', f'=-{sumif(DEP, "annual0")}')
    add('fin', 'كلف التمويل', lambda m: f'=-{sumif(FIN, f"m{m}")}', f'=-{sumif(FIN, "annual0")}')
    add('pbt', 'الربح قبل الضريبة',
        lambda m: f'={get_column_letter(m + 2)}{P["ebitda"]}+{get_column_letter(m + 2)}{P["dep"]}+{get_column_letter(m + 2)}{P["fin"]}',
        f'={BC}{P["ebitda"]}+{BC}{P["dep"]}+{BC}{P["fin"]}', True)
    add('tax', 'ضريبة الدخل (سنوية)', lambda m: None, f'=-MAX(0,{BC}{P["pbt"]})*{A["tax"]}')
    add('np', 'صافي الربح', lambda m: None, f'={BC}{P["pbt"]}+{BC}{P["tax"]}', True)
    add('margin_e', 'هامش الربح التشغيلي', lambda m: f'=IF({get_column_letter(m + 2)}{P["net"]}=0,"",{get_column_letter(m + 2)}{P["ebitda"]}/{get_column_letter(m + 2)}{P["net"]})',
        f'=IF({BC}{P["net"]}=0,"",{BC}{P["ebitda"]}/{BC}{P["net"]})')
    for key, label, per_month, base, bold in rows:
        r = P[key]
        style_body(pl.cell(row=r, column=1, value=label), font=F_BOLD if bold else F_BODY, align=RIGHT)
        fmt = FMT_PCT if key == 'margin_e' else FMT_M
        for m in range(12):
            style_body(pl.cell(row=r, column=m + 2, value=per_month(m)), fmt, font=F_BOLD if bold else F_BODY, align=RIGHT)
        if key == 'tax':
            year_f = f'=-MAX(0,{YC}{P["pbt"]})*{A["tax"]}'
        elif key == 'np':
            year_f = f'={YC}{P["pbt"]}+{YC}{P["tax"]}'
        elif key == 'margin_e':
            year_f = f'=IF({YC}{P["net"]}=0,"",{YC}{P["ebitda"]}/{YC}{P["net"]})'
        else:
            year_f = f'=SUM(B{r}:M{r})'
        style_body(pl.cell(row=r, column=14, value=year_f), fmt, font=F_BOLD, align=RIGHT)
        style_body(pl.cell(row=r, column=15, value=base), fmt, font=F_BODY, align=RIGHT)
        # الكلف سالبة بالقائمة، فتغيرها نسبة مباشرة (زيادة الكلفة موجبة)، والأرباح بالفرق على القيمة المطلقة
        chg = '' if key == 'margin_e' else f'=IF({BC}{r}=0,"",IF({BC}{r}<0,{YC}{r}/{BC}{r}-1,({YC}{r}-{BC}{r})/ABS({BC}{r})))'
        style_body(pl.cell(row=r, column=16, value=chg), FMT_PCT_SIGNED, align=CENTER)
        if bold:
            style_total(pl, r, 1, 16)
            for c in range(2, 16):
                pl.cell(row=r, column=c).number_format = fmt
    n = len(rows) + 3
    pl.cell(row=n, column=1, value='المبالغ بالملايين. الخصم يُستحق شهرياً على المبيعات الإجمالية.').font = F_NOTE
    set_widths(pl, [34] + [11] * 12 + [13, 13, 9])
    pl.freeze_panes = 'B2'
    PLr = lambda key, col: f"'{S_PL}'!${col}${P[key]}"

    # ---------- التدفق النقدي ----------
    cf = wb.create_sheet(S_CASH)
    heads = ['البند'] + MONTHS + [f'سنة {yr}']
    for j, h in enumerate(heads, start=1):
        cf.cell(row=1, column=j, value=h)
    style_header(cf, 1, 1, len(heads))
    labels = ['رقم الشهر', 'الرصيد الافتتاحي', 'المقبوضات من الوكلاء', 'الكلف النقدية للشهر', 'المدفوعات التشغيلية',
              'الاستثمارات الرأسمالية', 'سحوبات الشركاء', 'ضريبة الدخل وأقساط القروض', 'صافي التدفق', 'الرصيد الختامي',
              'أشهر تغطية الرواتب']
    K = {k: i for i, k in enumerate(['n', 'open', 'in', 'costs', 'out', 'capex', 'draw', 'other', 'net', 'close', 'cover'], start=2)}
    pay_rows = [P[f'g_{g}'] for g in groups] + [P['fin']]
    wage_groups = [g for g in groups if 'رواتب' in g]
    for key, r in K.items():
        style_body(cf.cell(row=r, column=1, value=labels[r - 2]), font=F_BOLD if key in ('net', 'close') else F_BODY, align=RIGHT)
    netrow = f"'{S_PL}'!$B${P['net']}:$M${P['net']}"
    for m in range(12):
        col = get_column_letter(m + 2)
        prev = get_column_letter(m + 1)
        pl_col = get_column_letter(m + 2)
        cells = {
            'n': (m + 1, '0', F_BODY),
            'open': (f'={A["cash0"]}' if m == 0 else f'={prev}{K["close"]}', FMT_M, F_BODY),
            'in': (f'=IF({col}{K["n"]}>{A["lag_in"]},INDEX({netrow},{col}{K["n"]}-{A["lag_in"]}),{A["ar0"]}/MAX({A["lag_in"]},1))',
                   FMT_M, F_BODY),
            'costs': ('=-(' + '+'.join(f"'{S_PL}'!{pl_col}{r}" for r in pay_rows) + ')', FMT_M, F_BODY),
            'out': (f'=IF({A["lag_out"]}=0,{col}{K["costs"]},IF({col}{K["n"]}>{A["lag_out"]},'
                    f'INDEX($B${K["costs"]}:$M${K["costs"]},{col}{K["n"]}-{A["lag_out"]}),{A["ap0"]}/{A["lag_out"]}))', FMT_M, F_BODY),
            'capex': (None, FMT_M, F_INPUT), 'draw': (None, FMT_M, F_INPUT), 'other': (None, FMT_M, F_INPUT),
            'net': (f'={col}{K["in"]}-{col}{K["out"]}-{col}{K["capex"]}-{col}{K["draw"]}-{col}{K["other"]}', FMT_M, F_BOLD),
            'close': (f'={col}{K["open"]}+{col}{K["net"]}', FMT_M, F_BOLD),
            'cover': ('=IFERROR(' + f'{col}{K["close"]}/(' + ('+'.join(f"-'{S_PL}'!{pl_col}{P['g_' + g]}" for g in wage_groups) or '0')
                      + '),"")', '0.0', F_BODY),
        }
        for key, (val, fmt, font) in cells.items():
            style_body(cf.cell(row=K[key], column=m + 2, value=val), fmt, font=font, align=RIGHT)
    for key in ('in', 'costs', 'out', 'capex', 'draw', 'other', 'net'):
        style_body(cf.cell(row=K[key], column=14, value=f'=SUM(B{K[key]}:M{K[key]})'), FMT_M, font=F_BOLD, align=RIGHT)
    style_body(cf.cell(row=K['close'], column=14, value=f'=M{K["close"]}'), FMT_M, font=F_BOLD, align=RIGHT)
    for key in ('net', 'close'):
        style_total(cf, K[key], 1, 14)
        for c in range(2, 15):
            cf.cell(row=K[key], column=c).number_format = FMT_M
    nr = K['cover'] + 2
    for t in ['المبالغ بالملايين. الخلايا الزرقاء (الاستثمارات والسحوبات والضريبة والأقساط) إدخال يدوي شهري.',
              'المقبوضات = صافي مبيعات الشهر بعد تأخير التحصيل، وذمم الوكلاء الافتتاحية تتحصل بالتساوي بأشهر التأخير.',
              'المدفوعات = الكلف النقدية (بدون الاندثار) بعد تأخير الدفع. أثر سعر الصرف على النقد بالدولار غير محسوب.',
              'سحوبات الشركاء وتوزيعاتهم تتطلب موافقة الشركاء.']:
        cf.cell(row=nr, column=1, value=t).font = F_NOTE
        nr += 1
    set_widths(cf, [34] + [11] * 12 + [13])
    cf.freeze_panes = 'B2'

    # ---------- السيناريوهات ----------
    sc = wb.create_sheet(S_SCEN)
    sc.merge_cells('A1:I1')
    sc['A1'], sc['A1'].font, sc['A1'].alignment = 'سيناريوهات الموازنة (كل خلية زرقاء تنعدل)', F_SECTION, CENTER
    heads = ['السيناريو', 'نمو الحجم', 'زيادة الأسعار', 'سعر الصرف الفعلي', 'التضخم', 'زيادة الرواتب',
             'صافي المبيعات', 'الربح قبل الضريبة', 'الهامش']
    for j, h in enumerate(heads, start=1):
        sc.cell(row=3, column=j, value=h)
    style_header(sc, 3, 1, len(heads))
    base_links = [f'={A["vol"]}', f'={A["price"]}', f'={A["fx_eff"]}', f'={A["infl"]}', f'={A["wage"]}']
    scen = [('الأساس (من الافتراضات)', base_links)]
    fx_hi = max(args.fx_market + 100, 1750)
    scen += [
        ('انخفاض الحجم 10%', [f'={A["vol"]}-0.1', f'={A["price"]}', f'={A["fx_eff"]}', f'={A["infl"]}', f'={A["wage"]}']),
        ('بدون زيادة أسعار', [f'={A["vol"]}', 0, f'={A["fx_eff"]}', f'={A["infl"]}', f'={A["wage"]}']),
        ('الدولار كله من السوق', [f'={A["vol"]}', f'={A["price"]}', f'={A["fx_mkt"]}', f'={A["infl"]}', f'={A["wage"]}']),
        (f'الدولار {fx_hi:,.0f}', [f'={A["vol"]}', f'={A["price"]}', fx_hi, f'={A["infl"]}', f'={A["wage"]}']),
        ('تضخم مضاعف', [f'={A["vol"]}', f'={A["price"]}', f'={A["fx_eff"]}', f'={A["infl"]}*2', f'={A["wage"]}']),
        ('الأسوأ مجتمعاً', [f'={A["vol"]}-0.1', 0, fx_hi, f'={A["infl"]}*2', f'={A["wage"]}']),
    ]
    drv, grp, a0, usd, bud = LN('driver'), LN('group'), LN('annual0'), LN('usd'), LN('budget')
    for i, (name, vals) in enumerate(scen, start=4):
        style_body(sc.cell(row=i, column=1, value=name), font=F_BOLD, align=RIGHT)
        for j, val in enumerate(vals, start=2):
            fmt = '#,##0' if j == 4 else FMT_PCT
            style_body(sc.cell(row=i, column=j, value=val), fmt, font=F_INPUT, align=CENTER)
        Bv, Cv, Dv, Ev, Fv = (f'$B{i}', f'$C{i}', f'$D{i}', f'$E{i}', f'$F{i}')
        mix = f'((1-{usd})*(1+{Ev})+{usd}*{Dv}/{fb})'
        amount = (f'{a0}*(({drv}="مبيعات")*(1+{Bv})*(1+{Cv})+({drv}="نسبة من المبيعات")*(1+{Bv})*(1+{Cv})'
                  f'+({drv}="حجم")*(1+{Bv})*{mix}+({drv}="تضخم")*{mix}+({drv}="رواتب")*(1+{Fv})*(1+{A["hc"]})+({drv}="ثابت"))'
                  f'+({drv}="يدوي")*{bud}')
        net = f'SUMPRODUCT(({grp}="{REV}")*{amount})*(1-{A["disc"]})'
        style_body(sc.cell(row=i, column=7, value=f'={net}'), FMT_M, align=RIGHT)
        style_body(sc.cell(row=i, column=8, value=f'=G{i}-SUMPRODUCT(({grp}<>"{REV}")*{amount})'), FMT_M, font=F_BOLD, align=RIGHT)
        style_body(sc.cell(row=i, column=9, value=f'=IF(G{i}=0,"",H{i}/G{i})'), FMT_PCT_SIGNED, align=CENTER)
    chk = len(scen) + 5
    sc.cell(row=chk, column=1, value='فحص: الأساس مقابل قائمة الدخل (لازم صفر)').font = F_NOTE
    style_body(sc.cell(row=chk, column=8, value=f'=ROUND(H4-{PLr("pbt", "N")},0)'), FMT_IQD, align=RIGHT)
    sc.cell(row=chk + 1, column=1, value='المبالغ بالملايين. السيناريو يعيد حساب كل البنود بمحركاتها، والبنود اليدوية تبقى كما هي.').font = F_NOTE
    set_widths(sc, [26, 11, 12, 14, 10, 12, 15, 16, 9])

    # ---------- الملخص ----------
    sm = wb.create_sheet(S_SUM, 1)
    sm.merge_cells('A1:D1')
    sm['A1'], sm['A1'].font, sm['A1'].alignment = COMPANY, F_TITLE, CENTER
    sm.merge_cells('A2:D2')
    sm['A2'], sm['A2'].font, sm['A2'].alignment = f'الموازنة التقديرية لسنة {yr}', F_SUB, CENTER
    sm.merge_cells('A3:D3')
    sm['A3'] = 'مسودة للنقاش. اعتماد الموازنة يتطلب موافقة المدير العام والشركاء.'
    sm['A3'].font, sm['A3'].alignment = F_NOTE, CENTER
    for j, h in enumerate(['المؤشر', f'سنة {yr}', 'الأساس السنوي', 'التغير'], start=1):
        sm.cell(row=5, column=j, value=h)
    style_header(sm, 5, 1, 4)
    kp = [('صافي المبيعات', 'net'), ('الربح قبل الاندثار والتمويل والضريبة', 'ebitda'), ('الربح قبل الضريبة', 'pbt'),
          ('صافي الربح', 'np')]
    for i, (label, key) in enumerate(kp, start=6):
        style_body(sm.cell(row=i, column=1, value=label), font=F_BOLD, align=RIGHT)
        style_body(sm.cell(row=i, column=2, value=f'={PLr(key, "N")}'), FMT_IQD, font=F_BOLD, align=RIGHT)
        style_body(sm.cell(row=i, column=3, value=f'={PLr(key, "O")}'), FMT_IQD, align=RIGHT)
        style_body(sm.cell(row=i, column=4, value=f'=IF(C{i}=0,"",(B{i}-C{i})/ABS(C{i}))'), FMT_PCT_SIGNED, align=CENTER)
    r = len(kp) + 6
    extra = [
        ('هامش الربح التشغيلي', f'={PLr("margin_e", "N")}', FMT_PCT_SIGNED),
        ('أدنى رصيد نقدي بالسنة', f"=MIN('{S_CASH}'!B{K['close']}:M{K['close']})", FMT_IQD),
        ('شهر أدنى رصيد', f"=INDEX('{S_CASH}'!B1:M1,MATCH(B{r + 1},'{S_CASH}'!B{K['close']}:M{K['close']},0))", None),
        ('الرصيد النقدي آخر السنة', f"='{S_CASH}'!M{K['close']}", FMT_IQD),
        ('سعر الصرف الفعلي المرجح', f'={A["fx_eff"]}', '#,##0'),
        ('الجزء الدولاري من الكلف', f"=SUM({LN('usd_part')})", FMT_IQD),
    ]
    for k, (label, f, fmt) in enumerate(extra, start=r):
        style_body(sm.cell(row=k, column=1, value=label), font=F_BOLD, align=RIGHT)
        style_body(sm.cell(row=k, column=2, value=f), fmt, font=F_BOLD, align=RIGHT)
    r += len(extra) + 1
    sm.cell(row=r, column=1, value='حساسية سعر الصرف').font = F_SECTION
    for j, h in enumerate(['سعر الصرف الفعلي', 'الربح قبل الضريبة', 'الفرق عن الموازنة'], start=1):
        sm.cell(row=r + 1, column=j, value=h)
    style_header(sm, r + 1, 1, 3)
    for k, fx in enumerate(sorted({args.fx_official, round(args.fx_market), 1600, 1750, 1850}), start=r + 2):
        style_body(sm.cell(row=k, column=1, value=fx), '#,##0', font=F_INPUT, align=CENTER)
        style_body(sm.cell(row=k, column=2, value=f"={PLr('pbt', 'N')}-SUM({LN('usd_part')})*(A{k}/{A['fx_eff']}-1)"), FMT_IQD, align=RIGHT)
        style_body(sm.cell(row=k, column=3, value=f"=B{k}-{PLr('pbt', 'N')}"), FMT_IQD, align=RIGHT)
    r = k + 2
    sm.cell(row=r, column=1, value='ملاحظات').font = F_SECTION
    notes = [
        'التضخم يُطبق على الجزء المحلي من كل بند، والجزء الدولاري يتحرك بسعر الصرف الفعلي مقابل سعر صرف الأساس.',
        'الإيرادات تنمو بالحجم والسعر معاً، والكلف المتغيرة بالحجم. زيادة الأسعار قرار تسعير منفصل (build_pricing.py).',
        'أرقام الأساس تلزم مطابقتها مع الدفاتر قبل الاعتماد. "غير متوفر" ليس صفراً.',
        'التوزيعات والقروض وبيع الأصول تتطلب موافقة الشركاء.',
    ]
    for k, t in enumerate(notes, start=r + 1):
        style_body(sm.cell(row=k, column=1, value=t), align=WRAP)
        sm.merge_cells(start_row=k, start_column=1, end_row=k, end_column=4)
        sm.row_dimensions[k].height = 30
    set_widths(sm, [40, 20, 20, 12])

    # ---------- المنهجية والاعتماد ----------
    me = wb.create_sheet(S_METHOD)
    me['A1'], me['A1'].font = 'المنهجية والاعتماد', F_TITLE
    method = [
        'الأساس السنوي = مبلغ الأساس ÷ أشهر الأساس × 12.',
        'المعاملات: مبيعات ونسبة من المبيعات = (1 + الحجم) × (1 + السعر). حجم = (1 + الحجم) × خليط. تضخم = خليط. رواتب = (1 + الزيادة) × (1 + تغير العدد). ثابت = 1.',
        'الخليط = (1 − نسبة الدولار) × (1 + التضخم) + نسبة الدولار × سعر الصرف الفعلي ÷ سعر صرف الأساس.',
        'التوزيع الشهري: المبيعات والكلف المتغيرة بالموسمية، والباقي بالتساوي، واليدوي كما يُكتب.',
        'خصم الوكلاء يُستحق شهرياً على المبيعات الإجمالية (IFRS 15).',
        'حساسية الصرف = الربح قبل الضريبة − الجزء الدولاري × (السعر ÷ السعر الفعلي − 1).',
        f'ملف الأساس: {args.base.name}. تاريخ الإعداد: {dt.date.today().isoformat()}.',
    ]
    for i, t in enumerate(method, start=3):
        style_body(me.cell(row=i, column=1, value=t), align=WRAP)
        me.merge_cells(start_row=i, start_column=1, end_row=i, end_column=4)
        me.row_dimensions[i].height = 30
    r = len(method) + 4
    for j, h in enumerate(['الدور', 'الاسم', 'التوقيع', 'التاريخ'], start=1):
        me.cell(row=r, column=j, value=h)
    style_header(me, r, 1, 4)
    for k, role in enumerate(['أعدّه: الإدارة المالية', 'راجعه: المدير العام', 'اعتمده: الشركاء'], start=r + 1):
        for j, val in enumerate([role, None, None, None], start=1):
            style_body(me.cell(row=k, column=j, value=val), font=F_BOLD if j == 1 else F_INPUT, align=RIGHT)
        me.row_dimensions[k].height = 24
    set_widths(me, [34, 26, 22, 16])

    for w in wb.worksheets:
        page_setup(w, COMPANY, '1:1' if w.title in (S_LINES, S_PL, S_CASH) else None)
        if w.title in (S_ASSUME, S_SUM):
            w.page_setup.fitToHeight = 1
    wb.properties.creator = wb.properties.lastModifiedBy = 'الإدارة المالية'
    wb.properties.title = f'الموازنة التقديرية {yr}'
    wb.save(args.out)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--base', type=Path, required=True)
    ap.add_argument('--out', type=Path, required=True)
    ap.add_argument('--year', type=int, default=2027)
    ap.add_argument('--fx-base', type=float, required=True, help='سعر الصرف اللي انحسبت بيه أرقام الأساس')
    ap.add_argument('--fx-base-basis', default='سعر قيد الدولار بدفاتر سنة الأساس')
    ap.add_argument('--fx-official', type=float, default=1520)
    ap.add_argument('--fx-market', type=float, required=True)
    ap.add_argument('--fx-market-basis', required=True)
    ap.add_argument('--market-share', type=float, required=True)
    ap.add_argument('--market-share-basis', default='نسبة مشتريات الدولار اللي ما تمر بالسعر الرسمي. يلزم تأكيدها')
    ap.add_argument('--volume-growth', type=float, default=0.0)
    ap.add_argument('--volume-basis')
    ap.add_argument('--price-increase', type=float, default=0.0)
    ap.add_argument('--price-basis')
    ap.add_argument('--discount', type=float, default=0.14)
    ap.add_argument('--inflation', type=float, default=0.03)
    ap.add_argument('--inflation-basis', default='صندوق النقد الدولي، آفاق الاقتصاد العالمي تشرين الأول 2026: 3% لسنة 2027 (قبل أثر قرار 544 بالكامل)')
    ap.add_argument('--wage-increase', type=float, default=0.0)
    ap.add_argument('--wage-basis')
    ap.add_argument('--headcount-change', type=float, default=0.0)
    ap.add_argument('--headcount-basis')
    ap.add_argument('--income-tax', type=float, default=0.15)
    ap.add_argument('--opening-cash', type=float)
    ap.add_argument('--opening-cash-basis')
    ap.add_argument('--opening-receivables', type=float)
    ap.add_argument('--opening-receivables-basis')
    ap.add_argument('--opening-payables', type=float)
    ap.add_argument('--dso', type=float, default=91)
    ap.add_argument('--dpo', type=float, default=30)
    ap.add_argument('--dpo-basis')
    ap.add_argument('--season', help='12 وزناً مفصولة بفواصل')
    ap.add_argument('--khazen', type=Path, help='مجلد تصدير الخازن لحساب الموسمية من المبيعات الشهرية')
    ap.add_argument('--check', action='store_true')
    ap.add_argument('--pdf', action='store_true')
    args = ap.parse_args()

    if args.season:
        args.season = [float(x) for x in args.season.split(',')]
        if len(args.season) != 12:
            sys.exit('--season لازم 12 رقم')
        args.season_basis = 'أوزان مُدخلة'
    elif args.khazen:
        args.season, args.season_basis = season_from_khazen(args.khazen)
    else:
        args.season, args.season_basis = [1] * 12, 'بالتساوي لعدم توفر الموسمية'

    lines = read_base(args.base)
    print(f'الأساس: {len(lines)} بند')
    build(args, lines)
    print(f'الملف: {args.out}')
    if args.check:
        style = house_style_issues(args.out)
        print('أسلوب الملفات:', 'سليم' if not style else style[:10])
        res = recalc_check(args.out)
        if res:
            errors, s = res
            print('فحص المعادلات:', 'لا توجد أخطاء' if not errors else errors[:10])
            for i in range(6, 16):
                label, v = s.cell(row=i, column=1).value, s.cell(row=i, column=2).value
                if label:
                    print(f'   {label}: {v:,.0f}' if isinstance(v, (int, float)) and abs(v) > 10 else f'   {label}: {v}')
            if errors:
                return 1
    if args.pdf:
        print('PDF:', export_pdf(args.out))
    return 0


if __name__ == '__main__':
    sys.exit(main())
