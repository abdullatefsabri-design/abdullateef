#!/usr/bin/env python3
"""مراجعة ملف Excel يرسله المستخدم (تكاليف أو تسعير) قبل مناقشته.

يقرأ الملف بالمعادلات والقيم المحفوظة معاً ويطلع:
  1. خريطة الملف: الأوراق، والعناوين، وعدد المعادلات والمدخلات، والمخفي، والروابط الخارجية.
  2. ملاحظات مصنفة (خطأ، تحذير، ملاحظة) بمرجع الخلية والاقتراح:
     أخطاء المعادلات، وأرقام مكتوبة بدل معادلات، ومعادلات مختلفة عن عمودها، وأرقام ثابتة داخل
     المعادلات (سعر صرف أو خصم)، وأكثر من سعر صرف، وغياب خصم الوكلاء، والهامش على الكلفة بدل
     السعر، ومجاميع ناقصة، وأرقام مخزونة كنص، وإشارات لخلايا فارغة.

الاستخدام:
  python review_workbook.py "ملف_التسعير.xlsx" [--fx 1520] [--fx-market 1668] [--discount 0.14] [--out "مراجعة_<الاسم>.xlsx"]

الملف الأصلي لا يُعدّل. يقبل xlsx وxlsm، وxls وods عبر LibreOffice إذا موجود.
الناتج المطبوع ملخص للنقاش، و--out ينشئ ملفاً بورقتين: الملاحظات (بعمود القرار) وخريطة الملف.
"""
import argparse
import re
import shutil
import subprocess
import sys
import tempfile
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_cost_workbook import (CENTER, F_BODY, F_INPUT, F_NOTE, F_SUB, F_TITLE, FILL_AMBER, FILL_RED,  # noqa: E402
                                 FILL_SECTION, RIGHT, house_style_issues, page_setup, set_widths, style_body,
                                 style_header)

ERR, WARN, NOTE = 'خطأ', 'تحذير', 'ملاحظة'
SEVERITY_ORDER = {ERR: 0, WARN: 1, NOTE: 2}
ERROR_VALUES = ('#REF!', '#DIV/0!', '#VALUE!', '#N/A', '#NAME?', '#NUM!', '#NULL!')
FX_WORDS = ('صرف', 'دولار', 'usd', 'rate', '$')
DISCOUNT_WORDS = ('خصم', 'discount')
MARGIN_WORDS = ('هامش', 'ربح', 'margin', 'markup')
FX_RANGE = (1000, 2000)
CELL_RE = re.compile(r"(\$?)([A-Z]{1,3})(\$?)(\d+)")
NUMTEXT_RE = re.compile(r'^\s*\(?-?[\d٠-٩][\d٠-٩,٬\.]*\)?\s*%?\s*$')
SUM_RE = re.compile(r"SUM\(\$?([A-Z]{1,3})\$?(\d+):\$?\1\$?(\d+)\)", re.I)
MARKUP_RE = re.compile(r"\(\s*\$?([A-Z]{1,3})\$?(\d+)\s*-\s*\$?([A-Z]{1,3})\$?(\d+)\s*\)\s*/\s*\$?([A-Z]{1,3})\$?(\d+)")
ARABIC_DIGITS = str.maketrans('٠١٢٣٤٥٦٧٨٩', '0123456789')


def soffice():
    return shutil.which('soffice') or shutil.which('libreoffice')


def as_xlsx(path, tmp):
    """xls وods تتحول لـ xlsx بنسخة مؤقتة. الأصل لا يُمس."""
    if path.suffix.lower() in ('.xlsx', '.xlsm'):
        return path
    exe = soffice()
    if not exe:
        sys.exit(f'الملف {path.suffix} يحتاج LibreOffice للتحويل، أو احفظه xlsx')
    subprocess.run([exe, '--headless', '--calc', '--convert-to', 'xlsx', '--outdir', tmp, str(path)],
                   check=True, capture_output=True, timeout=180)
    return Path(tmp) / (path.stem + '.xlsx')


def recalculated_values(path, tmp):
    """إذا الملف محفوظ بدون قيم محسوبة (مولد آلياً)، نعيد حسابه بنسخة مؤقتة."""
    exe = soffice()
    if not exe:
        return None
    out = Path(tmp) / 'recalc'
    out.mkdir(exist_ok=True)
    subprocess.run([exe, '--headless', '--calc', '--convert-to', 'xlsx', '--outdir', str(out), str(path)],
                   check=True, capture_output=True, timeout=180)
    return load_workbook(out / (path.stem + '.xlsx'), data_only=True)


def fmt_num(v):
    return f'{v:,.0f}' if abs(v) >= 100 else f'{v:g}'


def is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


class Review:
    def __init__(self, wb_f, wb_v, fx, discount, fx_market=None):
        self.wb_f, self.wb_v, self.fx, self.discount, self.fx_market = wb_f, wb_v, fx, discount, fx_market
        self.findings = []
        self.sheet_map = []
        self.fx_seen = defaultdict(list)       # القيمة -> الخلايا
        self.discount_seen = []

    def add(self, sev, ws, cell, kind, detail, fix):
        self.findings.append([sev, ws if isinstance(ws, str) else ws.title, cell, kind, detail, fix])

    # ---------- سياق الخلية ----------
    def header_row(self, ws):
        for r in range(1, min(ws.max_row, 15) + 1):
            texts = [c for c in ws[r] if isinstance(c.value, str) and c.value.strip() and not str(c.value).startswith('=')]
            if len(texts) >= 2:
                return r
        return None

    def context(self, ws, cell, hdr):
        parts = []
        for c in ws[cell.row]:
            if isinstance(c.value, str) and not c.value.startswith('='):
                parts.append(c.value)
        if hdr:
            h = ws.cell(row=hdr, column=cell.column).value
            if isinstance(h, str):
                parts.append(h)
        return ' '.join(parts).lower()

    # ---------- الفحوصات ----------
    def run(self):
        if getattr(self.wb_f, '_external_links', None):
            self.add(WARN, '(الملف)', '', 'روابط خارجية',
                     f'الملف مربوط بـ {len(self.wb_f._external_links)} ملف خارجي، وقيمه تتغير إذا تغير ذاك الملف أو انفقد',
                     'نسخ القيم المطلوبة لورقة مدخلات داخل الملف مع ذكر مصدرها')
        for ws in self.wb_f.worksheets:
            self.sheet(ws)
        self.cross_sheet()
        self.findings.sort(key=lambda f: (SEVERITY_ORDER[f[0]], f[1]))
        return self.findings

    def sheet(self, ws):
        wv = self.wb_v[ws.title]
        hdr = self.header_row(ws)
        n_formula = n_const = n_text = 0
        hard = defaultdict(list)          # (العمود، الرقم) -> خلايا
        numtext = defaultdict(list)
        empty_refs = defaultdict(list)
        col_formulas = defaultdict(dict)  # العمود -> {الصف: النمط}
        col_consts = defaultdict(dict)

        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v is None:
                    continue
                if isinstance(v, str) and v.startswith('='):
                    n_formula += 1
                    cached = wv[c.coordinate].value
                    if isinstance(cached, str) and cached in ERROR_VALUES or '#REF!' in v:
                        self.add(ERR, ws, c.coordinate, 'خطأ معادلة', f'النتيجة {cached or "#REF!"} بالمعادلة {v[:80]}',
                                 'تصحيح المرجع أو حماية القسمة بـ IFERROR بعد معرفة السبب')
                    if '[' in v and ']' in v:
                        self.add(WARN, ws, c.coordinate, 'رابط خارجي', v[:80], 'استبداله بقيمة موثقة بورقة المدخلات')
                    col_formulas[c.column][c.row] = self.pattern(v, c)
                    self.formula_numbers(ws, c, v, hdr, hard)
                    self.formula_refs(ws, c, v, empty_refs)
                    self.sum_range(ws, wv, c, v)
                    self.markup(ws, c, v, hdr)
                elif is_num(v):
                    n_const += 1
                    col_consts[c.column][c.row] = v
                    ctx = self.context(ws, c, hdr)
                    if FX_RANGE[0] <= v <= FX_RANGE[1] and any(w in ctx for w in FX_WORDS):
                        self.fx_seen[v].append(f'{ws.title}!{c.coordinate}')
                    if any(w in ctx for w in DISCOUNT_WORDS) and (abs(v - self.discount) < 1e-9 or abs(v - self.discount * 100) < 1e-9
                                                                   or 0 < v < 1 or 1 <= v <= 50):
                        self.discount_seen.append((f'{ws.title}!{c.coordinate}', v))
                elif isinstance(v, str):
                    n_text += 1
                    if NUMTEXT_RE.match(v) and any(ch.isdigit() for ch in v.translate(ARABIC_DIGITS)):
                        numtext[c.column].append(c.coordinate)
                    if any(w in v.lower() for w in DISCOUNT_WORDS):
                        self.discount_seen.append((f'{ws.title}!{c.coordinate}', None))

        self.column_consistency(ws, col_formulas, col_consts)
        for (col, num), cells in hard.items():
            kind, sev, fix = self.classify_number(num)
            more = f' و{len(cells) - 1} خلية مشابهة' if len(cells) > 1 else ''
            self.add(sev, ws, cells[0], kind, f'الرقم {fmt_num(num)} مكتوب داخل المعادلة بالعمود {get_column_letter(col)}{more}', fix)
        for col, cells in numtext.items():
            more = f' و{len(cells) - 1} خلية بنفس العمود' if len(cells) > 1 else ''
            self.add(WARN, ws, cells[0], 'رقم مخزون كنص', f'القيمة نص مو رقم، فالمجاميع والمعادلات تتجاهلها{more}',
                     'تحويلها لأرقام (Text to Columns أو ضرب بـ 1)')
        for target, cells in empty_refs.items():
            more = f' و{len(cells) - 1} معادلة ثانية' if len(cells) > 1 else ''
            self.add(NOTE, ws, cells[0], 'إشارة لخلية فارغة', f'المعادلة تأخذ {target} وهي فارغة، فتنحسب صفراً{more}',
                     'التأكد أن الفراغ مقصود، وإلا تعبئة القيمة')

        hidden_rows = sum(1 for d in ws.row_dimensions.values() if d.hidden)
        hidden_cols = [k for k, d in ws.column_dimensions.items() if d.hidden]
        if ws.sheet_state != 'visible':
            self.add(NOTE, ws, '', 'ورقة مخفية', f'الورقة {ws.title} مخفية ({ws.sheet_state})', 'مراجعة محتواها لأن معادلات ثانية ممكن تعتمد عليها')
        if hidden_rows or hidden_cols:
            self.add(NOTE, ws, '', 'صفوف أو أعمدة مخفية', f'{hidden_rows} صف و{len(hidden_cols)} عمود مخفي',
                     'إظهارها قبل الاعتماد حتى ما يفوت بند')
        headers = []
        if hdr:
            headers = [str(c.value).strip() for c in ws[hdr] if isinstance(c.value, str) and c.value.strip()][:20]
        self.sheet_map.append([ws.title, 'ظاهرة' if ws.sheet_state == 'visible' else 'مخفية', ws.dimensions,
                               hdr or '', n_formula, n_const, n_text, len(ws.merged_cells.ranges), ' | '.join(headers)])

    def pattern(self, formula, cell):
        """المعادلة بمراجع نسبية حتى نقارن معادلات العمود الواحد."""
        def rel(m):
            col = m.group(2) if m.group(1) else f'C[{column_index_from_string(m.group(2)) - cell.column}]'
            row = m.group(4) if m.group(3) else f'R[{int(m.group(4)) - cell.row}]'
            return f'{col}{row}'
        return CELL_RE.sub(rel, formula)

    def column_consistency(self, ws, col_formulas, col_consts):
        for col, rows in col_formulas.items():
            if len(rows) < 3:
                continue
            common, freq = Counter(rows.values()).most_common(1)[0]
            if freq / len(rows) < 0.6:
                continue
            first, last = min(rows), max(rows)
            for r, p in rows.items():
                if p != common:
                    coord = f'{get_column_letter(col)}{r}'
                    self.add(WARN, ws, coord, 'معادلة مختلفة عن عمودها',
                             f'{ws[coord].value[:70]} بينما أغلب العمود ({freq} من {len(rows)}) بنفس النمط الثاني',
                             'التأكد أن الاختلاف مقصود، وإلا نسخ معادلة العمود')
            for r, v in col_consts.get(col, {}).items():
                if first < r < last and (r - 1 in rows or r + 1 in rows):
                    self.add(ERR, ws, f'{get_column_letter(col)}{r}', 'رقم مكتوب بدل معادلة',
                             f'القيمة {fmt_num(v)} مكتوبة يدوياً وسط عمود معادلات، فما تتحدث إذا تغيرت المدخلات',
                             'إرجاع معادلة العمود، أو نقل الرقم لورقة المدخلات مع سببه')

    def formula_numbers(self, ws, cell, formula, hdr, hard):
        try:
            tokens = Tokenizer(formula).items
        except Exception:
            return
        for t in tokens:
            if t.type == 'OPERAND' and t.subtype == 'NUMBER':
                try:
                    num = float(t.value)
                except ValueError:
                    continue
                if num in (0, 1, 2, 3, 4, 10, 12, 100, 1000, 365, 30, 7) or num < 0:
                    continue
                hard[(cell.column, num)].append(cell.coordinate)

    def classify_number(self, num):
        if FX_RANGE[0] <= num <= FX_RANGE[1]:
            self.fx_seen[num].append('(داخل معادلة)')
            sev = WARN if self.known_rate(num) else ERR
            return ('سعر صرف داخل معادلة', sev,
                    f'نقل سعر الصرف لخلية مدخلات واحدة وربط المعادلات بيها. السعر الرسمي من 2026-10-07 هو {fmt_num(self.fx)}')
        if any(abs(num - x) < 1e-9 for x in (self.discount, 1 - self.discount, 1 + self.discount, self.discount * 100)):
            self.discount_seen.append(('(داخل معادلة)', num))
            return 'خصم داخل معادلة', WARN, 'نقل نسبة الخصم لخلية مدخلات واحدة'
        return 'رقم ثابت داخل معادلة', NOTE, 'نقل الرقم لخلية مدخلات بعنوان واضح ومصدر'

    def known_rate(self, v):
        return abs(v - self.fx) <= 0.5 or (self.fx_market is not None and abs(v - self.fx_market) <= 15)

    def formula_refs(self, ws, cell, formula, empty_refs):
        try:
            tokens = Tokenizer(formula).items
        except Exception:
            return
        for t in tokens:
            if t.type != 'OPERAND' or t.subtype != 'RANGE' or ':' in t.value or '!' in t.value:
                continue
            m = CELL_RE.fullmatch(t.value.replace('$', '').strip())
            if not m:
                continue
            target = f'{m.group(2)}{m.group(4)}'
            try:
                if ws[target].value is None and not any(target in r for r in ws.merged_cells.ranges):
                    empty_refs[target].append(cell.coordinate)
            except ValueError:
                continue

    def sum_range(self, ws, wv, cell, formula):
        for m in SUM_RE.finditer(formula):
            col, a, b = m.group(1).upper(), int(m.group(2)), int(m.group(3))
            if col != get_column_letter(cell.column) or cell.row <= b:
                continue
            missed = [r for r in range(b + 1, cell.row) if is_num(wv[f'{col}{r}'].value) or
                      (isinstance(ws[f'{col}{r}'].value, (int, float)))]
            if missed:
                self.add(ERR, ws, cell.coordinate, 'مجموع ناقص',
                         f'{m.group(0)} ما يشمل الصفوف {", ".join(map(str, missed[:5]))} وبيها أرقام',
                         f'توسيع المدى لحد {col}{cell.row - 1}')

    def markup(self, ws, cell, formula, hdr):
        m = MARKUP_RE.search(formula.replace(' ', ''))
        if not m:
            return
        minus, denom = f'{m.group(3)}{m.group(4)}', f'{m.group(5)}{m.group(6)}'
        label = self.context(ws, cell, hdr)
        if minus == denom and any(w in label for w in MARGIN_WORDS):
            self.add(NOTE, ws, cell.coordinate, 'الهامش على الكلفة',
                     f'{formula[:60]} يقسم على الكلفة (markup)، مثلاً 20% على الكلفة = 16.7% على السعر',
                     'توحيد التعريف: الهامش على صافي السعر بعد الخصم')

    def cross_sheet(self):
        fx_values = sorted(v for v in self.fx_seen)
        market = f'، والسوق الموازي حوالي {fmt_num(self.fx_market)}' if self.fx_market else ''
        if len(fx_values) > 1:
            where = '؛ '.join(f'{fmt_num(v)} ({", ".join(self.fx_seen[v][:3])})' for v in fx_values)
            in_cells = all(not c.startswith('(') for v in fx_values for c in self.fx_seen[v])
            two_known = len(fx_values) == 2 and all(self.known_rate(v) for v in fx_values)
            self.add(WARN if (two_known and in_cells) else ERR, '(الملف)', '', 'أكثر من سعر صرف',
                     f'الملف يستخدم {len(fx_values)} أسعار: {where}',
                     f'كل سعر بخلية مدخلات واحدة بعنوان واضح (رسمي أو سوق) ويبين وين ينطبق. الرسمي من 2026-10-07 هو '
                     f'{fmt_num(self.fx)} (قرار 544){market}')
        elif fx_values and not self.known_rate(fx_values[0]):
            v = fx_values[0]
            self.add(WARN, '(الملف)', self.fx_seen[v][0], 'سعر صرف قديم',
                     f'الملف على {fmt_num(v)}، والفرق للسعر الجديد {(self.fx / v - 1):.1%} على كل دولار',
                     f'تحديثه إلى {fmt_num(self.fx)} (الرسمي){market}، حسب مصدر الدولار لكل مادة')
        elif not fx_values:
            self.add(NOTE, '(الملف)', '', 'سعر الصرف غير ظاهر',
                     'ما لگيت سعر صرف بعنوان واضح، فالكلف بالدولار ممكن محولة يدوياً أو مكتوبة بالدينار',
                     'إضافة خلية سعر صرف وعمود عملة لكل مادة مستوردة')
        if not self.discount_seen:
            self.add(WARN, '(الملف)', '', 'خصم الوكلاء غير موجود',
                     f'ما بالملف خصم {self.discount:.0%}، فالهامش محسوب على سعر القائمة ويطلع أعلى من الحقيقي',
                     f'صافي السعر = سعر القائمة × {1 - self.discount:g}، والهامش يُحسب عليه')


def print_report(path, findings, sheet_map, fx_seen):
    print(f'# مراجعة {path.name}\n')
    print('## خريطة الملف')
    for name, state, dims, hdr, nf, nc, nt, _, heads in sheet_map:
        print(f'- {name} ({state}، {dims}): {nf} معادلة، {nc} رقم مدخل، {nt} نص. صف العناوين {hdr or "غير واضح"}')
        if heads:
            print(f'  العناوين: {heads[:200]}')
    if fx_seen:
        print('\nأسعار الصرف الظاهرة: ' + '، '.join(f'{fmt_num(v)} ({len(c)})' for v, c in sorted(fx_seen.items())))
    counts = Counter(f[0] for f in findings)
    print(f'\n## الملاحظات: {counts.get(ERR, 0)} خطأ، {counts.get(WARN, 0)} تحذير، {counts.get(NOTE, 0)} ملاحظة')
    for i, (sev, ws, cell, kind, detail, fix) in enumerate(findings[:40], start=1):
        where = f'{ws}!{cell}' if cell else ws
        print(f'{i}. [{sev}] {where}: {kind}. {detail}. الاقتراح: {fix}')
    if len(findings) > 40:
        print(f'... و{len(findings) - 40} ملاحظة ثانية بالملف')


def write_report(out, path, findings, sheet_map):
    wb = Workbook()
    ws = wb.active
    ws.title = 'الملاحظات'
    ws.merge_cells('A1:I1')
    ws['A1'], ws['A1'].font, ws['A1'].alignment = f'مراجعة ملف {path.name}', F_TITLE, CENTER
    ws.merge_cells('A2:I2')
    counts = Counter(f[0] for f in findings)
    ws['A2'] = f'{counts.get(ERR, 0)} خطأ، {counts.get(WARN, 0)} تحذير، {counts.get(NOTE, 0)} ملاحظة. عمود القرار يُعبأ أثناء النقاش.'
    ws['A2'].font, ws['A2'].alignment = F_SUB, CENTER
    heads = ['#', 'الدرجة', 'الورقة', 'الخلية', 'النوع', 'التفصيل', 'الاقتراح', 'القرار', 'ملاحظة المستخدم']
    for j, h in enumerate(heads, start=1):
        ws.cell(row=4, column=j, value=h)
    style_header(ws, 4, 1, len(heads))
    wrap = RIGHT
    for i, f in enumerate(findings, start=5):
        for j, v in enumerate([i - 4, *f, None, None], start=1):
            style_body(ws.cell(row=i, column=j, value=v), font=F_INPUT if j >= 8 else F_BODY,
                       align=CENTER if j in (1, 2, 4, 8) else wrap)
    last = max(len(findings) + 4, 5)
    dv = DataValidation(type='list', formula1='"مقبول,مرفوض,للنقاش"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f'H5:H{last}')
    ws.conditional_formatting.add(f'B5:B{last}', FormulaRule(formula=[f'B5="{ERR}"'], fill=FILL_RED))
    ws.conditional_formatting.add(f'B5:B{last}', FormulaRule(formula=[f'B5="{WARN}"'], fill=FILL_AMBER))
    set_widths(ws, [5, 9, 16, 9, 20, 60, 50, 11, 30])
    ws.freeze_panes = 'A5'

    sm = wb.create_sheet('خريطة الملف')
    heads = ['الورقة', 'الحالة', 'المدى', 'صف العناوين', 'معادلات', 'أرقام مدخلة', 'نصوص', 'خلايا مدمجة', 'العناوين']
    for j, h in enumerate(heads, start=1):
        sm.cell(row=1, column=j, value=h)
    style_header(sm, 1, 1, len(heads))
    for i, r in enumerate(sheet_map, start=2):
        for j, v in enumerate(r, start=1):
            style_body(sm.cell(row=i, column=j, value=v), align=CENTER if 2 <= j <= 8 else RIGHT)
    set_widths(sm, [22, 9, 14, 11, 10, 11, 9, 11, 90])
    sm.freeze_panes = 'A2'
    n = len(sheet_map) + 3
    sm.cell(row=n, column=1, value='الملف الأصلي لم يُعدّل. المراجعة على نسخة للقراءة.').font = F_NOTE

    for w in wb.worksheets:
        page_setup(w, 'شركة البركة للاستثمارات الصناعية', '4:4' if w.title == 'الملاحظات' else '1:1')
    wb.properties.creator = wb.properties.lastModifiedBy = 'الإدارة المالية'
    wb.save(out)
    issues = house_style_issues(out)
    if issues:
        print('تنبيه أسلوب:', issues[:5])


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('file', type=Path)
    ap.add_argument('--fx', type=float, default=1520, help='سعر الصرف المرجعي (افتراضياً 1520، قرار 544)')
    ap.add_argument('--fx-market', type=float, help='سعر السوق الموازي الحالي، حتى لا يُعتبر سعراً قديماً')
    ap.add_argument('--discount', type=float, default=0.14, help='خصم الوكلاء (افتراضياً 0.14)')
    ap.add_argument('--out', type=Path, help='ملف Excel للملاحظات وخريطة الملف')
    args = ap.parse_args()

    with tempfile.TemporaryDirectory() as tmp:
        src = as_xlsx(args.file, tmp)
        wb_f = load_workbook(src)
        wb_v = load_workbook(src, data_only=True)
        has_formulas = any(isinstance(c.value, str) and c.value.startswith('=')
                           for ws in wb_f.worksheets for row in ws.iter_rows() for c in row)
        has_cache = any(c.value is not None for ws in wb_v.worksheets for row in ws.iter_rows() for c in row
                        if isinstance(wb_f[ws.title][c.coordinate].value, str) and str(wb_f[ws.title][c.coordinate].value).startswith('='))
        if has_formulas and not has_cache:
            recalced = recalculated_values(src, tmp)
            if recalced:
                wb_v = recalced
                print('(الملف محفوظ بدون قيم محسوبة، فانعاد حسابه بنسخة مؤقتة)\n')
        review = Review(wb_f, wb_v, args.fx, args.discount, args.fx_market)
        findings = review.run()
    print_report(args.file, findings, review.sheet_map, review.fx_seen)
    if args.out:
        write_report(args.out, args.file, findings, review.sheet_map)
        print(f'\nالملف: {args.out}')
    return 0


if __name__ == '__main__':
    sys.exit(main())
