#!/usr/bin/env python3
"""يبني ملف التسعيرة الجديدة بمعادلات حية: الكلفة الواصلة لكل منتج على سعر الصرف (رسمي وسوق ويدوي)
والكمرك والضرائب والتضخم، والسعر المطلوب بعد خصم الوكلاء وكلفة الآجل والهامش المستهدف، وجسر الكلفة
(سعر المواد، الصرف والكمرك، التضخم)، والمقارنة بكلفة المستورد المنافس، وخانة سعر يدوي لكل منتج.

الاستخدام:
  python build_pricing.py --products products.csv --bom bom.csv --materials materials.csv \
      --margin 0.15 --conversion-per-kg 450 --conversion-basis "مصاريف التحويل أيلول ÷ الكغ المنتجة" \
      --fx-old 1450 --fx-market 1668 --fx-market-basis "بورصة الكفاح 2026-10-17" --market-share 0.3 \
      [--fx-official 1520] [--inflation 0.033] [--months 6] [--scrap 0.02] [--dso 91] [--round-to 250] \
      [--customs-raw 0.02] [--import-other 0] [--clearance 0.01] [--customs-fg 0.2] [--sales-tax 0] \
      [--review "مراجعة_<الملف>.xlsx"] --out "التسعيرة_الجديدة_<التاريخ>.xlsx" --check --pdf

كل الأسعار والنسب بخلايا زرقاء بورقة الافتراضات أو بالأعمدة الزرقاء، وتنعدل يدوياً بالملف.

الملفات (عناوين عربية أو إنكليزية):
  products : code/الرمز، name/المنتج، unit/الوحدة، list_price/سعر_القائمة_الحالي. اختياري: category/الفئة،
             kg_per_unit/كغ_للوحدة، volume/الكمية_الشهرية، conversion_per_kg/كلفة_التحويل_للكيلو،
             manual_price/السعر_اليدوي، import_price_usd/سعر_المستورد_بالدولار
  bom      : product_code/رمز_المنتج، component_code/رمز_المادة، qty/الكمية، unit/الوحدة (كغ أو غم أو قطعة)
  materials: code/الرمز، name/المادة، unit/الوحدة، currency/العملة (دولار أو دينار)، new_price/السعر_الجديد.
             اختياري: old_price/السعر_السابق، inflates/يتأثر_بالتضخم، fx_source/مصدر_الدولار (رسمي، سوق، مرجح، يدوي)،
             manual_fx/سعر_صرف_يدوي، direct_import/استيراد_مباشر، customs_pct/الرسم_الكمركي،
             other_pct/رسوم_أخرى، clearance_pct/التخليص، source/المصدر

القواعد: التضخم على الكلف المحلية فقط (مواد الدينار المؤشرة وكلفة التحويل). المواد بالدولار يحركها سعر الصرف وسعرها
بالدولار، حتى لا يُحسب الأثر مرتين. الكمرك على القيمة بسعر صرف الكمرك (افتراضياً الرسمي). الهامش على صافي السعر
بعد الخصم وكلفة الآجل، وقبل ضريبة المبيعات.
"""
import argparse
import csv
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_cost_workbook import (CENTER, F_BODY, F_BOLD, F_INPUT, F_NOTE, F_SECTION, F_SUB, F_TITLE,  # noqa: E402
                                 FILL_AMBER, FILL_GREEN, FILL_RED, FMT_PCT, RIGHT, export_pdf, house_style_issues,
                                 page_setup, recalc_check, set_widths, style_body, style_header, style_total)
from validate_costs import parse_number  # noqa: E402

COMPANY = 'شركة البركة للاستثمارات الصناعية'
FMT_IQD = '#,##0;[Red](#,##0);"-"'
FMT_DEC = '#,##0.00;[Red](#,##0.00);"-"'
FMT_KG = '#,##0.000;[Red](#,##0.000);"-"'
FMT_PCT_SIGNED = '0.0%;[Red]-0.0%;"-"'
WRAP = Alignment(horizontal='right', vertical='center', wrap_text=True)
S_ASSUME, S_SUM, S_PRICE, S_BOM, S_MAT, S_SCEN, S_REVIEW, S_METHOD = (
    'الافتراضات', 'الملخص', 'التسعيرة الجديدة', 'الوصفات', 'المواد', 'سيناريوهات الصرف', 'ملاحظات المراجعة', 'المنهجية والاعتماد')
USD, IQD = 'دولار', 'دينار'
YES, NO = 'نعم', 'لا'
BASIS_TARGET, BASIS_KEEP = 'الهامش المستهدف', 'الهامش السابق'
SRC_OFF, SRC_MKT, SRC_MIX, SRC_MAN = 'رسمي', 'سوق', 'مرجح', 'يدوي'

ALIASES = {
    'code': ('code', 'الرمز', 'رمز', 'item_code'),
    'name': ('name', 'المنتج', 'الاسم', 'المادة', 'اسم'),
    'unit': ('unit', 'الوحدة', 'وحدة'),
    'category': ('category', 'الفئة', 'فئة'),
    'kg_per_unit': ('kg_per_unit', 'كغ_للوحدة', 'الوزن'),
    'list_price': ('list_price', 'سعر_القائمة_الحالي', 'سعر_القائمة', 'السعر_الحالي', 'السعر'),
    'volume': ('volume', 'الكمية_الشهرية'),
    'conversion_per_kg': ('conversion_per_kg', 'كلفة_التحويل_للكيلو'),
    'manual_price': ('manual_price', 'السعر_اليدوي'),
    'import_price_usd': ('import_price_usd', 'سعر_المستورد_بالدولار'),
    'product_code': ('product_code', 'رمز_المنتج'),
    'component_code': ('component_code', 'رمز_المادة', 'المكون'),
    'qty': ('qty', 'الكمية', 'quantity'),
    'currency': ('currency', 'العملة'),
    'new_price': ('new_price', 'السعر_الجديد', 'price', 'السعر'),
    'old_price': ('old_price', 'السعر_السابق'),
    'inflates': ('inflates', 'يتأثر_بالتضخم'),
    'fx_source': ('fx_source', 'مصدر_الدولار'),
    'manual_fx': ('manual_fx', 'سعر_صرف_يدوي'),
    'direct_import': ('direct_import', 'استيراد_مباشر'),
    'customs_pct': ('customs_pct', 'الرسم_الكمركي'),
    'other_pct': ('other_pct', 'رسوم_أخرى'),
    'clearance_pct': ('clearance_pct', 'التخليص'),
    'source': ('source', 'المصدر'),
}
KG_UNITS = {'كغ', 'كيلو', 'kg', 'كيلوغرام'}
G_UNITS = {'غم', 'g', 'غرام', 'gm'}
TRUE_WORDS = {'نعم', 'yes', 'y', '1', 'true'}


def read_rows(path, fields):
    with open(path, encoding='utf-8-sig', newline='') as f:
        rows = list(csv.DictReader(f))
    if not rows:
        sys.exit(f'{path}: فارغ')
    keys = {k.strip().lower(): k for k in rows[0]}
    mapping = {}
    for field in fields:
        for alias in ALIASES[field]:
            if alias.lower() in keys and keys[alias.lower()] not in mapping.values():
                mapping[field] = keys[alias.lower()]
                break
    return [{f: (r.get(src) or '').strip() for f, src in mapping.items()} for r in rows], mapping


def currency(v):
    v = (v or '').strip().lower()
    if v in ('usd', 'دولار', 'دولار أميركي', 'دولار امريكي', '$'):
        return USD
    if v in ('iqd', 'دينار', 'دينار عراقي', ''):
        return IQD
    sys.exit(f'عملة غير معروفة: {v}')


def flag(v, default):
    v = (v or '').strip().lower()
    return (YES if v in TRUE_WORDS else NO) if v else default


def load_inputs(args):
    products, pm = read_rows(args.products, ['code', 'name', 'unit', 'category', 'kg_per_unit', 'list_price', 'volume',
                                             'conversion_per_kg', 'manual_price', 'import_price_usd'])
    bom, _ = read_rows(args.bom, ['product_code', 'component_code', 'qty', 'unit'])
    mats, _ = read_rows(args.materials, ['code', 'name', 'unit', 'currency', 'fx_source', 'manual_fx', 'direct_import',
                                         'customs_pct', 'other_pct', 'clearance_pct', 'new_price', 'old_price', 'inflates', 'source'])
    for need in ('code', 'name'):
        if need not in pm:
            sys.exit(f'products: العمود {need} مطلوب')

    problems = []
    mat_by_code = {}
    for m in mats:
        m['currency'] = currency(m.get('currency'))
        m['new_price'] = parse_number(m.get('new_price'))
        if m['new_price'] is None:
            problems.append(f'المادة {m["code"]}: بدون سعر جديد')
            continue
        old = parse_number(m.get('old_price'))
        m['old_price'] = m['new_price'] if old is None else old
        is_usd = m['currency'] == USD
        m['inflates'] = flag(m.get('inflates'), NO if is_usd else YES)
        m['direct_import'] = flag(m.get('direct_import'), YES if is_usd else NO)
        unit = (m.get('unit') or 'كغ').strip()
        m['unit'] = 'كغ' if unit.lower() in KG_UNITS else unit
        src = (m.get('fx_source') or '').strip().lower()
        m['fx_source'] = {'رسمي': SRC_OFF, 'official': SRC_OFF, 'سوق': SRC_MKT, 'market': SRC_MKT,
                          'يدوي': SRC_MAN, 'manual': SRC_MAN}.get(src, SRC_MIX)
        m['manual_fx'] = parse_number(m.get('manual_fx'))
        if m['fx_source'] == SRC_MAN and m['manual_fx'] is None:
            problems.append(f'المادة {m["code"]}: مصدر الدولار يدوي بدون سعر صرف يدوي')
        for k in ('customs_pct', 'other_pct', 'clearance_pct'):
            m[k] = parse_number(m.get(k))
        mat_by_code[m['code']] = m

    lines = []
    for b in bom:
        qty = parse_number(b.get('qty'))
        mat = mat_by_code.get(b.get('component_code'))
        if mat is None:
            problems.append(f'الوصفة {b.get("product_code")}: المادة {b.get("component_code")} غير موجودة بملف المواد')
            continue
        if qty is None:
            problems.append(f'الوصفة {b.get("product_code")}/{b.get("component_code")}: كمية غير رقمية')
            continue
        unit = (b.get('unit') or mat['unit']).strip().lower()
        if unit in G_UNITS and mat['unit'] == 'كغ':
            qty /= 1000
        lines.append({'product': b['product_code'], 'component': b['component_code'], 'qty': qty})

    with_bom = {l['product'] for l in lines}
    for p in products:
        for k in ('list_price', 'kg_per_unit', 'volume', 'conversion_per_kg', 'manual_price', 'import_price_usd'):
            p[k] = parse_number(p.get(k))
        if p['code'] not in with_bom:
            problems.append(f'المنتج {p["code"]} ({p["name"]}): بدون وصفة، فكلفة مواده صفر')
    if problems:
        print('مشاكل بالمدخلات، صحّحها قبل البناء:')
        for x in problems[:40]:
            print('  -', x)
        sys.exit(1)
    return products, lines, list(mat_by_code.values())


def write_table(ws, cols, rows, center=()):
    """cols: [(المفتاح، العنوان، التنسيق)]، rows: [{المفتاح: (القيمة، النوع)}]. النوع in مدخل أزرق، bold عريض."""
    for j, (_, h, _) in enumerate(cols, start=1):
        ws.cell(row=1, column=j, value=h)
    style_header(ws, 1, 1, len(cols))
    ws.row_dimensions[1].height = 45
    for i, r in enumerate(rows, start=2):
        for j, (k, _, fmt) in enumerate(cols, start=1):
            v, kind = r[k]
            font = F_INPUT if kind == 'in' else (F_BOLD if kind == 'bold' else F_BODY)
            style_body(ws.cell(row=i, column=j, value=v), fmt, font=font, align=CENTER if k in center else RIGHT)


def letters(cols):
    return {k: get_column_letter(j) for j, (k, _, _) in enumerate(cols, start=1)}


def build(args, products, lines, mats):
    wb = Workbook()
    wb.remove(wb.active)
    A = {}

    # ---------- الافتراضات ----------
    ws = wb.create_sheet(S_ASSUME)
    ws.merge_cells('A1:C1')
    ws['A1'], ws['A1'].font, ws['A1'].alignment = COMPANY, F_TITLE, CENTER
    ws.merge_cells('A2:C2')
    ws['A2'], ws['A2'].font, ws['A2'].alignment = 'افتراضات التسعيرة الجديدة', F_SUB, CENTER
    ws.merge_cells('A3:C3')
    ws['A3'] = 'الخلايا الزرقاء إدخال يدوي، وكل الأوراق تتحدث منها. الخلية الفارغة تنحسب صفراً.'
    ws['A3'].font, ws['A3'].alignment = F_NOTE, CENTER
    for j, h in enumerate(['البند', 'القيمة', 'الأساس والمصدر'], start=1):
        ws.cell(row=5, column=j, value=h)
    style_header(ws, 5, 1, 3)
    est = 'تقدير افتراضي، يلزم تأكيده من الإنتاج'
    customs_note = ('قرار مجلس الوزراء 957 (نافذ من بداية 2026): المواد الأولية الصناعية 0.5% إلى 2% حسب Shafaq. '
                    'غير مؤكد لرمز HS بعينه، يلزم التأكيد من المخلص الكمركي')
    items = [
        ('h_fx', 'سعر الصرف', None, None, None),
        ('fx_off', 'سعر الصرف الرسمي', args.fx_official, '#,##0', 'قرار مجلس الوزراء 544، نافذ من 2026-10-07 (البيع للمستفيد النهائي)'),
        ('fx_mkt', 'سعر صرف السوق الموازي', args.fx_market, '#,##0', args.fx_market_basis),
        ('mkt_share', 'نسبة الدولار المشترى من السوق', args.market_share, FMT_PCT, args.market_share_basis),
        ('fx_new', 'سعر الصرف الفعلي المرجح', None, '#,##0', 'تُحسب: الرسمي × (1 − نسبة السوق) + السوق × نسبة السوق. ينطبق على المواد المصنفة مرجح'),
        ('fx_old', 'سعر الصرف بالتسعيرة الحالية', args.fx_old, '#,##0', args.fx_old_basis),
        ('h_tax', 'الكمرك والضرائب', None, None, None),
        ('cust_fx', 'سعر الصرف لاحتساب الكمرك', None, '#,##0', 'افتراضياً الرسمي. يُكتب يدوياً إذا الكمرك يعتمد سعراً آخر'),
        ('cust_raw', 'الرسم الكمركي على المواد الأولية المستوردة', args.customs_raw, FMT_PCT, customs_note),
        ('tax_raw', 'رسوم وضرائب استيراد أخرى على المواد', args.import_other, FMT_PCT, 'يُدخل من بيان الكمرك أو المخلص'),
        ('clear', 'التخليص والنقل الداخلي للمواد المستوردة', args.clearance, FMT_PCT, 'نسبة من قيمة المادة، من فواتير التخليص والنقل'),
        ('cust_fg', 'الرسم الكمركي على الأنابيب والملحقات البلاستيكية المستوردة', args.customs_fg, FMT_PCT,
         'لمقارنة كلفة المستورد المنافس. غير مؤكد، يلزم التأكيد من جدول التعرفة (HS 3917)'),
        ('tax_fg', 'رسوم أخرى على المنتج المستورد', args.import_other_fg, FMT_PCT, 'يُدخل من بيان الكمرك أو المخلص'),
        ('imp_fx', 'سعر صرف المستورد المنافس', None, '#,##0', 'افتراضياً سعر السوق. يُكتب يدوياً إذا مختلف'),
        ('sales_tax', 'ضريبة المبيعات على منتجات الشركة', args.sales_tax, FMT_PCT,
         'ضريبة المبيعات بالعراق على سلع وخدمات محددة (iraq-law.md §2) ولا تشمل الأنابيب حسب المعلوم. يلزم التحقق'),
        ('income_tax', 'ضريبة الدخل على الشركات', args.income_tax, FMT_PCT, 'القانون 113 لسنة 1982 المعدل: 15% (iraq-law.md §1)'),
        ('h_cost', 'الكلفة والتضخم', None, None, None),
        ('infl', 'التضخم السنوي', args.inflation, FMT_PCT, args.inflation_basis),
        ('months', 'أشهر صلاحية التسعيرة', args.months, '0', 'المدة حتى المراجعة القادمة للأسعار'),
        ('pinfl', 'تضخم فترة الصلاحية', None, FMT_PCT, 'تُحسب: (1 + التضخم السنوي) أس (الأشهر ÷ 12)، ناقص 1'),
        ('cinfl', 'تضخم كلفة التحويل للفترة', None, FMT_PCT, 'افتراضياً نفس تضخم الفترة. يُكتب يدوياً إذا معروفة زيادة رواتب أو طاقة'),
        ('conv', 'كلفة التحويل للكيلو (دينار)', args.conversion_per_kg, '#,##0', args.conversion_basis),
        ('scrap', 'نسبة الهدر غير المسترد', args.scrap, FMT_PCT, args.scrap_basis or est),
        ('h_price', 'السعر والهامش', None, None, None),
        ('disc', 'خصم الوكلاء', args.discount, FMT_PCT, 'سياسة الخصم الحالية (نصف سنوي)'),
        ('margin', 'الهامش المستهدف بعد الخصم وكلفة الآجل', args.margin, FMT_PCT, args.margin_basis),
        ('dso', 'أيام التحصيل', args.dso, '0', 'متوسط أيام تحصيل الوكلاء، خط الأساس 2026-09-30'),
        ('cof', 'كلفة المال السنوية', None, FMT_PCT, 'افتراضياً بقدر التضخم. تُكتب يدوياً بكلفة التمويل الفعلية إن وجدت'),
        ('credit', 'كلفة الآجل من صافي السعر', None, FMT_PCT, 'تُحسب: (1 + كلفة المال) أس (أيام التحصيل ÷ 365)، ناقص 1'),
        ('round', 'تقريب السعر لأقرب (دينار)', args.round_to, '#,##0', 'سياسة التقريب بقائمة الأسعار'),
        ('basis', 'أساس السعر المقترح', BASIS_TARGET, None,
         f'{BASIS_TARGET}: السعر يحقق الهامش المستهدف. {BASIS_KEEP}: السعر يرتفع بقدر ارتفاع الكلفة حتى يبقى هامش كل منتج كما هو'),
        ('allow_down', 'السماح بتخفيض السعر', NO, None, 'لا: المنتج اللي هامشه أعلى من المطلوب يبقى بسعره الحالي'),
    ]
    R = {key: i for i, (key, *_) in enumerate(items, start=6)}
    for key, i in R.items():
        A[key] = f"'{S_ASSUME}'!$B${i}"
    formulas = {
        'fx_new': f'=B{R["fx_off"]}*(1-B{R["mkt_share"]})+B{R["fx_mkt"]}*B{R["mkt_share"]}',
        'cust_fx': f'=B{R["fx_off"]}',
        'imp_fx': f'=B{R["fx_mkt"]}',
        'pinfl': f'=(1+B{R["infl"]})^(B{R["months"]}/12)-1',
        'cinfl': f'=B{R["pinfl"]}',
        'cof': f'=B{R["infl"]}',
        'credit': f'=(1+B{R["cof"]})^(B{R["dso"]}/365)-1',
    }
    calc_only = {'fx_new', 'pinfl', 'credit'}   # معادلات بحتة، والباقي مدخل حتى لو بيه معادلة افتراضية
    for key, i in R.items():
        _, label, val, fmt, basis = items[i - 6]
        if key.startswith('h_'):
            ws.cell(row=i, column=1, value=label).font = F_SECTION
            continue
        style_body(ws.cell(row=i, column=1, value=label), font=F_BOLD, align=RIGHT)
        style_body(ws.cell(row=i, column=2, value=formulas.get(key, val)), fmt,
                   font=F_BODY if key in calc_only else F_INPUT, align=CENTER)
        style_body(ws.cell(row=i, column=3, value=basis), align=WRAP)
    for key, opts in (('basis', f'"{BASIS_TARGET},{BASIS_KEEP}"'), ('allow_down', f'"{YES},{NO}"')):
        dv = DataValidation(type='list', formula1=opts, allow_blank=False)
        ws.add_data_validation(dv)
        dv.add(f'B{R[key]}')
    set_widths(ws, [44, 16, 80])

    # ---------- المواد ----------
    mt = wb.create_sheet(S_MAT)
    mcols = [
        ('code', 'الرمز', None), ('name', 'المادة', None), ('unit', 'الوحدة', None), ('cur', 'العملة', None),
        ('src', 'مصدر الدولار', None), ('mfx', 'سعر صرف يدوي', '#,##0'), ('direct', 'استيراد مباشر', None),
        ('cust', 'الرسم الكمركي', FMT_PCT), ('other', 'رسوم أخرى', FMT_PCT), ('clear', 'التخليص والنقل', FMT_PCT),
        ('old', 'السعر السابق', FMT_DEC), ('new', 'السعر الجديد', FMT_DEC), ('infl', 'يتأثر بالتضخم', None),
        ('fx', 'سعر الصرف المطبق', '#,##0'), ('old_iqd', 'الواصل السابق بالدينار', FMT_IQD),
        ('new_base', 'الواصل الجديد قبل التضخم', FMT_IQD), ('infl_iqd', 'زيادة التضخم', FMT_IQD),
        ('new_iqd', 'الواصل الجديد بالدينار', FMT_IQD), ('price_eff', 'أثر سعر المادة', FMT_IQD),
        ('chg', 'التغير', FMT_PCT_SIGNED), ('source', 'المصدر', None),
    ]
    M = letters(mcols)
    rows = []
    for i, m in enumerate(mats, start=2):
        x = {k: f'{M[k]}{i}' for k in M}
        duty = f'IF({x["direct"]}="{YES}",{x["cust"]}+{x["other"]}+{x["clear"]},0)'
        old_fx = f'IF({x["cur"]}="{USD}",{A["fx_old"]},1)'
        cust_fx = f'IF({x["cur"]}="{USD}",{A["cust_fx"]},1)'
        pct = lambda k, default: (m[k], 'in') if m[k] is not None else (f'=IF({x["direct"]}="{YES}",{A[default]},0)', 'in')
        rows.append({
            'code': (m['code'], ''), 'name': (m.get('name'), ''), 'unit': (m['unit'], ''), 'cur': (m['currency'], 'in'),
            'src': (m['fx_source'] if m['currency'] == USD else '', 'in'), 'mfx': (m['manual_fx'], 'in'),
            'direct': (m['direct_import'], 'in'),
            'cust': pct('customs_pct', 'cust_raw'), 'other': pct('other_pct', 'tax_raw'), 'clear': pct('clearance_pct', 'clear'),
            'old': (m['old_price'], 'in'), 'new': (m['new_price'], 'in'), 'infl': (m['inflates'], 'in'),
            'fx': (f'=IF({x["cur"]}="{USD}",IF({x["src"]}="{SRC_OFF}",{A["fx_off"]},IF({x["src"]}="{SRC_MKT}",{A["fx_mkt"]},'
                   f'IF({x["src"]}="{SRC_MAN}",{x["mfx"]},{A["fx_new"]}))),1)', ''),
            'old_iqd': (f'={x["old"]}*{old_fx}*(1+{duty})', ''),
            'new_base': (f'={x["new"]}*{x["fx"]}+IF({x["direct"]}="{YES}",{x["new"]}*{cust_fx}*({x["cust"]}+{x["other"]})'
                         f'+{x["new"]}*{x["fx"]}*{x["clear"]},0)', ''),
            'infl_iqd': (f'=IF({x["infl"]}="{YES}",{x["new"]}*{x["fx"]}*{A["pinfl"]},0)', ''),
            'new_iqd': (f'={x["new_base"]}+{x["infl_iqd"]}', 'bold'),
            'price_eff': (f'=({x["new"]}-{x["old"]})*{old_fx}*(1+{duty})', ''),
            'chg': (f'=IF({x["old_iqd"]}=0,"",{x["new_iqd"]}/{x["old_iqd"]}-1)', ''),
            'source': (m.get('source'), ''),
        })
    write_table(mt, mcols, rows, center={'code', 'unit', 'cur', 'src', 'direct', 'infl'})
    n_mat = len(mats) + 1
    for k, opts in (('cur', f'"{USD},{IQD}"'), ('src', f'"{SRC_OFF},{SRC_MKT},{SRC_MIX},{SRC_MAN}"'),
                    ('direct', f'"{YES},{NO}"'), ('infl', f'"{YES},{NO}"')):
        dv = DataValidation(type='list', formula1=opts, allow_blank=True)
        mt.add_data_validation(dv)
        dv.add(f'{M[k]}2:{M[k]}{n_mat}')
    set_widths(mt, [13, 28, 7, 7, 9, 9, 8, 9, 9, 9, 10, 10, 8, 9, 12, 12, 10, 12, 11, 8, 30])
    mt.freeze_panes = 'C2'
    MAT = lambda k: f"'{S_MAT}'!${M[k]}$2:${M[k]}${n_mat}"

    # ---------- الوصفات ----------
    bm = wb.create_sheet(S_BOM)
    bcols = [
        ('prod', 'رمز المنتج', None), ('comp', 'رمز المادة', None), ('name', 'المادة', None), ('qty', 'الكمية للوحدة', FMT_KG),
        ('unit', 'الوحدة', None), ('cur', 'العملة', None), ('u_old', 'الواصل السابق للوحدة', FMT_IQD),
        ('u_new', 'الواصل الجديد للوحدة', FMT_IQD), ('u_infl', 'تضخم الوحدة', FMT_IQD), ('u_price', 'أثر السعر للوحدة', FMT_IQD),
        ('old', 'الكلفة السابقة', FMT_IQD), ('e_price', 'أثر سعر المادة', FMT_IQD), ('e_fx', 'أثر الصرف والكمرك', FMT_IQD),
        ('e_infl', 'أثر التضخم', FMT_IQD), ('new', 'الكلفة الجديدة', FMT_IQD), ('usd', 'منها بالدولار', FMT_IQD), ('kg', 'كغ', FMT_KG),
    ]
    B = letters(bcols)
    look = lambda k, i: f'=INDEX({MAT(k)},MATCH(${B["comp"]}{i},{MAT("code")},0))'
    rows = []
    for i, l in enumerate(lines, start=2):
        x = {k: f'{B[k]}{i}' for k in B}
        rows.append({
            'prod': (l['product'], ''), 'comp': (l['component'], ''), 'name': (look('name', i), ''), 'qty': (l['qty'], 'in'),
            'unit': (look('unit', i), ''), 'cur': (look('cur', i), ''), 'u_old': (look('old_iqd', i), ''),
            'u_new': (look('new_iqd', i), ''), 'u_infl': (look('infl_iqd', i), ''), 'u_price': (look('price_eff', i), ''),
            'old': (f'={x["qty"]}*{x["u_old"]}', ''), 'e_price': (f'={x["qty"]}*{x["u_price"]}', ''),
            'e_infl': (f'={x["qty"]}*{x["u_infl"]}', ''), 'new': (f'={x["qty"]}*{x["u_new"]}', ''),
            'e_fx': (f'={x["new"]}-{x["old"]}-{x["e_price"]}-{x["e_infl"]}', ''),
            'usd': (f'=IF({x["cur"]}="{USD}",{x["qty"]}*({x["u_new"]}-{x["u_infl"]}),0)', ''),
            'kg': (f'=IF({x["unit"]}="كغ",{x["qty"]},0)', ''),
        })
    write_table(bm, bcols, rows, center={'prod', 'comp', 'unit', 'cur'})
    n_bom = len(lines) + 1
    set_widths(bm, [13, 13, 26, 10, 7, 7, 11, 11, 10, 10, 12, 11, 11, 11, 12, 12, 8])
    bm.freeze_panes = 'C2'
    bsum = lambda k, i: f"SUMIFS('{S_BOM}'!${B[k]}$2:${B[k]}${n_bom},'{S_BOM}'!${B['prod']}$2:${B['prod']}${n_bom},$A{i})"

    # ---------- التسعيرة الجديدة ----------
    pr = wb.create_sheet(S_PRICE)
    sc = f'(1-{A["scrap"]})'
    d, c, mg = A['disc'], A['credit'], A['margin']
    pcols = [
        ('code', 'الرمز', None), ('name', 'المنتج', None), ('cat', 'الفئة', None), ('unit', 'الوحدة', None),
        ('kg', 'كغ للوحدة', FMT_KG), ('convkg', 'كلفة التحويل للكيلو', FMT_IQD), ('vol', 'الكمية الشهرية', '#,##0;;"-"'),
        ('mat0', 'كلفة المواد السابقة', FMT_IQD), ('conv0', 'كلفة التحويل السابقة', FMT_IQD), ('cost0', 'الكلفة السابقة', FMT_IQD),
        ('e_price', 'أثر سعر المواد', FMT_IQD), ('e_fx', 'أثر الصرف والكمرك', FMT_IQD), ('e_infl', 'أثر التضخم', FMT_IQD),
        ('cost', 'الكلفة الجديدة', FMT_IQD), ('usd', 'منها بالدولار', FMT_IQD),
        ('price', 'سعر القائمة الحالي', FMT_IQD), ('net', 'الصافي الحالي بعد الخصم', FMT_IQD),
        ('m0', 'الهامش السابق', FMT_PCT_SIGNED), ('m_now', 'الهامش الحالي على الكلفة الجديدة', FMT_PCT_SIGNED),
        ('p_target', 'السعر للهامش المستهدف', FMT_IQD), ('p_keep', 'السعر للحفاظ على الهامش السابق', FMT_IQD),
        ('p_new', 'السعر المقترح', FMT_IQD), ('p_manual', 'السعر اليدوي', FMT_IQD), ('p_final', 'السعر النهائي', FMT_IQD),
        ('chg', 'نسبة التغيير', FMT_PCT_SIGNED), ('m_new', 'الهامش بالسعر النهائي', FMT_PCT_SIGNED),
        ('p_tax', 'السعر شامل ضريبة المبيعات', FMT_IQD), ('status', 'الحالة', None), ('rev', 'أثر الإيراد الشهري', FMT_IQD),
        ('imp_usd', 'سعر المستورد بالدولار', FMT_DEC), ('imp_iqd', 'كلفة المستورد الواصلة', FMT_IQD),
        ('vs_imp', 'صافي سعرنا مقابل المستورد', FMT_PCT_SIGNED),
        ('cost_kg', 'كلفة الكيلو الجديدة', FMT_IQD), ('net_kg', 'صافي سعر الكيلو الحالي', FMT_IQD), ('w', 'وزن الترجيح', '#,##0'),
    ]
    L = letters(pcols)
    rows = []
    for i, p in enumerate(products, start=2):
        x = {k: f'{L[k]}{i}' for k in L}
        val = lambda v: (v, 'in')
        rows.append({
            'code': (p['code'], ''), 'name': (p['name'], ''), 'cat': (p.get('category') or '', ''), 'unit': (p.get('unit') or '', ''),
            'kg': val(p['kg_per_unit']) if p['kg_per_unit'] is not None else (f'={bsum("kg", i)}', ''),
            'convkg': val(p['conversion_per_kg']) if p['conversion_per_kg'] is not None else (f'={A["conv"]}', ''),
            'vol': val(p['volume']),
            'mat0': (f'={bsum("old", i)}/{sc}', ''), 'conv0': (f'={x["kg"]}*{x["convkg"]}', ''),
            'cost0': (f'={x["mat0"]}+{x["conv0"]}', ''),
            'e_price': (f'={bsum("e_price", i)}/{sc}', ''), 'e_fx': (f'={bsum("e_fx", i)}/{sc}', ''),
            'e_infl': (f'={bsum("e_infl", i)}/{sc}+{x["conv0"]}*{A["cinfl"]}', ''),
            'cost': (f'={x["cost0"]}+{x["e_price"]}+{x["e_fx"]}+{x["e_infl"]}', 'bold'),
            'usd': (f'={bsum("usd", i)}/{sc}', ''),
            'price': val(p['list_price']),
            'net': (f'=IF({x["price"]}="","",{x["price"]}*(1-{d}))', ''),
            'm0': (f'=IF({x["price"]}="","",({x["net"]}*(1-{c})-{x["cost0"]})/{x["net"]})', ''),
            'm_now': (f'=IF({x["price"]}="","",({x["net"]}*(1-{c})-{x["cost"]})/{x["net"]})', ''),
            'p_target': (f'={x["cost"]}/(1-{mg}-{c})/(1-{d})', ''),
            'p_keep': (f'=IF({x["price"]}="","",{x["price"]}*{x["cost"]}/{x["cost0"]})', ''),
            'p_new': (f'=MAX(ROUNDUP(IF(AND({A["basis"]}="{BASIS_KEEP}",{x["price"]}<>""),{x["p_keep"]},{x["p_target"]})/{A["round"]},0)'
                      f'*{A["round"]},IF(AND({A["allow_down"]}="{NO}",{x["price"]}<>""),{x["price"]},0))', ''),
            'p_manual': val(p['manual_price']),
            'p_final': (f'=IF({x["p_manual"]}<>"",{x["p_manual"]},{x["p_new"]})', 'bold'),
            'chg': (f'=IF({x["price"]}="","",{x["p_final"]}/{x["price"]}-1)', 'bold'),
            'm_new': (f'=({x["p_final"]}*(1-{d})*(1-{c})-{x["cost"]})/({x["p_final"]}*(1-{d}))', ''),
            'p_tax': (f'={x["p_final"]}*(1+{A["sales_tax"]})', ''),
            'status': (f'=IF({x["price"]}="","بدون سعر حالي",IF({x["net"]}*(1-{c})<{x["cost"]},"يبيع تحت الكلفة",'
                       f'IF({x["m_now"]}<{mg},"تحت الهامش المستهدف","ضمن الهامش")))', ''),
            'rev': (f'=IF(OR({x["price"]}="",{x["vol"]}=""),"",({x["p_final"]}-{x["price"]})*(1-{d})*{x["vol"]})', ''),
            'imp_usd': val(p['import_price_usd']),
            'imp_iqd': (f'=IF({x["imp_usd"]}="","",{x["imp_usd"]}*{A["imp_fx"]}*(1+{A["cust_fg"]}+{A["tax_fg"]}))', ''),
            'vs_imp': (f'=IF({x["imp_usd"]}="","",{x["p_final"]}*(1-{d})/{x["imp_iqd"]}-1)', ''),
            'cost_kg': (f'=IF({x["kg"]}>0,{x["cost"]}/{x["kg"]},"")', ''),
            'net_kg': (f'=IF(AND({x["price"]}<>"",{x["kg"]}>0),{x["net"]}/{x["kg"]},"")', ''),
            'w': (f'=IF({x["vol"]}="",1,{x["vol"]})', ''),
        })
    write_table(pr, pcols, rows, center={'code', 'cat', 'unit', 'status'})
    last = len(products) + 1
    tot = last + 1
    pr.cell(row=tot, column=2, value='الإجمالي المرجح بالكمية')
    sums = ('mat0', 'conv0', 'cost0', 'e_price', 'e_fx', 'e_infl', 'cost', 'usd')
    for k in sums:
        pr[f'{L[k]}{tot}'] = f'=SUMPRODUCT({L[k]}2:{L[k]}{last},${L["w"]}$2:${L["w"]}${last})'
    pr[f'{L["rev"]}{tot}'] = f'=SUM({L["rev"]}2:{L["rev"]}{last})'
    style_total(pr, tot, 1, len(pcols))
    for k in (*sums, 'rev'):
        pr[f'{L[k]}{tot}'].number_format = FMT_IQD
    st = L['status']
    for text, fill in (('يبيع تحت الكلفة', FILL_RED), ('تحت الهامش المستهدف', FILL_AMBER), ('ضمن الهامش', FILL_GREEN)):
        pr.conditional_formatting.add(f'{st}2:{st}{last}', FormulaRule(formula=[f'{st}2="{text}"'], fill=fill))
    set_widths(pr, [13, 30, 12, 8, 9, 10, 10, 12, 12, 12, 11, 11, 11, 12, 12, 12, 12, 10, 11, 12, 12, 12, 12, 12,
                    9, 10, 12, 17, 13, 11, 12, 11, 11, 11, 8])
    pr.freeze_panes = 'C2'
    pr.auto_filter.ref = f'A1:{L["w"]}{last}'
    P = lambda k: f"'{S_PRICE}'!${L[k]}$2:${L[k]}${last}"

    # ---------- سيناريوهات الصرف ----------
    sn = wb.create_sheet(S_SCEN)
    sn.merge_cells('A1:E1')
    sn['A1'], sn['A1'].font, sn['A1'].alignment = 'أثر سعر الصرف على الكلفة والسعر المطلوب', F_SECTION, CENTER
    heads = ['سعر الصرف الفعلي', 'الكلفة الكلية المرجحة', 'تغير الكلفة عن السابقة', 'التغير المطلوب للهامش المستهدف (مرجح)',
             'منتجات تحت الكلفة بالسعر الحالي']
    for j, h in enumerate(heads, start=1):
        sn.cell(row=3, column=j, value=h)
    style_header(sn, 3, 1, len(heads))
    rates = sorted({args.fx_old, args.fx_official, round(args.fx_market), 1450, 1500, 1550, 1600, 1700, 1750})
    cost_at = lambda r: f'(({P("cost")}-{P("usd")})+{P("usd")}*A{r}/{A["fx_new"]})'
    for r, fx in enumerate(rates, start=4):
        vals = [(fx, '#,##0', F_INPUT),
                (f'=SUMPRODUCT({P("w")},{cost_at(r)})', FMT_IQD, F_BODY),
                (f'=B{r}/SUMPRODUCT({P("w")},{P("cost0")})-1', FMT_PCT_SIGNED, F_BODY),
                (f'=SUMPRODUCT(({P("price")}>0)*{P("w")}*{cost_at(r)})/(1-{mg}-{c})/(1-{d})/SUMPRODUCT(({P("price")}>0)*{P("w")}*{P("price")})-1',
                 FMT_PCT_SIGNED, F_BOLD),
                (f'=SUMPRODUCT(({P("price")}>0)*({P("price")}*(1-{d})*(1-{c})<{cost_at(r)}))', '0', F_BODY)]
        for j, (v, fmt, font) in enumerate(vals, start=1):
            style_body(sn.cell(row=r, column=j, value=v), fmt, font=font, align=CENTER)
    sn.cell(row=len(rates) + 5, column=1,
            value='كل سيناريو يحرك كل كلف الدولار بنسبة السعر للسعر الفعلي المرجح. التغير قبل التقريب، ومرجح بالكمية إن وجدت وإلا بالتساوي. التضخم ثابت.').font = F_NOTE
    set_widths(sn, [16, 20, 18, 26, 24])

    # ---------- ملاحظات المراجعة ----------
    n_review = 0
    if args.review:
        rv = load_workbook(args.review, data_only=True)['الملاحظات']
        src_rows = [r for r in rv.iter_rows(min_row=5, values_only=True) if r and r[0] is not None]
        rs = wb.create_sheet(S_REVIEW)
        heads = ['#', 'الدرجة', 'الورقة', 'الخلية', 'النوع', 'التفصيل', 'الاقتراح', 'القرار', 'ملاحظة المستخدم']
        for j, h in enumerate(heads, start=1):
            rs.cell(row=1, column=j, value=h)
        style_header(rs, 1, 1, len(heads))
        for i, r in enumerate(src_rows, start=2):
            for j, v in enumerate(list(r)[:9], start=1):
                style_body(rs.cell(row=i, column=j, value=v), align=CENTER if j in (1, 2, 4, 8) else WRAP)
        n_review = len(src_rows)
        set_widths(rs, [5, 9, 16, 9, 20, 55, 45, 11, 30])
        rs.freeze_panes = 'A2'

    # ---------- الملخص ----------
    sm = wb.create_sheet(S_SUM, 1)
    sm.merge_cells('A1:D1')
    sm['A1'], sm['A1'].font, sm['A1'].alignment = COMPANY, F_TITLE, CENTER
    sm.merge_cells('A2:D2')
    sm['A2'] = (f'=CONCATENATE("التسعيرة الجديدة: الصرف الرسمي ",TEXT({A["fx_off"]},"#,##0"),"، والسوق ",TEXT({A["fx_mkt"]},"#,##0"),'
                f'"، والفعلي المرجح ",TEXT({A["fx_new"]},"#,##0"),"، وتضخم سنوي ",TEXT({A["infl"]},"0.0%"))')
    sm['A2'].font, sm['A2'].alignment = F_SUB, CENTER
    sm.merge_cells('A3:D3')
    sm['A3'] = 'الأسعار المقترحة توصية، واعتمادها يتطلب موافقة المدير العام والشركاء.'
    sm['A3'].font, sm['A3'].alignment = F_NOTE, CENTER
    for j, h in enumerate(['المؤشر', 'القيمة', 'الملاحظة'], start=1):
        sm.cell(row=5, column=j, value=h)
    style_header(sm, 5, 1, 3)
    wt = f'SUMPRODUCT(({P("price")}>0)*{P("w")}'
    base = f'SUMPRODUCT({P("w")},{P("cost0")})'
    kpis = [
        ('عدد المنتجات', f'=COUNTA({P("code")})', '0', ''),
        ('منها بسعر حالي', f'=COUNT({P("price")})', '0', ''),
        ('تبيع تحت الكلفة الجديدة', f'=COUNTIF({P("status")},"يبيع تحت الكلفة")', '0', 'صافي السعر بعد الخصم وكلفة الآجل أقل من الكلفة'),
        ('تحت الهامش المستهدف', f'=COUNTIF({P("status")},"تحت الهامش المستهدف")', '0', ''),
        ('بسعر يدوي', f'=COUNT({P("p_manual")})', '0', 'السعر اليدوي يتقدم على المقترح'),
        ('تغير الكلفة (مرجح)', f'=SUMPRODUCT({P("w")},{P("cost")})/{base}-1', FMT_PCT_SIGNED, 'الكلفة الجديدة مقابل السابقة'),
        ('منه: سعر المواد', f'=SUMPRODUCT({P("w")},{P("e_price")})/{base}', FMT_PCT_SIGNED, 'تغير أسعار المواد بعملتها'),
        ('منه: الصرف والكمرك', f'=SUMPRODUCT({P("w")},{P("e_fx")})/{base}', FMT_PCT_SIGNED, 'المواد المستوردة والمسعرة بالدولار'),
        ('منه: التضخم', f'=SUMPRODUCT({P("w")},{P("e_infl")})/{base}', FMT_PCT_SIGNED, 'المواد المحلية وكلفة التحويل'),
        ('الهامش الحالي على الكلفة الجديدة (مرجح)',
         f'={wt}*({P("price")}*(1-{d})*(1-{c})-{P("cost")}))/{wt}*{P("price")}*(1-{d}))', FMT_PCT_SIGNED, ''),
        ('التغير بالسعر النهائي (مرجح)', f'={wt}*{P("p_final")})/{wt}*{P("price")})-1', FMT_PCT_SIGNED, 'بعد التقريب والأسعار اليدوية'),
        ('الهامش بالسعر النهائي (مرجح)',
         f'={wt}*({P("p_final")}*(1-{d})*(1-{c})-{P("cost")}))/{wt}*{P("p_final")}*(1-{d}))', FMT_PCT_SIGNED, ''),
        ('الهامش بعد ضريبة الدخل (تقريبي)', None, FMT_PCT_SIGNED, 'قبل المصاريف الثابتة غير المحملة على المنتج'),
        ('أثر الإيراد الشهري', f'=SUM({P("rev")})', FMT_IQD, 'صافي بعد الخصم، للمنتجات اللي إلها كمية شهرية فقط'),
        ('أغلى من المستورد الواصل', f'=COUNTIF({P("vs_imp")},">0")', '0', 'للمنتجات اللي إلها سعر مستورد'),
    ]
    row_of = {label: i for i, (label, *_) in enumerate(kpis, start=6)}
    for i, (label, f, fmt, note) in enumerate(kpis, start=6):
        if f is None:
            f = f'=B{row_of["الهامش بالسعر النهائي (مرجح)"]}*(1-{A["income_tax"]})'
        style_body(sm.cell(row=i, column=1, value=label), font=F_BOLD, align=RIGHT)
        style_body(sm.cell(row=i, column=2, value=f), fmt, font=F_BOLD, align=CENTER)
        style_body(sm.cell(row=i, column=3, value=note), font=F_NOTE, align=WRAP)
    r = len(kpis) + 7
    cats = sorted({p.get('category') for p in products if p.get('category')})
    if cats:
        for j, h in enumerate(['الفئة', 'عدد المنتجات', 'متوسط التغير بالسعر النهائي', 'تحت الكلفة'], start=1):
            sm.cell(row=r, column=j, value=h)
        style_header(sm, r, 1, 4)
        for k, cat in enumerate(cats, start=r + 1):
            style_body(sm.cell(row=k, column=1, value=cat), font=F_BOLD, align=RIGHT)
            style_body(sm.cell(row=k, column=2, value=f'=COUNTIF({P("cat")},A{k})'), '0', align=CENTER)
            style_body(sm.cell(row=k, column=3, value=f'=IFERROR(AVERAGEIF({P("cat")},A{k},{P("chg")}),"")'), FMT_PCT_SIGNED, align=CENTER)
            style_body(sm.cell(row=k, column=4, value=f'=COUNTIFS({P("cat")},A{k},{P("status")},"يبيع تحت الكلفة")'), '0', align=CENTER)
        r += len(cats) + 2
    sm.cell(row=r, column=1, value='ملاحظات').font = F_SECTION
    notes = [
        'التضخم يُطبق على المواد المحلية المؤشرة وكلفة التحويل فقط، والمواد بالدولار يغيرها سعر الصرف وسعرها بالدولار، حتى لا يُحسب الأثر مرتين.',
        'الكمرك يُحتسب على قيمة المادة بسعر صرف الكمرك، والتخليص نسبة من قيمتها. النسب مدخلات يدوية تلزم مطابقتها مع بيانات الكمرك.',
        'المخزون الموجود يبقى بكلفته الدفترية (IAS 2)، والتسعير على كلفة الاستبدال.',
        'كلفة التحويل موزعة بالكيلو، وهي تقدير لحين توفر كلفة كل خط (طاقة وساعات).',
        'تغيير قائمة الأسعار والخصم يتطلب موافقة المدير العام والشركاء، ويُبلغ الوكلاء بالكتابة قبل النفاذ.',
    ]
    if n_review:
        notes.insert(0, f'التسعيرة مبنية بعد مراجعة ملف التسعيرة الحالي: {n_review} ملاحظة بورقة {S_REVIEW}.')
    for k, t in enumerate(notes, start=r + 1):
        style_body(sm.cell(row=k, column=1, value=t), align=WRAP)
        sm.merge_cells(start_row=k, start_column=1, end_row=k, end_column=4)
        sm.row_dimensions[k].height = 30
    set_widths(sm, [40, 18, 46, 14])

    # ---------- المنهجية والاعتماد ----------
    me = wb.create_sheet(S_METHOD)
    me['A1'], me['A1'].font = 'المنهجية والاعتماد', F_TITLE
    method = [
        'الواصل للكيلو = السعر × سعر الصرف المطبق + الكمرك والرسوم (السعر × سعر صرف الكمرك × النسبة) + التخليص (نسبة من القيمة) + زيادة التضخم للمواد المحلية.',
        'سعر الصرف المطبق حسب مصدر دولار المادة: رسمي، أو سوق، أو مرجح بنسبة الشراء من السوق، أو يدوي.',
        'كلفة المواد للوحدة = مجموع (الكمية × الواصل للكيلو) ÷ (1 − الهدر غير المسترد).',
        'كلفة التحويل للوحدة = كغ الوحدة × كلفة التحويل للكيلو، وتزيد بتضخم كلفة التحويل للفترة.',
        'جسر الكلفة: الكلفة السابقة + أثر سعر المواد + أثر الصرف والكمرك + أثر التضخم = الكلفة الجديدة.',
        'صافي السعر = سعر القائمة × (1 − الخصم). كلفة الآجل = صافي السعر × كلفة الآجل من الافتراضات.',
        'الهامش = (صافي السعر − كلفة الآجل − الكلفة) ÷ صافي السعر، قبل ضريبة المبيعات وضريبة الدخل.',
        'السعر المطلوب = الكلفة ÷ (1 − الهامش المستهدف − كلفة الآجل) ÷ (1 − الخصم)، ثم يُقرب للأعلى. السعر اليدوي يتقدم عليه.',
        'كلفة المستورد الواصلة = سعره بالدولار × سعر صرف المستورد × (1 + الكمرك والرسوم على المنتج المستورد).',
        f'المدخلات: {args.products.name} و{args.bom.name} و{args.materials.name}.',
    ]
    for i, t in enumerate(method, start=3):
        style_body(me.cell(row=i, column=1, value=t), align=WRAP)
        me.merge_cells(start_row=i, start_column=1, end_row=i, end_column=4)
        me.row_dimensions[i].height = 30
    r = len(method) + 4
    for j, h in enumerate(['الدور', 'الاسم', 'التوقيع', 'التاريخ'], start=1):
        me.cell(row=r, column=j, value=h)
    style_header(me, r, 1, 4)
    for k, role in enumerate(['أعدّه: الإدارة المالية', 'راجعه: مدير الإنتاج', 'اعتمده: المدير العام', 'اعتمده: الشركاء'], start=r + 1):
        for j, v in enumerate([role, None, None, None], start=1):
            style_body(me.cell(row=k, column=j, value=v), font=F_BOLD if j == 1 else F_INPUT, align=RIGHT)
        me.row_dimensions[k].height = 24
    set_widths(me, [34, 26, 22, 16])

    for w in wb.worksheets:
        page_setup(w, COMPANY, '1:1' if w.title in (S_PRICE, S_BOM, S_MAT, S_REVIEW) else None)
        if w.title in (S_ASSUME, S_SUM):
            w.page_setup.fitToHeight = 1
    wb.properties.creator = wb.properties.lastModifiedBy = 'الإدارة المالية'
    wb.properties.title = 'التسعيرة الجديدة'
    wb.save(args.out)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--products', type=Path, required=True)
    ap.add_argument('--bom', type=Path, required=True)
    ap.add_argument('--materials', type=Path, required=True)
    ap.add_argument('--out', type=Path, required=True)
    ap.add_argument('--margin', type=float, required=True, help='الهامش المستهدف على صافي السعر، مثلاً 0.15')
    ap.add_argument('--margin-basis', default='يحدده المدير العام والشركاء')
    ap.add_argument('--conversion-per-kg', type=float, required=True, help='كلفة التحويل للكيلو بالدينار')
    ap.add_argument('--conversion-basis', required=True, help='أساس الرقم ومصدره')
    ap.add_argument('--fx-old', type=float, required=True, help='سعر الصرف المستخدم بالتسعيرة الحالية')
    ap.add_argument('--fx-old-basis', default='سعر الصرف المستخدم بملف التسعيرة الحالي')
    ap.add_argument('--fx-official', type=float, default=1520, help='السعر الرسمي (قرار 544)')
    ap.add_argument('--fx-market', type=float, required=True, help='سعر السوق الموازي بتاريخ التسعيرة')
    ap.add_argument('--fx-market-basis', required=True, help='مصدر سعر السوق وتاريخه')
    ap.add_argument('--market-share', type=float, required=True, help='نسبة الدولار المشترى من السوق للمواد المصنفة مرجح (0 إلى 1)')
    ap.add_argument('--market-share-basis', default='نسبة مشتريات الدولار اللي ما تمر بالسعر الرسمي. يلزم تأكيدها من المشتريات والحسابات')
    ap.add_argument('--customs-raw', type=float, help='الرسم الكمركي على المواد الأولية (يترك فارغاً إذا غير معروف)')
    ap.add_argument('--import-other', type=float, help='رسوم استيراد أخرى على المواد')
    ap.add_argument('--clearance', type=float, help='التخليص والنقل الداخلي كنسبة من قيمة المادة')
    ap.add_argument('--customs-fg', type=float, help='الرسم الكمركي على المنتج البلاستيكي المستورد المنافس')
    ap.add_argument('--import-other-fg', type=float, help='رسوم أخرى على المنتج المستورد')
    ap.add_argument('--sales-tax', type=float, default=0.0)
    ap.add_argument('--income-tax', type=float, default=0.15)
    ap.add_argument('--discount', type=float, default=0.14)
    ap.add_argument('--inflation', type=float, default=0.033)
    ap.add_argument('--inflation-basis', default='الجهاز المركزي للإحصاء: التضخم السنوي آب 2026 بنسبة 3.3% (منقول بالصحافة). يلزم التحقق من آخر نشرة')
    ap.add_argument('--months', type=int, default=6)
    ap.add_argument('--scrap', type=float, default=0.02)
    ap.add_argument('--scrap-basis')
    ap.add_argument('--dso', type=float, default=91)
    ap.add_argument('--round-to', type=float, default=250)
    ap.add_argument('--review', type=Path, help='ملف المراجعة من review_workbook.py بعد تعبئة القرارات')
    ap.add_argument('--check', action='store_true')
    ap.add_argument('--pdf', action='store_true')
    args = ap.parse_args()

    products, lines, mats = load_inputs(args)
    print(f'المدخلات: {len(products)} منتج، {len(lines)} سطر وصفة، {len(mats)} مادة')
    build(args, products, lines, mats)
    print(f'الملف: {args.out}')
    if args.check:
        style = house_style_issues(args.out)
        print('أسلوب الملفات:', 'سليم' if not style else style[:10])
        res = recalc_check(args.out)
        if res:
            errors, s = res
            print('فحص المعادلات:', 'لا توجد أخطاء' if not errors else errors[:10])
            for i in range(6, 21):
                label, v = s.cell(row=i, column=1).value, s.cell(row=i, column=2).value
                if label:
                    print(f'   {label}: {v:.4g}' if isinstance(v, float) else f'   {label}: {v}')
            if errors:
                return 1
    if args.pdf:
        print('PDF:', export_pdf(args.out))
    return 0


if __name__ == '__main__':
    sys.exit(main())
