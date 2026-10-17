#!/usr/bin/env python3
"""مطابقة الخازن ويه Odoo بتاريخ قطع واحد، وإخراج ملف Excel بأسلوب الإدارة المالية.

الاستخدام:
  python recon_khazen_odoo.py --khazen <مجلد الخازن_تصدير_...> --odoo <مجلد تصدير Odoo>
      --cutoff 2026-09-30 --out "مطابقة_الخازن_Odoo_2026-09-30.xlsx" [--map ../assets/khazen_odoo_map.csv] [--check] [--pdf]

من الخازن: 11 (الحركات)، و14 (قائمة الوكلاء)، و16 (رؤوس فواتير البيع). الأرصدة تُحسب من الحركات لتاريخ القطع.
من Odoo: قيود اليومية، إما ملفات Cowork (12_قيود_اليومية_السنة_الحالية_Q*.xlsx بعناوين عربية)
أو تصدير odoo_fetch.py (11_قيود_اليومية_السنة_الحالية.csv بأسماء الحقول التقنية).
الفرق لكل وكيل أو شريك يُفصل إلى جزء سعر صرف (قيمة أسطر الدولار بالنظامين) وفرق متبقٍ يحتاج تحقق.
"""
import argparse
import csv
import glob
import os
import re
import sys
from collections import Counter, defaultdict
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_cost_workbook import (CENTER, F_BODY, F_BOLD, F_NOTE, F_SECTION, F_SUB, F_TITLE, FILL_AMBER,  # noqa: E402
                                 FILL_GREEN, FILL_RED, FMT_PCT, RIGHT, export_pdf, house_style_issues, page_setup,
                                 recalc_check, set_widths, style_body, style_header, style_total)
from openpyxl.formatting.rule import FormulaRule  # noqa: E402

SKILL_DIR = Path(__file__).resolve().parent.parent
IQD_FMT = '#,##0;[Red](#,##0);"-"'
USD_FMT = '#,##0.00;[Red](#,##0.00);"-"'
KHAZEN_IQD = 'دينار عراقي'
OPENING_JOURNAL = 'القيود الافتتاحية'
COMPANY = 'شركة البركة للاستثمارات الصناعية'

# عناوين ملفات Cowork العربية ← أسماء موحدة
COWORK_HEADERS = {'التاريخ': 'date', 'قيد اليومية': 'move', 'دفتر اليومية': 'journal', 'الحساب': 'account',
                  'الشريك': 'partner', 'المدين': 'debit', 'الدائن': 'credit', 'الرصيد': 'balance',
                  'العملة': 'currency', 'المبلغ بالعملة': 'amount_currency'}
FETCH_HEADERS = {'date': 'date', 'move_name': 'move', 'journal_id': 'journal', 'account_id': 'account',
                 'partner_id': 'partner', 'debit': 'debit', 'credit': 'credit', 'balance': 'balance',
                 'currency_id': 'currency', 'amount_currency': 'amount_currency'}


def num(v):
    if v in (None, '', 'None'):
        return Decimal(0)
    return Decimal(str(v))


def norm_name(s):
    s = re.sub('[أإآ]', 'ا', str(s or '')).replace('ة', 'ه').replace('ى', 'ي').replace('ـ', '')
    s = re.sub('[–—‐-]', '-', s)
    return re.sub(r'\s+', ' ', s).strip()


def read_csv(path):
    with open(path, encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))


def load_odoo(folder, cutoff):
    """يرجع أسطر Odoo لحد تاريخ القطع بحقول موحدة."""
    rows = []
    fetch = glob.glob(os.path.join(folder, '11_قيود_اليومية_السنة_الحالية*.csv'))
    cowork = sorted(glob.glob(os.path.join(folder, '12_قيود_اليومية_السنة_الحالية*.xlsx')))
    if fetch:
        for path in fetch:
            for r in read_csv(path):
                rows.append({v: r.get(k, '') for k, v in FETCH_HEADERS.items()})
    elif cowork:
        from openpyxl import load_workbook
        for path in cowork:
            ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
            it = ws.iter_rows(values_only=True)
            idx = {COWORK_HEADERS[h]: i for i, h in enumerate(next(it)) if h in COWORK_HEADERS}
            for row in it:
                rec = {k: row[i] for k, i in idx.items()}
                if hasattr(rec['date'], 'date'):
                    rec['date'] = rec['date'].date().isoformat()
                rows.append(rec)
    else:
        raise SystemExit(f'ما لكيت قيود يومية Odoo بالمجلد {folder}')
    for r in rows:
        r['date'] = str(r['date'])[:10]
        r['code'] = str(r['account'] or '').split(' ')[0]
        r['partner'] = r['partner'] or ''
        r['currency'] = r['currency'] or ''
    return [r for r in rows if r['date'] <= cutoff]


def load_khazen(folder, cutoff):
    lines = [r for r in read_csv(os.path.join(folder, '11_الحركات_السنة_الحالية.csv')) if r['EntryDate'] <= cutoff]
    agents = {r['AgentCode']: r['AgentName'] for r in read_csv(os.path.join(folder, '14_أرصدة_الوكلاء.csv'))}
    bills = [r for r in read_csv(os.path.join(folder, '16_فواتير_البيع_رأس.csv')) if r['BillDate'] <= cutoff]
    return lines, agents, bills


def khazen_balances(lines, codes=None, prefix=None):
    """لكل حساب: الرصيد بالدينار المعادل، وجزء أسطر الدولار بالدينار، والرصيد بالدولار."""
    out = defaultdict(lambda: {'lc': Decimal(0), 'usd_lc': Decimal(0), 'usd': Decimal(0), 'iqd': Decimal(0), 'name': ''})
    for r in lines:
        c = r['AccountCode']
        if (codes is not None and c not in codes) or (prefix and not c.startswith(prefix)):
            continue
        b = out[c]
        b['name'] = r['AccountName'].strip()
        lc = num(r['DebitLC']) - num(r['CreditLC'])
        b['lc'] += lc
        if r['Currency'] == KHAZEN_IQD:
            b['iqd'] += num(r['Debit']) - num(r['Credit'])
        else:
            b['usd_lc'] += lc
            b['usd'] += num(r['Debit']) - num(r['Credit'])
    return out


def odoo_group(rows, key, pred):
    out = defaultdict(lambda: {'bal': Decimal(0), 'usd_bal': Decimal(0), 'usd': Decimal(0), 'iqd_lines_usd': 0})
    for r in rows:
        if not pred(r):
            continue
        g = out[key(r)]
        g['bal'] += num(r['balance'])
        if r['currency'] == 'USD':
            g['usd_bal'] += num(r['balance'])
            g['usd'] += num(r['amount_currency'])
    return out


def mode_rate(pairs):
    c = Counter(round(b / a) for a, b in pairs if a)
    return c.most_common(1)[0][0] if c else None


# ---------- بناء الملف ----------

def detail_sheet(wb, title, headers, widths, rows, fmts, diff_cols, status_col):
    """جدول تفاصيل بمعادلات للفرق والحالة. diff_cols: {عمود: معادلة بصيغة {r}}."""
    ws = wb.create_sheet(title)
    ws.append(headers)
    style_header(ws, 1, 1, len(headers))
    for i, values in enumerate(rows, start=2):
        for j, v in enumerate(values, start=1):
            style_body(ws.cell(row=i, column=j, value=v), fmts[j - 1])
        for col, formula in diff_cols.items():
            style_body(ws.cell(row=i, column=col, value=formula.format(r=i)), fmts[col - 1],
                       align=CENTER if col == status_col else None)
    last = len(rows) + 1
    total = last + 1
    ws.cell(row=total, column=1, value='الإجمالي')
    for j, fmt in enumerate(fmts, start=1):
        if fmt in (IQD_FMT, USD_FMT) and last >= 2:
            L = ws.cell(row=1, column=j).column_letter
            ws.cell(row=total, column=j, value=f'=SUM({L}2:{L}{last})').number_format = fmt
    style_total(ws, total, 1, len(headers))
    if status_col and last >= 2:
        L = ws.cell(row=1, column=status_col).column_letter
        rng = f'{L}2:{L}{last}'
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{L}2="مطابق"'], fill=FILL_GREEN))
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{L}2="سعر صرف فقط"'], fill=FILL_AMBER))
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{L}2="يحتاج تحقق"'], fill=FILL_RED))
    set_widths(ws, widths)
    ws.freeze_panes = 'B2'
    return ws, total


STATUS = '=IF(ABS({d}{{r}})<{tol},"مطابق",IF(ABS({res}{{r}})<{tol},"سعر صرف فقط","يحتاج تحقق"))'


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--khazen', required=True)
    ap.add_argument('--odoo', required=True)
    ap.add_argument('--cutoff', required=True, help='YYYY-MM-DD')
    ap.add_argument('--out', required=True)
    ap.add_argument('--map', default=str(SKILL_DIR / 'assets' / 'khazen_odoo_map.csv'))
    ap.add_argument('--odoo-sales-accounts', default='411100001,411100002')
    ap.add_argument('--official-rate', type=float, default=1520)
    ap.add_argument('--tolerance', type=float, default=1000, help='حد الأهمية بالدينار للفروقات (افتراضياً 1,000 للتقريب)')
    ap.add_argument('--check', action='store_true')
    ap.add_argument('--pdf', action='store_true')
    a = ap.parse_args()

    tol = int(a.tolerance)
    kl, agents, bills = load_khazen(a.khazen, a.cutoff)
    ol = load_odoo(a.odoo, a.cutoff)
    k_rate = mode_rate([(num(r['Debit']) - num(r['Credit']), num(r['DebitLC']) - num(r['CreditLC']))
                        for r in kl if r['Currency'] != KHAZEN_IQD])
    o_rate = mode_rate([(num(r['amount_currency']), num(r['balance'])) for r in ol if r['currency'] == 'USD'])

    # --- الوكلاء ---
    kb = khazen_balances(kl, codes=set(agents))
    k_by_name = defaultdict(lambda: {'lc': Decimal(0), 'usd_lc': Decimal(0), 'usd': Decimal(0)})
    for code, b in kb.items():
        k = k_by_name[norm_name(agents[code])]
        for f in ('lc', 'usd_lc', 'usd'):
            k[f] += b[f]
    ob = odoo_group(ol, lambda r: norm_name(r['partner']), lambda r: r['code'].startswith(('1614', '1615')))
    agent_rows = []
    for name in sorted(set(k_by_name) | set(ob)):
        k, o = k_by_name.get(name), ob.get(name)
        if not (k and k['lc']) and not (o and o['bal']):
            continue
        k = k or {'lc': Decimal(0), 'usd_lc': Decimal(0), 'usd': Decimal(0)}
        o = o or {'bal': Decimal(0), 'usd_bal': Decimal(0), 'usd': Decimal(0)}
        note = '' if (name in k_by_name and name in ob) else ('بالخازن فقط' if name in k_by_name else 'بـ Odoo فقط')
        if not name:
            note = 'بـ Odoo بدون اسم شريك'
        agent_rows.append([name or '(بدون شريك)', float(k['lc']), float(o['bal']), None, float(o['usd_bal'] - k['usd_lc']),
                           None, None, float(k['usd']), float(o['usd']), note])

    # --- الشركاء ---
    kp = khazen_balances(kl, prefix='2112')
    op = odoo_group(ol, lambda r: r['code'], lambda r: r['code'].startswith('2112'))
    partner_rows = []
    for code in sorted(set(kp) | set(op)):
        k = kp.get(code, {'lc': Decimal(0), 'usd_lc': Decimal(0), 'usd': Decimal(0), 'name': ''})
        o = op.get(code, {'bal': Decimal(0), 'usd_bal': Decimal(0), 'usd': Decimal(0)})
        # Odoo يقيّم دولار الشريك بسعره بدون أسطر دولار منفصلة أحياناً، فجزء الصرف يُقدّر من رصيد الخازن بالدولار
        fx = float(k['usd']) * (o_rate - k_rate) if (o_rate and k_rate and k['usd']) else float(o['usd_bal'] - k['usd_lc'])
        partner_rows.append([f'{code} {k["name"]}'.strip(), float(k['lc']), float(o['bal']), None, fx, None, None,
                             float(k['usd']), float(o['usd']), ''])

    # --- النقد والمصارف ---
    cash_rows, usd_on_iqd = [], []
    mapping = read_csv(a.map)
    k_cash = khazen_balances(kl, codes={m['khazen_code'] for m in mapping if m['kind'] == 'cash'})
    for m in (m for m in mapping if m['kind'] == 'cash'):
        kc = k_cash.get(m['khazen_code'])
        is_usd = m['khazen_currency'] != KHAZEN_IQD
        k_native = (kc['usd'] if is_usd else kc['iqd']) if kc else Decimal(0)
        lines = [r for r in ol if r['code'] == m['odoo_code']]
        o_native = sum(num(r['amount_currency']) if is_usd else num(r['balance']) for r in lines)
        n_usd = sum(1 for r in lines if not is_usd and r['currency'] == 'USD')
        if n_usd and abs(o_native - k_native) >= tol:  # تُذكر فقط إذا الحساب بيه فرق ويه الخازن
            usd_on_iqd.append((m['label'], n_usd, sum(num(r['balance']) for r in lines if r['currency'] == 'USD'),
                               o_native - k_native))
        cash_rows.append([m['label'], m['khazen_code'], m['odoo_code'], 'دولار' if is_usd else 'دينار',
                          float(k_native), float(o_native), None, None,
                          f'{n_usd} سطر بعملة الدولار على حساب دينار' if n_usd else ''])

    # --- المبيعات الشهرية ---
    ks = defaultdict(Decimal)
    for r in bills:
        ks[r['BillDate'][:7]] += num(r['NetTotal']) * int(r['Sign']) * (num(r['ExchangeRate']) or 1)
    sales_codes = tuple(a.odoo_sales_accounts.split(','))
    os_, opb_rev = defaultdict(Decimal), Decimal(0)
    for r in ol:
        if r['code'] in sales_codes:
            if r['journal'] == OPENING_JOURNAL or str(r['move']).startswith('OPB/'):
                opb_rev -= num(r['balance'])
            else:
                os_[r['date'][:7]] -= num(r['balance'])
    sales_rows = [[m, float(ks.get(m, 0)), float(os_.get(m, 0)), None, None] for m in sorted(set(ks) | set(os_))]

    # --- البنود المفتوحة (محسوبة من البيانات) ---
    findings = []
    opb_dates = sorted({r['date'] for r in ol if r['journal'] == OPENING_JOURNAL and r['code'][:1] == '4'})
    if opb_rev and opb_dates:
        before = sum(v for m, v in os_.items() if m < opb_dates[0][:7])
        findings.append(('إيراد مكرر محتمل بـ Odoo',
                         f'قيد افتتاحي بتاريخ {opb_dates[0]} حمّل {opb_rev:,.0f} على حسابات المبيعات، بينما فواتير المبيعات قبل '
                         f'هذا التاريخ موجودة أصلاً بـ Odoo بمبلغ {before:,.0f}. أرباح Odoo للسنة غير صالحة قبل التصحيح.',
                         float(opb_rev), 'المحاسب ومزوّد Odoo: عكس أحد الطرفين بقيد تصحيح موثق', 'يتطلب موافقة المدير المالي'))
    diff_acc = defaultdict(Decimal)
    for r in ol:
        if 'فروقات الأرصدة الافتتاحية' in str(r['account']):
            diff_acc[str(r['account']).strip()] += num(r['balance'])
    for acc, v in diff_acc.items():
        if abs(v) >= 1:
            findings.append(('رصيد فروقات افتتاحية بـ Odoo', f'الحساب {acc} رصيده {v:,.0f}، أي أرصدة افتتاحية ما توازنت '
                             'وانحطت بحساب وسيط. يضخّم الأصول وحقوق الملكية لحد ما يتفسر.', float(v),
                             'المحاسب: تحليل الرصيد ونقله لحساباته الصحيحة', 'يتطلب موافقة المدير المالي'))
    findings.append(('سعر صرف مختلف بين النظامين', f'الخازن يقيّم الدولار غالباً بـ {k_rate:,}، وOdoo بـ {o_rate:,}، '
                     f'والسعر الرسمي من 2026-10-07 هو {a.official_rate:,.0f} (قرار 544).', None,
                     'الإدارة: اعتماد سعر إقفال موثق واحد للنظامين', 'يتطلب موافقة الإدارة'))
    for label, cnt, amt, gap in usd_on_iqd:
        findings.append(('فرق صندوق وبيه دفعات دولار من حساب دينار', f'{label}: الرصيد بـ Odoo يختلف عن الخازن بـ {gap:,.0f}، '
                         f'والحساب بيه {cnt} سطراً بعملة الدولار قيمتها {amt:,.0f} دينار. ممكن تكون دفعات دولار صحيحة من صندوق '
                         'الدينار، ويلزم التحقق إن نفس الدولار ما انحسب من صندوق الدولار أيضاً.', float(gap),
                         'المحاسب: جرد الصندوق ومطابقة حركاته بين النظامين', ''))
    residual = [(r[0], r[2] - r[1] - r[4]) for r in agent_rows]
    unexplained = sorted((x for x in residual if abs(x[1]) >= tol), key=lambda x: -abs(x[1]))
    if unexplained:
        top = '، '.join(f'{n_} ({v:,.0f})' for n_, v in unexplained[:5])
        findings.append(('فروقات وكلاء غير مفسرة بسعر الصرف', f'{len(unexplained)} وكيلاً، أكبرها: {top}.',
                         float(sum(v for _, v in unexplained)), 'المحاسب: كشف حساب الوكيل من النظامين لنفس الفترة', ''))
    sundry = [r for r in agent_rows if r[0] == '(بدون شريك)' or 'مبيعات' in r[0]]
    for r in sundry:
        findings.append(('حركة على حسابات الوكلاء بدون وكيل', f'"{r[0]}" برصيد {r[2]:,.0f} بـ Odoo على حسابات الوكلاء.',
                         r[2], 'المحاسب: إعادة توجيهها للوكيل أو الحساب الصحيح', ''))

    # --- الملف ---
    wb = Workbook()
    wb.remove(wb.active)
    ag_ws, ag_tot = detail_sheet(
        wb, 'الوكلاء',
        ['الوكيل', 'الخازن (دينار معادل)', 'Odoo (دينار)', 'الفرق', 'منه سعر صرف', 'الفرق المتبقي', 'الحالة',
         'الخازن (دولار)', 'Odoo (دولار)', 'ملاحظة'],
        [30, 17, 17, 15, 14, 15, 13, 12, 12, 20], agent_rows,
        [None, IQD_FMT, IQD_FMT, IQD_FMT, IQD_FMT, IQD_FMT, None, USD_FMT, USD_FMT, None],
        {4: '=C{r}-B{r}', 6: '=D{r}-E{r}', 7: STATUS.format(d='D', res='F', tol=tol)}, 7)
    pt_ws, pt_tot = detail_sheet(
        wb, 'الشركاء',
        ['الحساب', 'الخازن (دينار معادل)', 'Odoo (دينار)', 'الفرق', 'منه سعر صرف', 'الفرق المتبقي', 'الحالة',
         'الخازن (دولار)', 'Odoo (دولار)', 'ملاحظة'],
        [36, 17, 17, 15, 14, 15, 13, 12, 12, 10], partner_rows,
        [None, IQD_FMT, IQD_FMT, IQD_FMT, IQD_FMT, IQD_FMT, None, USD_FMT, USD_FMT, None],
        {4: '=C{r}-B{r}', 6: '=D{r}-E{r}', 7: STATUS.format(d='D', res='F', tol=tol)}, 7)
    cash_fmts = [None, None, None, None, USD_FMT, USD_FMT, USD_FMT, None, None]
    cs_ws, _ = detail_sheet(
        wb, 'النقد والمصارف',
        ['الحساب', 'رمز الخازن', 'رمز Odoo', 'العملة', 'الخازن (بعملته)', 'Odoo (بعملته)', 'الفرق', 'الحالة', 'ملاحظة'],
        [36, 12, 12, 8, 17, 17, 15, 12, 34], cash_rows, cash_fmts,
        {7: '=F{r}-E{r}', 8: '=IF(ABS(G{r})<IF(D{r}="دولار",1,' + str(tol) + '),"مطابق","يحتاج تحقق")'}, 8)
    total_row = len(cash_rows) + 2
    for col in (5, 6, 7):  # مجموع عملتين ما له معنى، فيُستبدل بمجموع كل عملة
        cs_ws.cell(row=total_row, column=col, value=None)
    cs_ws.cell(row=total_row, column=1, value='المجموع حسب العملة بالملخص')
    sl_ws, sl_tot = detail_sheet(
        wb, 'المبيعات الشهرية', ['الشهر', 'الخازن (صافي الفواتير)', 'Odoo (فواتير بدون القيد الافتتاحي)', 'الفرق', 'الحالة'],
        [12, 22, 26, 16, 12], sales_rows, [None, IQD_FMT, IQD_FMT, IQD_FMT, None],
        {4: '=C{r}-B{r}', 5: '=IF(ABS(D{r})<' + str(tol) + ',"مطابق","يحتاج تحقق")'}, 5)
    sl_ws.cell(row=sl_tot + 2, column=1, value=f'القيد الافتتاحي بـ Odoo على حسابات المبيعات (خارج الجدول): {opb_rev:,.0f}').font = F_NOTE

    fd = wb.create_sheet('البنود المفتوحة')
    fd.append(['#', 'البند', 'الوصف', 'المبلغ (دينار)', 'الإجراء والمسؤول', 'الموافقة'])
    style_header(fd, 1, 1, 6)
    for i, (title, desc, amt, action, approval) in enumerate(findings, start=1):
        for j, v in enumerate([i, title, desc, amt, action, approval or '-'], start=1):
            style_body(fd.cell(row=i + 1, column=j, value=v), IQD_FMT if j == 4 else None, align=RIGHT if j in (2, 3, 5) else CENTER)
        fd.row_dimensions[i + 1].height = 48
    set_widths(fd, [5, 28, 70, 18, 38, 22])

    sm = wb.create_sheet('الملخص', 0)
    set_widths(sm, [30, 19, 19, 17, 17, 17, 14])
    sm.merge_cells('A1:G1')
    sm['A1'], sm['A1'].font, sm['A1'].alignment = COMPANY, F_TITLE, CENTER
    sm.merge_cells('A2:G2')
    sm['A2'], sm['A2'].font, sm['A2'].alignment = f'مطابقة الخازن ويه Odoo بتاريخ {a.cutoff}', F_SUB, CENTER
    sm.merge_cells('A3:G3')
    sm['A3'] = (f'سعر الدولار بالخازن {k_rate:,} | بـ Odoo {o_rate:,} | الرسمي من 2026-10-07 {a.official_rate:,.0f}. '
                'المبالغ بالدينار إلا ما ذُكر.')
    sm['A3'].font, sm['A3'].alignment = F_NOTE, CENTER
    hdr = ['المجال', 'الخازن', 'Odoo', 'الفرق', 'منه سعر صرف', 'الفرق المتبقي', 'الحالة']
    for j, h in enumerate(hdr, start=1):
        sm.cell(row=5, column=j, value=h)
    style_header(sm, 5, 1, 7)
    cash_iqd = [i + 2 for i, r in enumerate(cash_rows) if r[3] == 'دينار']
    cash_usd = [i + 2 for i, r in enumerate(cash_rows) if r[3] == 'دولار']
    sumif = lambda col, rows_: '=' + ('+'.join(f"'النقد والمصارف'!{col}{r}" for r in rows_) if rows_ else '0')
    lines = [
        ('ذمم الوكلاء', f"='الوكلاء'!B{ag_tot}", f"='الوكلاء'!C{ag_tot}", f"='الوكلاء'!E{ag_tot}", IQD_FMT),
        ('جاري الشركاء', f"='الشركاء'!B{pt_tot}", f"='الشركاء'!C{pt_tot}", f"='الشركاء'!E{pt_tot}", IQD_FMT),
        ('النقد والمصارف بالدينار', sumif('E', cash_iqd), sumif('F', cash_iqd), '0', IQD_FMT),
        ('النقد بالدولار (دولار)', sumif('E', cash_usd), sumif('F', cash_usd), '0', USD_FMT),
        ('المبيعات من بداية السنة', f"='المبيعات الشهرية'!B{sl_tot}", f"='المبيعات الشهرية'!C{sl_tot}", '0', IQD_FMT),
    ]
    for i, (label, kf, of, fx, fmt) in enumerate(lines, start=6):
        vals = [label, kf, of, f'=C{i}-B{i}', '=' + fx.lstrip('='), f'=D{i}-E{i}',
                f'=IF(ABS(D{i})<{1 if fmt == USD_FMT else tol},"مطابق",IF(ABS(F{i})<{1 if fmt == USD_FMT else tol},"سعر صرف فقط","يحتاج تحقق"))']
        for j, v in enumerate(vals, start=1):
            style_body(sm.cell(row=i, column=j, value=v), fmt if 2 <= j <= 6 else None,
                       font=F_BOLD if j == 1 else F_BODY, align=CENTER if j == 7 else None)
    rng = 'G6:G10'
    sm.conditional_formatting.add(rng, FormulaRule(formula=['G6="مطابق"'], fill=FILL_GREEN))
    sm.conditional_formatting.add(rng, FormulaRule(formula=['G6="سعر صرف فقط"'], fill=FILL_AMBER))
    sm.conditional_formatting.add(rng, FormulaRule(formula=['G6="يحتاج تحقق"'], fill=FILL_RED))
    sm.cell(row=12, column=1, value='الوكلاء حسب الحالة').font = F_SECTION
    for j, h in enumerate(['الحالة', 'عدد الوكلاء', 'الفرق المتبقي'], start=1):
        sm.cell(row=13, column=j, value=h)
    style_header(sm, 13, 1, 3)
    for i, st in enumerate(['مطابق', 'سعر صرف فقط', 'يحتاج تحقق'], start=14):
        style_body(sm.cell(row=i, column=1, value=st), font=F_BOLD)
        style_body(sm.cell(row=i, column=2, value=f"=COUNTIF('الوكلاء'!G2:G{ag_tot - 1},A{i})"), '0')
        style_body(sm.cell(row=i, column=3, value=f"=SUMIF('الوكلاء'!G2:G{ag_tot - 1},A{i},'الوكلاء'!F2:F{ag_tot - 1})"), IQD_FMT)
    sm.cell(row=18, column=1, value=f"=\"البنود المفتوحة: \"&COUNTA('البنود المفتوحة'!B2:B200)&\" (التفاصيل بورقة البنود المفتوحة)\"").font = F_NOTE
    sm.cell(row=19, column=1, value='منه سعر صرف = قيمة أسطر الدولار بـ Odoo ناقص قيمتها بالخازن. الفرق المتبقي هو اللي يحتاج تحقق.').font = F_NOTE
    sm.cell(row=20, column=1, value='أرصدة الخازن محسوبة من حركاته لتاريخ القطع، وأرصدة Odoo من قيود اليومية المرحّلة لنفس التاريخ.').font = F_NOTE

    md = wb.create_sheet('المنهجية والاعتماد')
    set_widths(md, [24, 90])
    md['A1'], md['A1'].font = 'المنهجية والاعتماد', F_TITLE
    info = [('تاريخ القطع', a.cutoff), ('مصدر الخازن', os.path.basename(os.path.normpath(a.khazen))),
            ('مصدر Odoo', os.path.basename(os.path.normpath(a.odoo))),
            ('أسطر الخازن', f'{len(kl):,}'), ('أسطر Odoo', f'{len(ol):,}'),
            ('ربط الوكلاء', 'بالاسم بعد توحيد الهمزات والتاء المربوطة والياء والشرطات والمسافات'),
            ('ربط الحسابات', f'ملف الربط {os.path.basename(a.map)} للنقد والمصارف، ورمز الحساب نفسه للشركاء'),
            ('المبيعات', 'الخازن: صافي رؤوس فواتير البيع بالإشارة. Odoo: حسابات المبيعات بدون دفتر القيود الافتتاحية'),
            ('حد الأهمية', f'{tol:,} دينار (ودولار واحد لحسابات الدولار). الفرق الأقل منه يُعد تقريباً')]
    for i, (k, v) in enumerate(info, start=3):
        style_body(md.cell(row=i, column=1, value=k), font=F_BOLD, align=RIGHT)
        style_body(md.cell(row=i, column=2, value=v), align=RIGHT)
    r0 = len(info) + 5
    md.cell(row=r0, column=1, value='الاعتماد').font = F_SECTION
    for j, h in enumerate(['الدور', 'الاسم والتوقيع والتاريخ'], start=1):
        md.cell(row=r0 + 1, column=j, value=h)
    style_header(md, r0 + 1, 1, 2)
    for i, role in enumerate(['أعدّه', 'راجعه', 'اعتمده (المدير المالي)'], start=r0 + 2):
        style_body(md.cell(row=i, column=1, value=role), font=F_BOLD, align=RIGHT)
        style_body(md.cell(row=i, column=2, value=None))
        md.row_dimensions[i].height = 30

    for ws in wb.worksheets:
        page_setup(ws, COMPANY, '1:1' if ws.title not in ('الملخص', 'المنهجية والاعتماد') else None)
    wb.properties.creator = wb.properties.lastModifiedBy = 'الإدارة المالية'
    wb.properties.title = f'مطابقة الخازن وOdoo {a.cutoff}'
    out = Path(a.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f'تم إنشاء {out} | وكلاء {len(agent_rows)} | شركاء {len(partner_rows)} | حسابات نقد {len(cash_rows)} | بنود {len(findings)}')

    if a.check:
        style = house_style_issues(out)
        print('أسلوب الملفات: ' + ('سليم' if not style else f'{len(style)} مخالفة'))
        for s in style[:20]:
            print('  ', s)
        result = recalc_check(out)
        errors = []
        if result is None:
            print('تنبيه: LibreOffice غير متوفر، لم يُفحص الحساب.')
        else:
            errors, summary = result
            print('فحص المعادلات: ' + ('لا توجد أخطاء' if not errors else f'{len(errors)} خطأ'))
            for e in errors[:20]:
                print('  ', e)
            for i in range(6, 11):
                vals = [summary.cell(row=i, column=j).value for j in range(1, 8)]
                print('   ', ' | '.join(f'{v:,.0f}' if isinstance(v, (int, float)) else str(v) for v in vals))
        if errors or style:
            sys.exit(2)
    if a.pdf:
        pdf = export_pdf(out)
        print(f'PDF: {pdf}' if pdf else 'تنبيه: LibreOffice غير متوفر، لم يُصدَّر PDF.')


if __name__ == '__main__':
    main()
