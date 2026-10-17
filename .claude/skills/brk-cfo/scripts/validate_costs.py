#!/usr/bin/env python3
"""فحص بيانات التكاليف قبل بناء التقرير.

يكشف: أخطاء الترميز، العملة وسعر الصرف، المبالغ المفقودة، التكرار، تجزئة المشتريات
لتفادي حد الموافقة، الإنفاق الرأسمالي المصنّف كمصروف، والقيم الشاذة.

الاستخدام:
  python validate_costs.py costs.csv [--codes ../assets/cost_codes.csv] [--base IQD]
      [--fx USD=1520] [--approval-limit 5000000] [--capex-threshold 1000000]
      [--period-start 2026-07-01] [--period-end 2026-09-30]

يقرأ CSV أو XLSX بعناوين عربية أو إنكليزية (ومنها تصدير قيود Odoo: Date/Account/Partner/Label/Debit/Credit).
رمز الخروج 1 إذا وُجد "خطأ"، و0 إذا كانت تحذيرات وملاحظات فقط.
"""
import argparse
import csv
import datetime as dt
import re
import statistics
import sys
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path

SKILL_DIR = Path(__file__).resolve().parent.parent
DEFAULT_CODES = SKILL_DIR / 'assets' / 'cost_codes.csv'

ERROR, WARNING, INFO = 'خطأ', 'تحذير', 'ملاحظة'
SEVERITY_ORDER = {ERROR: 0, WARNING: 1, INFO: 2}

# أسماء الأعمدة المقبولة لكل حقل (المقارنة بدون حساسية لحالة الأحرف)
ALIASES = {
    'date': ['date', 'التاريخ', 'تاريخ'],
    'ref': ['ref', 'reference', 'journal entry', 'number', 'رقم المرجع', 'المرجع', 'رقم المستند', 'رقم الفاتورة'],
    'cost_center': ['cost_center', 'cost center', 'analytic account', 'analytic distribution', 'مركز الكلفة', 'مركز التكلفة'],
    'cost_code': ['cost_code', 'code', 'account', 'رمز الكلفة', 'الرمز', 'الحساب'],
    'vendor': ['vendor', 'supplier', 'partner', 'المورد', 'المجهز', 'الشريك'],
    'description': ['description', 'label', 'الوصف', 'البيان'],
    'qty': ['qty', 'quantity', 'الكمية'],
    'unit': ['unit', 'uom', 'الوحدة'],
    'unit_price': ['unit_price', 'price', 'unit price', 'سعر الوحدة', 'السعر'],
    'amount': ['amount', 'المبلغ'],
    'debit': ['debit', 'مدين'],
    'credit': ['credit', 'دائن'],
    'tax': ['tax', 'الضريبة', 'الضريبة/الرسوم', 'الرسوم'],
    'currency': ['currency', 'العملة'],
    'fx_rate': ['fx_rate', 'exchange rate', 'rate', 'سعر الصرف'],
    'payment_status': ['payment_status', 'status', 'payment status', 'حالة الدفع'],
    'notes': ['notes', 'ملاحظات'],
}

# كلمات تدل على شراء أصل وليس مصروفاً (IAS 16)، ما لم يكن الوصف خدمة على الأصل
CAPEX_WORDS = ('جهاز', 'مولدة', 'سيارة', 'معدات', 'أثاث', 'حاسبة', 'حاسوب', 'لابتوب', 'سيرفر',
               'مكيف', 'رافعة', 'generator', 'vehicle', 'laptop', 'server', 'equipment', 'forklift')
SERVICE_WORDS = ('صيانة', 'تصليح', 'اشتراك', 'إيجار', 'ايجار', 'تأجير', 'وقود', 'قطع غيار',
                 'maintenance', 'repair', 'subscription', 'rent', 'fuel')
# بنود قيمها مدوّرة بطبيعتها (رواتب، إيجارات، فوائد) فلا تُفحص للتدوير
ROUND_EXEMPT_CATEGORIES = {31, 36}
ROUND_EXEMPT_CODES = {3306}

ARABIC_DIGITS = str.maketrans('٠١٢٣٤٥٦٧٨٩٫', '0123456789.')


@dataclass
class Issue:
    severity: str
    row: int        # رقم الصف في سجل التكاليف (الصف 2 = أول قيد)
    field: str
    message: str


def _norm(header):
    return re.sub(r'\s+', ' ', str(header or '')).strip().lower()


def _read_table(path):
    path = Path(path)
    if path.suffix.lower() in ('.xlsx', '.xlsm'):
        from openpyxl import load_workbook
        ws = load_workbook(path, read_only=True, data_only=True).active
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    else:
        with open(path, encoding='utf-8-sig', newline='') as f:
            rows = list(csv.reader(f))
    rows = [r for r in rows if any(c not in (None, '') for c in r)]
    return (rows[0], rows[1:]) if rows else ([], [])


def _map_headers(headers):
    lookup = {_norm(alias): key for key, names in ALIASES.items() for alias in names}
    mapping = {}
    for i, h in enumerate(headers):
        key = lookup.get(_norm(h))
        if key and key not in mapping.values():
            mapping[i] = key
    return mapping


def parse_date(value):
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    s = str(value or '').strip().translate(ARABIC_DIGITS)
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%Y/%m/%d', '%d-%m-%Y', '%m/%d/%Y'):
        try:
            return dt.datetime.strptime(s[:10], fmt).date()
        except ValueError:
            pass
    return None


def parse_number(value):
    if value is None or value == '':
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().translate(ARABIC_DIGITS).replace(',', '').replace('٬', '')
    negative = s.startswith('(') and s.endswith(')')
    try:
        n = float(s.strip('()'))
    except ValueError:
        return None
    return -n if negative else n


def parse_code(value):
    # يقبل "3201" أو "3201 المواد الأولية" كما في تصدير Odoo
    m = re.match(r'\s*(\d+)', str(value or '').translate(ARABIC_DIGITS))
    return int(m.group(1)) if m else None


def parse_fx(pairs):
    """['USD=1520', 'EUR=1650'] -> {'USD': 1520.0, 'EUR': 1650.0}"""
    rates = {}
    for pair in pairs or []:
        cur, _, rate = pair.partition('=')
        if not rate or parse_number(rate) is None:
            raise SystemExit(f'صيغة سعر صرف غير صحيحة: {pair} (المطلوب مثل USD=1520)')
        rates[cur.strip().upper()] = parse_number(rate)
    return rates


def load_codes(path=DEFAULT_CODES):
    with open(path, encoding='utf-8-sig', newline='') as f:
        codes = {}
        for rec in csv.DictReader(f):
            code = parse_code(rec['code'])
            rec['level'] = int(rec.get('level') or (1 if code < 100 else 2))
            codes[code] = rec
    return codes


def load_rows(path):
    headers, data = _read_table(path)
    mapping = _map_headers(headers)
    missing = {'date', 'cost_code'} - set(mapping.values())
    if missing:
        raise SystemExit('أعمدة إلزامية مفقودة: %s | العناوين الموجودة: %s' % (', '.join(sorted(missing)), headers))
    rows = []
    for raw in data:
        rec = {key: (raw[i] if i < len(raw) else None) for i, key in mapping.items()}
        qty, price = parse_number(rec.get('qty')), parse_number(rec.get('unit_price'))
        amount = parse_number(rec.get('amount'))
        if amount is None and ('debit' in rec or 'credit' in rec):
            amount = (parse_number(rec.get('debit')) or 0.0) - (parse_number(rec.get('credit')) or 0.0)
        if qty is None and price is None and amount is not None:
            qty, price = 1.0, amount
        elif qty is None and price is not None:
            qty = 1.0
        rows.append({
            'date': parse_date(rec.get('date')),
            'ref': str(rec.get('ref') or '').strip(),
            'cost_center': str(rec.get('cost_center') or '').strip(),
            'cost_code': parse_code(rec.get('cost_code')),
            'raw_code': str(rec.get('cost_code') or '').strip(),
            'vendor': str(rec.get('vendor') or '').strip(),
            'description': str(rec.get('description') or '').strip(),
            'qty': qty,
            'unit': str(rec.get('unit') or '').strip(),
            'unit_price': price,
            'tax': parse_number(rec.get('tax')) or 0.0,
            'currency': str(rec.get('currency') or '').strip().upper(),
            'fx_rate': parse_number(rec.get('fx_rate')),
            'payment_status': str(rec.get('payment_status') or '').strip(),
            'notes': str(rec.get('notes') or '').strip(),
        })
    return rows


def prepare_rows(rows, base='IQD', fx=None):
    """يكمل العملة وسعر الصرف ويحسب المبالغ. يرجع ملاحظات عن الافتراضات."""
    fx = fx or {}
    notes = []
    for i, r in enumerate(rows, start=2):
        if not r['currency']:
            r['currency'] = base
            notes.append(Issue(INFO, i, 'العملة', f'العملة غير محددة، وافتُرضت {base}'))
        if r['fx_rate']:
            r['fx_source'] = 'القيد'
        elif r['currency'] == base:
            r['fx_rate'], r['fx_source'] = 1.0, 'العملة الأساسية'
        elif r['currency'] in fx:
            r['fx_rate'], r['fx_source'] = fx[r['currency']], 'افتراضي'
            notes.append(Issue(INFO, i, 'سعر الصرف',
                               f'استُخدم السعر الافتراضي {fx[r["currency"]]:,.2f} لعملة {r["currency"]}، '
                               'والمعيار IAS 21 يطلب سعر تاريخ المعاملة'))
        else:
            r['fx_source'] = None
        if r['qty'] is not None and r['unit_price'] is not None:
            r['total'] = r['qty'] * r['unit_price'] + r['tax']
            r['amount_base'] = r['total'] * r['fx_rate'] if r['fx_rate'] else None
        else:
            r['total'] = r['amount_base'] = None
    return notes


def validate(rows, codes, base='IQD', approval_limit=None, capex_threshold=1_000_000,
             period_start=None, period_end=None):
    """يفترض أن prepare_rows استُدعيت. يرجع قائمة Issue."""
    issues = []
    add = lambda sev, row, field, msg: issues.append(Issue(sev, row, field, msg))
    round_floor, round_unit = (10_000_000, 1_000_000) if base == 'IQD' else (10_000, 1_000)

    for i, r in enumerate(rows, start=2):
        if r['date'] is None:
            add(ERROR, i, 'التاريخ', 'تاريخ مفقود أو بصيغة غير مفهومة')
        else:
            if (period_start and r['date'] < period_start) or (period_end and r['date'] > period_end):
                add(WARNING, i, 'التاريخ', f'التاريخ {r["date"]} خارج فترة التقرير')
            if r['date'].weekday() == 4:
                add(INFO, i, 'التاريخ', 'قيد بيوم جمعة، يلزم التحقق من أنه ليس إدخالاً خارج الدوام بدون مبرر')

        code = r['cost_code']
        if code is None:
            add(ERROR, i, 'رمز الكلفة', f'رمز الكلفة مفقود أو غير رقمي: "{r["raw_code"]}"')
        elif code not in codes:
            add(ERROR, i, 'رمز الكلفة', f'الرمز {code} غير موجود في دليل الترميز')
        elif codes[code]['level'] == 1:
            add(INFO, i, 'رمز الكلفة', f'استُخدم رمز فئة رئيسية ({code}) بدل بند تفصيلي')

        if r['total'] is None:
            add(ERROR, i, 'المبلغ', 'الكمية أو سعر الوحدة مفقود')
        elif r['total'] <= 0:
            add(WARNING, i, 'المبلغ', 'مبلغ صفري أو سالب، وإذا كان إشعاراً دائناً يُوضَّح ذلك بالوصف')

        if r['fx_source'] is None:
            add(ERROR, i, 'سعر الصرف', f'لا يوجد سعر صرف لعملة {r["currency"]}، ويُضاف بالقيد أو بالخيار --fx {r["currency"]}=...')
        elif r['currency'] == 'USD' and base == 'IQD' and not 1000 <= r['fx_rate'] <= 2500:
            add(WARNING, i, 'سعر الصرف', f'سعر صرف الدولار {r["fx_rate"]:,.2f} غير منطقي')

        if not r['vendor']:
            add(WARNING, i, 'المورد', 'المورد غير محدد')
        if not r['ref']:
            add(WARNING, i, 'رقم المرجع', 'بدون رقم مستند، وكل قيد يحتاج مستنداً مؤيداً')

        base_amt = r['amount_base']
        if base_amt is not None:
            text = r['description'].lower()
            is_opex = code in codes and codes[code]['capex_opex'] == 'تشغيلي'
            looks_like_asset = any(w in text for w in CAPEX_WORDS) and not any(w in text for w in SERVICE_WORDS)
            if is_opex and capex_threshold and base_amt >= capex_threshold and looks_like_asset:
                add(WARNING, i, 'رمز الكلفة', 'يبدو شراء أصل مصنّفاً كمصروف تشغيلي، وتلزم مراجعة الرسملة حسب IAS 16')
            exempt = code in ROUND_EXEMPT_CODES or (code or 0) // 100 in ROUND_EXEMPT_CATEGORIES
            # يُفحص التدوير بالعملة الأساسية فقط، لأن تحويل عملة أجنبية قد ينتج رقماً مدوّراً بالصدفة
            if r['currency'] == base and not exempt and base_amt >= round_floor and base_amt % round_unit == 0:
                add(INFO, i, 'المبلغ', f'مبلغ مدوّر ({base_amt:,.0f})، يلزم التحقق من وجود فاتورة تفصيلية')

    # التكرار: نفس المورد ورقم المستند والمبلغ، أو نفس المورد والتاريخ والمبلغ بدون رقم مستند
    seen = {}
    for i, r in enumerate(rows, start=2):
        if r['total'] is None:
            continue
        key = (r['vendor'], r['ref'], round(r['total'], 2)) if r['ref'] else (r['vendor'], r['date'], round(r['total'], 2))
        if key in seen:
            add(WARNING, i, 'تكرار', f'تكرار محتمل مع الصف {seen[key]} (نفس المورد والمستند والمبلغ)')
        else:
            seen[key] = i

    # تجزئة المشتريات: عدة فواتير لنفس المورد بنفس اليوم، كل وحدة تحت حد الموافقة ومجموعها فوقه
    if approval_limit:
        groups = defaultdict(list)
        for i, r in enumerate(rows, start=2):
            if r['vendor'] and r['date'] and r['amount_base'] is not None and r['amount_base'] < approval_limit:
                groups[(r['vendor'], r['date'])].append((i, r['amount_base']))
        for (vendor, day), items in groups.items():
            total = sum(a for _, a in items)
            if len(items) >= 2 and total >= approval_limit:
                rows_txt = '، '.join(str(n) for n, _ in items)
                for n, _ in items:
                    add(WARNING, n, 'تجزئة', f'تجزئة محتملة لتفادي حد الموافقة: {len(items)} فواتير من "{vendor}" '
                                             f'بتاريخ {day} مجموعها {total:,.0f} (الصفوف {rows_txt})')

    # القيم الشاذة: مبلغ أكبر من 3 أضعاف الوسيط لنفس البند (عند وجود 5 قيود فأكثر)
    by_code = defaultdict(list)
    for i, r in enumerate(rows, start=2):
        if r['amount_base'] and r['amount_base'] > 0:
            by_code[r['cost_code']].append((i, r['amount_base']))
    for code, items in by_code.items():
        if len(items) >= 5:
            median = statistics.median(a for _, a in items)
            for n, a in items:
                if a > 3 * median:
                    add(WARNING, n, 'المبلغ', f'مبلغ شاذ: {a:,.0f} أكبر من 3 أضعاف الوسيط ({median:,.0f}) للبند {code}')

    issues.sort(key=lambda x: (SEVERITY_ORDER[x.severity], x.row))
    return issues


def common_args(parser):
    parser.add_argument('costs', help='ملف التكاليف CSV أو XLSX')
    parser.add_argument('--codes', default=str(DEFAULT_CODES), help='دليل الترميز CSV')
    parser.add_argument('--base', default='IQD', help='العملة الأساسية (افتراضياً IQD)')
    parser.add_argument('--fx', action='append', default=[], help='سعر صرف افتراضي مثل USD=1520 (يتكرر)')
    parser.add_argument('--approval-limit', type=float, help='حد الموافقة بالعملة الأساسية لكشف تجزئة المشتريات')
    parser.add_argument('--capex-threshold', type=float, default=1_000_000,
                        help='حد الرسملة بالعملة الأساسية (افتراضياً 1,000,000)')
    parser.add_argument('--period-start', type=parse_date, help='بداية الفترة YYYY-MM-DD')
    parser.add_argument('--period-end', type=parse_date, help='نهاية الفترة YYYY-MM-DD')


def run_checks(args):
    rows = load_rows(args.costs)
    codes = load_codes(args.codes)
    notes = prepare_rows(rows, args.base.upper(), parse_fx(args.fx))
    issues = validate(rows, codes, args.base.upper(), args.approval_limit, args.capex_threshold,
                      args.period_start, args.period_end)
    issues = sorted(issues + notes, key=lambda x: (SEVERITY_ORDER[x.severity], x.row))
    return rows, codes, issues


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    common_args(parser)
    args = parser.parse_args()
    rows, _, issues = run_checks(args)
    counts = {s: sum(1 for x in issues if x.severity == s) for s in (ERROR, WARNING, INFO)}
    print(f'القيود: {len(rows)} | أخطاء: {counts[ERROR]} | تحذيرات: {counts[WARNING]} | ملاحظات: {counts[INFO]}')
    for x in issues:
        print(f'[{x.severity}] الصف {x.row} | {x.field}: {x.message}')
    sys.exit(1 if counts[ERROR] else 0)


if __name__ == '__main__':
    main()
