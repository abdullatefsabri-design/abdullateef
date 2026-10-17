#!/usr/bin/env python3
"""يبني دفتر تكاليف احترافي (xlsx) من ملف قيود التكاليف.

الأوراق: الملخص، سجل التكاليف، الموازنة، الانحراف، ملاحظات الفحص، القوائم، المنهجية والاعتماد.
كل المجاميع معادلات حية مربوطة بسجل التكاليف، والمدخلات بخط أزرق والمعادلات بخط أسود.

الاستخدام:
  python build_cost_workbook.py costs.csv --company "اسم الشركة" --out report.xlsx
      [--budget budget.csv] [--codes ../assets/cost_codes.csv] [--base IQD] [--fx USD=1520]
      [--period-start 2026-07-01 --period-end 2026-09-30] [--approval-limit 5000000]
      [--capex-threshold 1000000] [--prepared-by "الاسم"] [--check] [--pdf]

ملف الموازنة (اختياري): cost_code, month (YYYY-MM), amount بالعملة الأساسية.
إذا غاب عمود الشهر يوزَّع المبلغ بالتساوي على أشهر الفترة.
--check يعيد حساب الملف بـ LibreOffice ويفحص أخطاء المعادلات، و--pdf يصدّر نسخة PDF.
"""
import argparse
import csv
import datetime as dt
import shutil
import subprocess
import sys
import tempfile
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

sys.path.insert(0, str(Path(__file__).resolve().parent))
from validate_costs import (ERROR, INFO, WARNING, common_args, parse_code, parse_date,  # noqa: E402
                            parse_number, run_checks)

# ---------- الهوية البصرية ----------
FONT_NAME = 'Arial'
NAVY, GREY, WHITE = '1F3864', '595959', 'FFFFFF'
FILL_HEADER = PatternFill('solid', fgColor=NAVY)
FILL_SECTION = PatternFill('solid', fgColor='D9E1F2')
FILL_TILE = PatternFill('solid', fgColor='EEF2F8')
FILL_TOTAL = PatternFill('solid', fgColor='F2F2F2')
FILL_RED = PatternFill('solid', fgColor='F8CBAD')
FILL_AMBER = PatternFill('solid', fgColor='FFE699')
FILL_GREEN = PatternFill('solid', fgColor='C6EFCE')
FILL_LILAC = PatternFill('solid', fgColor='E4DFEC')
THIN = Side(style='thin', color='BFBFBF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOTAL_BORDER = Border(top=Side(style='thin', color='000000'), bottom=Side(style='double', color='000000'))

F_TITLE = Font(name=FONT_NAME, size=16, bold=True, color=NAVY)
F_SUB = Font(name=FONT_NAME, size=12, bold=True, color=GREY)
F_SECTION = Font(name=FONT_NAME, size=12, bold=True, color=NAVY)
F_HEADER = Font(name=FONT_NAME, size=10, bold=True, color=WHITE)
F_BODY = Font(name=FONT_NAME, size=10, color='000000')
F_INPUT = Font(name=FONT_NAME, size=10, color='0000FF')     # مدخلات
F_BOLD = Font(name=FONT_NAME, size=10, bold=True)
F_TILE_LABEL = Font(name=FONT_NAME, size=9, color=GREY)
F_TILE_VALUE = Font(name=FONT_NAME, size=14, bold=True, color=NAVY)
F_NOTE = Font(name=FONT_NAME, size=9, italic=True, color=GREY)

CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
RIGHT = Alignment(horizontal='right', vertical='center', wrap_text=True)

FMT_PCT = '0.0%'
FMT_ORIG = '#,##0.00;[Red](#,##0.00);"–"'
FMT_RATE = '#,##0.00'
FMT_DATE = 'yyyy-mm-dd'
FMT_MONTH = 'yyyy-mm'

S_SUMMARY, S_REG, S_BUD, S_VAR = 'الملخص', 'سجل التكاليف', 'الموازنة', 'الانحراف'
S_ISSUES, S_LISTS, S_METHOD = 'ملاحظات الفحص', 'القوائم', 'المنهجية والاعتماد'
PAY_STATUSES = ['مدفوع', 'غير مدفوع', 'جزئي']


def q(sheet):
    return f"'{sheet}'"


REG, BUD, LST = q(S_REG), q(S_BUD), q(S_LISTS)

# أعمدة سجل التكاليف: (العنوان، العرض)
REG_COLS = [
    ('التاريخ', 11), ('رقم المرجع', 13), ('مركز الكلفة', 16), ('رمز الكلفة', 10), ('البند', 26),
    ('رمز الفئة', 8), ('الفئة الرئيسية', 22), ('طبيعة الكلفة', 11), ('السلوك', 9), ('تشغيلي/رأسمالي', 11),
    ('المورد', 22), ('الوصف', 30), ('الكمية', 8), ('الوحدة', 8), ('سعر الوحدة', 14), ('المجموع', 15),
    ('الضريبة/الرسوم', 12), ('الإجمالي', 15), ('العملة', 7), ('سعر الصرف', 10),
    ('الإجمالي بالعملة الأساسية', 17), ('الشهر', 9), ('حالة الدفع', 11), ('ملاحظات', 22),
    ('نتيجة الفحص عند الإعداد', 40),
]
C = {name: get_column_letter(i) for i, (name, _) in enumerate(REG_COLS, start=1)}


def base_fmt(base):
    return '#,##0;[Red](#,##0);"–"' if base == 'IQD' else '#,##0.00;[Red](#,##0.00);"–"'


def style_header(ws, row, first_col, last_col):
    for col in range(first_col, last_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.font, cell.fill, cell.alignment, cell.border = F_HEADER, FILL_HEADER, CENTER, BORDER


def style_body(cell, fmt=None, font=F_BODY, align=None):
    cell.font, cell.border = font, BORDER
    if fmt:
        cell.number_format = fmt
    if align:
        cell.alignment = align


def style_total(ws, row, first_col, last_col):
    for col in range(first_col, last_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.font, cell.fill, cell.border = F_BOLD, FILL_TOTAL, TOTAL_BORDER


def page_setup(ws, company, title_rows=None, landscape=True):
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = 'landscape' if landscape else 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.6
    ws.print_options.horizontalCentered = True
    if title_rows:
        ws.print_title_rows = title_rows
    ws.oddHeader.center.text = company
    ws.oddHeader.center.size = 9
    ws.oddFooter.right.text = 'سري — للاستخدام الداخلي'
    ws.oddFooter.center.text = 'صفحة &P من &N'
    ws.oddFooter.left.text = '&D'
    for part in (ws.oddFooter.right, ws.oddFooter.center, ws.oddFooter.left):
        part.size = 8


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def month_list(start, end):
    months, y, m = [], start.year, start.month
    while (y, m) <= (end.year, end.month):
        months.append(dt.date(y, m, 1))
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return months


def load_budget(path, months):
    """يرجع {code: {month_date: amount}}."""
    with open(path, encoding='utf-8-sig', newline='') as f:
        reader = csv.DictReader(f)
        fields = {(h or '').strip().lower(): h for h in reader.fieldnames}
        get = lambda rec, *names: next((rec[fields[n]] for n in names if n in fields), None)
        budget = defaultdict(lambda: defaultdict(float))
        for rec in reader:
            code = parse_code(get(rec, 'cost_code', 'code', 'رمز الكلفة', 'الرمز'))
            amount = parse_number(get(rec, 'amount', 'المبلغ', 'الموازنة'))
            if code is None or not amount:
                continue
            month_raw = get(rec, 'month', 'الشهر')
            if month_raw:
                d = parse_date(str(month_raw).strip()[:7] + '-01')
                if d is None:
                    raise SystemExit(f'شهر غير مفهوم في ملف الموازنة: {month_raw}')
                budget[code][d] += amount
            else:
                for m in months:
                    budget[code][m] += amount / len(months)
    return budget


# ---------- الأوراق ----------

def build_lists(wb, codes, centers, currencies):
    ws = wb.create_sheet(S_LISTS)
    headers = ['الرمز', 'البند', 'Name', 'المستوى', 'الأب', 'طبيعة الكلفة', 'السلوك', 'تشغيلي/رأسمالي']
    ws.append(headers)
    for code in sorted(codes):
        c = codes[code]
        ws.append([code, c['name_ar'], c['name_en'], c['level'], parse_code(c['parent']),
                   c['nature'] or None, c['behavior'] or None, c['capex_opex']])
    style_header(ws, 1, 1, len(headers))
    for row in ws.iter_rows(min_row=2, max_col=len(headers)):
        for cell in row:
            style_body(cell, font=F_BOLD if row[3].value == 1 else F_BODY)
    for col, title, values in (('J', 'مراكز الكلفة', centers), ('L', 'العملات', currencies),
                               ('N', 'حالات الدفع', PAY_STATUSES)):
        ws[f'{col}1'] = title
        style_header(ws, 1, ws[f'{col}1'].column, ws[f'{col}1'].column)
        for i, v in enumerate(values, start=2):
            style_body(ws.cell(row=i, column=ws[f'{col}1'].column, value=v))
    set_widths(ws, [8, 36, 28, 8, 7, 12, 10, 13, 2, 22, 2, 10, 2, 12])
    ws.freeze_panes = 'A2'
    return len(codes) + 1, len(centers) + 1, len(currencies) + 1


def build_register(wb, rows, row_issues, base, n_codes_last, n_centers_last, n_cur_last):
    ws = wb.create_sheet(S_REG)
    ws.append([name for name, _ in REG_COLS])
    style_header(ws, 1, 1, len(REG_COLS))
    ws.row_dimensions[1].height = 30
    bfmt = base_fmt(base)
    for r_idx, r in enumerate(rows, start=2):
        n = r_idx
        values = {
            'التاريخ': r['date'], 'رقم المرجع': r['ref'] or None, 'مركز الكلفة': r['cost_center'] or None,
            'رمز الكلفة': r['cost_code'] if r['cost_code'] is not None else (r['raw_code'] or None),
            'البند': f'=IF(${C["رمز الكلفة"]}{n}="","",IFERROR(VLOOKUP(${C["رمز الكلفة"]}{n},{LST}!$A:$H,2,FALSE),"⚠ رمز غير معرّف"))',
            'رمز الفئة': f'=IF(${C["رمز الكلفة"]}{n}="","",IFERROR(IF(${C["رمز الكلفة"]}{n}>=100,INT(${C["رمز الكلفة"]}{n}/100),${C["رمز الكلفة"]}{n}),""))',
            'الفئة الرئيسية': f'=IF(${C["رمز الفئة"]}{n}="","",IFERROR(VLOOKUP(${C["رمز الفئة"]}{n},{LST}!$A:$B,2,FALSE),""))',
            'طبيعة الكلفة': f'=IF(${C["رمز الكلفة"]}{n}="","",IFERROR(VLOOKUP(${C["رمز الكلفة"]}{n},{LST}!$A:$H,6,FALSE)&"",""))',
            'السلوك': f'=IF(${C["رمز الكلفة"]}{n}="","",IFERROR(VLOOKUP(${C["رمز الكلفة"]}{n},{LST}!$A:$H,7,FALSE)&"",""))',
            'تشغيلي/رأسمالي': f'=IF(${C["رمز الكلفة"]}{n}="","",IFERROR(VLOOKUP(${C["رمز الكلفة"]}{n},{LST}!$A:$H,8,FALSE),""))',
            'المورد': r['vendor'] or None, 'الوصف': r['description'] or None,
            'الكمية': r['qty'], 'الوحدة': r['unit'] or None, 'سعر الوحدة': r['unit_price'],
            'المجموع': f'=IF(OR(${C["الكمية"]}{n}="",${C["سعر الوحدة"]}{n}=""),"",${C["الكمية"]}{n}*${C["سعر الوحدة"]}{n})',
            'الضريبة/الرسوم': r['tax'] or None,
            'الإجمالي': f'=IF(${C["المجموع"]}{n}="","",${C["المجموع"]}{n}+N(${C["الضريبة/الرسوم"]}{n}))',
            'العملة': r['currency'], 'سعر الصرف': r['fx_rate'],
            'الإجمالي بالعملة الأساسية': f'=IF(OR(${C["الإجمالي"]}{n}="",${C["سعر الصرف"]}{n}=""),"",${C["الإجمالي"]}{n}*${C["سعر الصرف"]}{n})',
            'الشهر': f'=IF(${C["التاريخ"]}{n}="","",DATE(YEAR(${C["التاريخ"]}{n}),MONTH(${C["التاريخ"]}{n}),1))',
            'حالة الدفع': r['payment_status'] or None, 'ملاحظات': r['notes'] or None,
            'نتيجة الفحص عند الإعداد': '؛ '.join(row_issues.get(n, [])) or None,
        }
        formula_cols = {'البند', 'رمز الفئة', 'الفئة الرئيسية', 'طبيعة الكلفة', 'السلوك', 'تشغيلي/رأسمالي',
                        'المجموع', 'الإجمالي', 'الإجمالي بالعملة الأساسية', 'الشهر', 'نتيجة الفحص عند الإعداد'}
        fmts = {'التاريخ': FMT_DATE, 'سعر الوحدة': FMT_ORIG, 'المجموع': FMT_ORIG, 'الضريبة/الرسوم': FMT_ORIG,
                'الإجمالي': FMT_ORIG, 'سعر الصرف': FMT_RATE, 'الإجمالي بالعملة الأساسية': bfmt,
                'الشهر': FMT_MONTH, 'الكمية': 'General'}
        for c_idx, (name, _) in enumerate(REG_COLS, start=1):
            cell = ws.cell(row=n, column=c_idx, value=values[name])
            style_body(cell, fmts.get(name), F_BODY if name in formula_cols else F_INPUT)
    last = max(len(rows) + 1, 2)
    set_widths(ws, [w for _, w in REG_COLS])

    table = Table(displayName='tblCosts', ref=f'A1:{get_column_letter(len(REG_COLS))}{last}')
    table.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = 'B2'

    # قوائم منسدلة تمتد 500 صف إضافي لاستيعاب قيود جديدة
    upto = last + 500
    for col, formula in ((C['رمز الكلفة'], f'={LST}!$A$2:$A${n_codes_last}'),
                         (C['مركز الكلفة'], f'={LST}!$J$2:$J${n_centers_last}'),
                         (C['العملة'], f'={LST}!$L$2:$L${n_cur_last}'),
                         (C['حالة الدفع'], f'={LST}!$N$2:$N${len(PAY_STATUSES) + 1}')):
        dv = DataValidation(type='list', formula1=formula, allow_blank=True, showErrorMessage=True,
                            errorStyle='warning', errorTitle='قيمة خارج القائمة',
                            error='القيمة غير موجودة في ورقة القوائم. أضفها هناك أولاً أو اختر من القائمة.')
        dv.add(f'{col}2:{col}{upto}')
        ws.add_data_validation(dv)

    rng = lambda col: f'{col}2:{col}{upto}'
    ws.conditional_formatting.add(rng(C['البند']), FormulaRule(formula=[f'ISNUMBER(SEARCH("⚠",{C["البند"]}2))'], fill=FILL_RED))
    ws.conditional_formatting.add(rng(C['سعر الصرف']), FormulaRule(
        formula=[f'AND({C["العملة"]}2<>"",{C["العملة"]}2<>"{base}",{C["سعر الصرف"]}2="")'], fill=FILL_RED))
    ws.conditional_formatting.add(rng(C['حالة الدفع']), FormulaRule(formula=[f'{C["حالة الدفع"]}2="غير مدفوع"'], fill=FILL_AMBER))
    res = C['نتيجة الفحص عند الإعداد']
    ws.conditional_formatting.add(rng(res), FormulaRule(formula=[f'ISNUMBER(SEARCH("{ERROR}",{res}2))'], fill=FILL_RED))
    ws.conditional_formatting.add(rng(res), FormulaRule(formula=[f'ISNUMBER(SEARCH("{WARNING}",{res}2))'], fill=FILL_AMBER))
    return ws


def build_budget(wb, codes, budget, months, base):
    ws = wb.create_sheet(S_BUD)
    bud_codes = sorted(set(c for c in codes if codes[c]['level'] == 2) | set(budget))
    headers = ['الرمز', 'البند', 'رمز الفئة', 'تشغيلي/رأسمالي'] + [m for m in months] + ['الإجمالي']
    ws.append(headers)
    style_header(ws, 1, 1, len(headers))
    for i, m in enumerate(months, start=5):
        ws.cell(row=1, column=i).number_format = FMT_MONTH
    bfmt = base_fmt(base)
    first_m, last_m = get_column_letter(5), get_column_letter(4 + len(months))
    tot_col = get_column_letter(5 + len(months))
    for r, code in enumerate(bud_codes, start=2):
        name = codes[code]['name_ar'] if code in codes else '⚠ رمز غير معرّف'
        ws.cell(row=r, column=1, value=code)
        ws.cell(row=r, column=2, value=name)
        ws.cell(row=r, column=3, value=f'=IF(A{r}>=100,INT(A{r}/100),A{r})')
        ws.cell(row=r, column=4, value=f'=IFERROR(VLOOKUP(A{r},{LST}!$A:$H,8,FALSE),"")')
        for i, m in enumerate(months, start=5):
            amount = budget.get(code, {}).get(m)
            ws.cell(row=r, column=i, value=round(amount, 2) if amount else None)
        ws.cell(row=r, column=5 + len(months), value=f'=SUM({first_m}{r}:{last_m}{r})')
        is_cat = code in codes and codes[code]['level'] == 1
        for col in range(1, 6 + len(months)):
            cell = ws.cell(row=r, column=col)
            font = F_INPUT if 5 <= col < 5 + len(months) else (F_BOLD if is_cat else F_BODY)
            style_body(cell, bfmt if col >= 5 else None, font)
    total_row = len(bud_codes) + 2
    ws.cell(row=total_row, column=2, value='الإجمالي')
    for col in range(5, 6 + len(months)):
        L = get_column_letter(col)
        ws.cell(row=total_row, column=col, value=f'=SUM({L}2:{L}{total_row - 1})').number_format = bfmt
    style_total(ws, total_row, 1, 5 + len(months))
    set_widths(ws, [8, 34, 8, 12] + [14] * len(months) + [16])
    ws.freeze_panes = 'C2'
    ws.cell(row=total_row + 2, column=2, value='الأرقام الزرقاء مدخلات: عدّلها هنا وتتحدث أوراق الملخص والانحراف تلقائياً.').font = F_NOTE
    return tot_col, total_row


def build_variance(wb, codes, rows, budget, base, bud_tot_col):
    ws = wb.create_sheet(S_VAR)
    used = {r['cost_code'] for r in rows if isinstance(r['cost_code'], int)} | set(budget)
    var_codes = sorted(c for c in used if c is not None)
    headers = ['الرمز', 'البند', 'الفئة الرئيسية', 'الموازنة', 'الفعلي', 'الوفر / (التجاوز)', 'نسبة الاستخدام', 'الحالة']
    ws.append(headers)
    style_header(ws, 1, 1, len(headers))
    bfmt = base_fmt(base)
    for r, code in enumerate(var_codes, start=2):
        cat = code // 100 if code >= 100 else code
        ws.cell(row=r, column=1, value=code)
        ws.cell(row=r, column=2, value=codes[code]['name_ar'] if code in codes else '⚠ رمز غير معرّف')
        ws.cell(row=r, column=3, value=codes[cat]['name_ar'] if cat in codes else '')
        ws.cell(row=r, column=4, value=f'=SUMIFS({BUD}!${bud_tot_col}:${bud_tot_col},{BUD}!$A:$A,A{r})')
        ws.cell(row=r, column=5, value=f'=SUMIFS({REG}!${C["الإجمالي بالعملة الأساسية"]}:${C["الإجمالي بالعملة الأساسية"]},{REG}!${C["رمز الكلفة"]}:${C["رمز الكلفة"]},A{r})')
        ws.cell(row=r, column=6, value=f'=D{r}-E{r}')
        ws.cell(row=r, column=7, value=f'=IF(D{r}=0,"",E{r}/D{r})')
        ws.cell(row=r, column=8, value=f'=IF(AND(D{r}=0,E{r}>0),"بدون موازنة للبند",IF(E{r}>D{r},"تجاوز",IF(E{r}>=0.9*D{r},"قريب من الحد","ضمن الموازنة")))')
        for col, fmt in zip(range(1, 9), [None, None, None, bfmt, bfmt, bfmt, FMT_PCT, None]):
            style_body(ws.cell(row=r, column=col), fmt, align=CENTER if col == 8 else None)
    total_row = len(var_codes) + 2
    ws.cell(row=total_row, column=2, value='الإجمالي')
    for col in (4, 5, 6):
        L = get_column_letter(col)
        ws.cell(row=total_row, column=col, value=f'=SUM({L}2:{L}{total_row - 1})').number_format = bfmt
    ws.cell(row=total_row, column=7, value=f'=IF(D{total_row}=0,"",E{total_row}/D{total_row})').number_format = FMT_PCT
    style_total(ws, total_row, 1, 8)
    status = f'H2:H{total_row - 1}'
    for text, fill in (('تجاوز', FILL_RED), ('قريب من الحد', FILL_AMBER), ('ضمن الموازنة', FILL_GREEN),
                       ('بدون موازنة للبند', FILL_LILAC)):
        ws.conditional_formatting.add(status, FormulaRule(formula=[f'H2="{text}"'], fill=fill))
    set_widths(ws, [8, 34, 24, 16, 16, 18, 12, 18])
    ws.freeze_panes = 'C2'
    notes = ['الوفر / (التجاوز) = الموازنة − الفعلي؛ القيمة بين قوسين وبالأحمر تعني تجاوزاً.',
             'قريب من الحد = استُخدم 90% أو أكثر من موازنة البند.',
             'إذا كانت الموازنة على مستوى الفئة الرئيسية فقارن على مستوى الفئة في ورقة الملخص.']
    for i, text in enumerate(notes):
        ws.cell(row=total_row + 2 + i, column=2, value=text).font = F_NOTE
    return ws


def build_issues(wb, issues):
    ws = wb.create_sheet(S_ISSUES)
    headers = ['#', 'الخطورة', 'الصف في السجل', 'الحقل', 'الملاحظة']
    ws.append(headers)
    style_header(ws, 1, 1, len(headers))
    for i, x in enumerate(issues, start=1):
        ws.append([i, x.severity, x.row, x.field, x.message])
        for col in range(1, 6):
            style_body(ws.cell(row=i + 1, column=col), align=CENTER if col < 5 else RIGHT)
    if not issues:
        ws.append([None, None, None, None, 'لا توجد ملاحظات — البيانات اجتازت كل الفحوصات.'])
    last = len(issues) + 1
    for text, fill in ((ERROR, FILL_RED), (WARNING, FILL_AMBER), (INFO, FILL_SECTION)):
        ws.conditional_formatting.add(f'B2:B{max(last, 2)}', FormulaRule(formula=[f'B2="{text}"'], fill=fill))
    set_widths(ws, [5, 10, 12, 14, 95])
    ws.freeze_panes = 'A2'
    return ws


def build_summary(wb, codes, rows, budget, centers, months, opts, bud_tot_col, bud_total_row):
    ws = wb.create_sheet(S_SUMMARY, 0)
    base = opts.base
    bfmt = base_fmt(base)
    U, Fc, Dc = C['الإجمالي بالعملة الأساسية'], C['رمز الفئة'], C['رمز الكلفة']
    U_rng = f'{REG}!${U}:${U}'
    set_widths(ws, [9, 34, 19, 19, 19, 19, 19, 3])

    ws.merge_cells('A1:G1')
    ws['A1'] = opts.company
    ws['A1'].font, ws['A1'].alignment = F_TITLE, CENTER
    ws.merge_cells('A2:G2')
    ws['A2'] = 'تقرير التكاليف — الفعلي مقابل الموازنة'
    ws['A2'].font, ws['A2'].alignment = F_SUB, CENTER
    ws.merge_cells('A3:G3')
    ws['A3'] = (f'الفترة: {months[0]:%Y-%m} إلى {months[-1]:%Y-%m}   |   العملة الأساسية: {base}   |   '
                f'تاريخ الإعداد: {dt.date.today():%Y-%m-%d}')
    ws['A3'].font, ws['A3'].alignment = F_NOTE, CENTER
    ws.row_dimensions[1].height = 28

    # الفئات التي فيها فعلي أو موازنة: التشغيلية أولاً ثم الرأسمالية
    cat_of = lambda c: c // 100 if c >= 100 else c
    active = {cat_of(r['cost_code']) for r in rows if isinstance(r['cost_code'], int)} | {cat_of(c) for c in budget}
    cats = [c for c in sorted(active) if c in codes]
    opex = [c for c in cats if codes[c]['capex_opex'] == 'تشغيلي']
    capex = [c for c in cats if codes[c]['capex_opex'] == 'رأسمالي']

    # --- جدول الفئات (يُبنى أولاً لأن البطاقات تشير لمجاميعه) ---
    row = 9
    ws.cell(row=row, column=1, value='التكاليف حسب الفئة الرئيسية').font = F_SECTION
    row += 1
    cat_header_row = row
    for col, h in enumerate(['الرمز', 'الفئة', 'الفعلي', 'الموازنة', 'الوفر / (التجاوز)', 'نسبة الاستخدام', 'الحصة'], start=1):
        ws.cell(row=row, column=col, value=h)
    style_header(ws, row, 1, 7)
    grand_ref = {}

    def cat_block(cat_list, label):
        nonlocal row
        start = row + 1
        for cat in cat_list:
            row += 1
            ws.cell(row=row, column=1, value=cat)
            ws.cell(row=row, column=2, value=codes[cat]['name_ar'])
            ws.cell(row=row, column=3, value=f'=SUMIFS({U_rng},{REG}!${Fc}:${Fc},$A{row})')
            ws.cell(row=row, column=4, value=f'=SUMIFS({BUD}!${bud_tot_col}:${bud_tot_col},{BUD}!$C:$C,$A{row})')
            ws.cell(row=row, column=5, value=f'=D{row}-C{row}')
            ws.cell(row=row, column=6, value=f'=IF(D{row}=0,"",C{row}/D{row})')
            ws.cell(row=row, column=7, value=f'=IF(GRAND=0,"",C{row}/GRAND)')
            for col, fmt in zip(range(1, 8), [None, None, bfmt, bfmt, bfmt, FMT_PCT, FMT_PCT]):
                style_body(ws.cell(row=row, column=col), fmt)
        row += 1
        ws.cell(row=row, column=2, value=label)
        for col in (3, 4, 5):
            L = get_column_letter(col)
            ws.cell(row=row, column=col, value=f'=SUM({L}{start}:{L}{row - 1})' if cat_list else 0).number_format = bfmt
        ws.cell(row=row, column=6, value=f'=IF(D{row}=0,"",C{row}/D{row})').number_format = FMT_PCT
        ws.cell(row=row, column=7, value=f'=IF(GRAND=0,"",C{row}/GRAND)').number_format = FMT_PCT
        style_total(ws, row, 1, 7)
        return start, row

    opex_start, opex_total = cat_block(opex, 'إجمالي التكاليف التشغيلية')
    capex_total = None
    if capex:
        _, capex_total = cat_block(capex, 'إجمالي الإنفاق الرأسمالي')
    row += 1
    grand_row = row
    ws.cell(row=row, column=2, value='الإجمالي العام')
    parts = [opex_total] + ([capex_total] if capex_total else [])
    for col in (3, 4, 5):
        L = get_column_letter(col)
        ws.cell(row=row, column=col, value='=' + '+'.join(f'{L}{p}' for p in parts)).number_format = bfmt
    ws.cell(row=row, column=6, value=f'=IF(D{row}=0,"",C{row}/D{row})').number_format = FMT_PCT
    ws.cell(row=row, column=7, value=f'=IF(C{row}=0,"",C{row}/C{row})').number_format = FMT_PCT
    style_total(ws, row, 1, 7)
    for r in range(cat_header_row + 1, grand_row + 1):   # استبدال المرجع المؤقت للإجمالي العام
        cell = ws.cell(row=r, column=7)
        if isinstance(cell.value, str):
            cell.value = cell.value.replace('GRAND', f'$C${grand_row}')
    ws.conditional_formatting.add(f'F{cat_header_row + 1}:F{grand_row}',
                                  FormulaRule(formula=[f'AND(ISNUMBER(F{cat_header_row + 1}),F{cat_header_row + 1}>1)'],
                                              font=Font(name=FONT_NAME, size=10, bold=True, color='C00000')))

    # --- بطاقات المؤشرات ---
    Wc = C['حالة الدفع']
    tiles = [
        ('التكاليف التشغيلية (فعلي)', f'=C{opex_total}', bfmt),
        ('الموازنة التشغيلية', f'=D{opex_total}', bfmt),
        ('الوفر / (التجاوز)', f'=E{opex_total}', bfmt),
        ('نسبة استخدام الموازنة', f'=F{opex_total}', FMT_PCT),
        ('الإنفاق الرأسمالي', f'=C{capex_total}' if capex_total else 0, bfmt),
        ('قيود غير مسددة بالكامل', f'=SUMIFS({U_rng},{REG}!${Wc}:${Wc},"غير مدفوع")+SUMIFS({U_rng},{REG}!${Wc}:${Wc},"جزئي")', bfmt),
    ]
    for i, (label, formula, fmt) in enumerate(tiles):
        col = 2 + i  # بطاقة لكل عمود من B إلى G
        lab = ws.cell(row=5, column=col, value=label)
        val = ws.cell(row=6, column=col, value=formula)
        lab.font, lab.fill, lab.alignment = F_TILE_LABEL, FILL_TILE, CENTER
        val.font, val.fill, val.alignment, val.number_format = F_TILE_VALUE, FILL_TILE, CENTER, fmt
    ws['A5'], ws['A6'] = 'المؤشرات', None
    ws['A5'].font = F_SECTION
    ws.row_dimensions[5].height, ws.row_dimensions[6].height = 30, 30
    ws.cell(row=7, column=2, value=f'=COUNT({REG}!$A:$A)&" قيد  |  أخطاء الفحص: "&COUNTIF(\'{S_ISSUES}\'!$B:$B,"{ERROR}")&"  |  تحذيرات: "&COUNTIF(\'{S_ISSUES}\'!$B:$B,"{WARNING}")').font = F_NOTE

    # --- حسب مركز الكلفة ---
    row = grand_row + 2
    ws.cell(row=row, column=1, value='التكاليف حسب مركز الكلفة').font = F_SECTION
    row += 1
    for col, h in enumerate(['', 'مركز الكلفة', 'تشغيلي', 'رأسمالي', 'الإجمالي', 'الحصة'], start=1):
        ws.cell(row=row, column=col, value=h or None)
    style_header(ws, row, 2, 6)
    Cc, Jc = C['مركز الكلفة'], C['تشغيلي/رأسمالي']
    c_start = row + 1
    for center in centers:
        row += 1
        ws.cell(row=row, column=2, value=center)
        ws.cell(row=row, column=3, value=f'=SUMIFS({U_rng},{REG}!${Cc}:${Cc},$B{row},{REG}!${Jc}:${Jc},"تشغيلي")')
        ws.cell(row=row, column=4, value=f'=SUMIFS({U_rng},{REG}!${Cc}:${Cc},$B{row},{REG}!${Jc}:${Jc},"رأسمالي")')
        ws.cell(row=row, column=5, value=f'=C{row}+D{row}')
        ws.cell(row=row, column=6, value=f'=IF($C${grand_row}=0,"",E{row}/$C${grand_row})')
        for col, fmt in zip(range(2, 7), [None, bfmt, bfmt, bfmt, FMT_PCT]):
            style_body(ws.cell(row=row, column=col), fmt)
    row += 1
    ws.cell(row=row, column=2, value='الإجمالي')
    for col in (3, 4, 5):
        L = get_column_letter(col)
        ws.cell(row=row, column=col, value=f'=SUM({L}{c_start}:{L}{row - 1})').number_format = bfmt
    ws.cell(row=row, column=6, value=f'=IF($C${grand_row}=0,"",E{row}/$C${grand_row})').number_format = FMT_PCT
    style_total(ws, row, 2, 6)
    ws.cell(row=row + 1, column=2,
            value=f'=IF(ROUND(E{row}-$C${grand_row},0)=0,"✓ مطابق للإجمالي العام","⚠ فرق "&TEXT($C${grand_row}-E{row},"#,##0")&" عن الإجمالي العام: راجع القيود بدون مركز كلفة أو برمز غير معرّف")').font = F_NOTE

    # --- حسب الشهر (التشغيلي) ---
    row += 3
    ws.cell(row=row, column=1, value='التكاليف التشغيلية حسب الشهر').font = F_SECTION
    row += 1
    month_header = row
    for col, h in enumerate(['', 'الشهر', 'الفعلي', 'الموازنة', 'الوفر / (التجاوز)', 'نسبة الاستخدام'], start=1):
        ws.cell(row=row, column=col, value=h or None)
    style_header(ws, row, 2, 6)
    Vc = C['الشهر']
    m_start = row + 1
    for i, m in enumerate(months):
        row += 1
        bud_col = get_column_letter(5 + i)
        ws.cell(row=row, column=2, value=m)
        ws.cell(row=row, column=3, value=f'=SUMIFS({U_rng},{REG}!${Vc}:${Vc},$B{row},{REG}!${Jc}:${Jc},"تشغيلي")')
        ws.cell(row=row, column=4, value=f'=SUMIFS({BUD}!${bud_col}$2:${bud_col}${bud_total_row - 1},{BUD}!$D$2:$D${bud_total_row - 1},"تشغيلي")')
        ws.cell(row=row, column=5, value=f'=D{row}-C{row}')
        ws.cell(row=row, column=6, value=f'=IF(D{row}=0,"",C{row}/D{row})')
        for col, fmt in zip(range(2, 7), [FMT_MONTH, bfmt, bfmt, bfmt, FMT_PCT]):
            style_body(ws.cell(row=row, column=col), fmt, align=CENTER if col == 2 else None)
    m_end = row
    row += 1
    ws.cell(row=row, column=2, value='الإجمالي')
    for col in (3, 4, 5):
        L = get_column_letter(col)
        ws.cell(row=row, column=col, value=f'=SUM({L}{m_start}:{L}{m_end})').number_format = bfmt
    ws.cell(row=row, column=6, value=f'=IF(D{row}=0,"",C{row}/D{row})').number_format = FMT_PCT
    style_total(ws, row, 2, 6)
    ws.cell(row=row + 1, column=2,
            value=f'=IF(ROUND(C{row}-C{opex_total},0)=0,"✓ مطابق لإجمالي التشغيلي","⚠ فرق "&TEXT(C{opex_total}-C{row},"#,##0")&" عن إجمالي التشغيلي: راجع القيود خارج الفترة أو برمز غير معرّف")').font = F_NOTE

    # --- الرسوم البيانية ---
    if opex:
        chart = BarChart()
        chart.type, chart.style = 'bar', 10
        chart.title = 'الفعلي مقابل الموازنة حسب الفئة (مليون)'
        chart.y_axis.numFmt = '#,##0,,"م"'
        chart.add_data(Reference(ws, min_col=3, max_col=4, min_row=cat_header_row, max_row=opex_total - 1), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=2, min_row=cat_header_row + 1, max_row=opex_total - 1))
        chart.height, chart.width = 9, 16
        ws.add_chart(chart, 'I5')
    chart2 = BarChart()
    chart2.type, chart2.style = 'col', 10
    chart2.title = 'التكاليف التشغيلية الشهرية (مليون)'
    chart2.y_axis.numFmt = '#,##0,,"م"'
    chart2.add_data(Reference(ws, min_col=3, max_col=4, min_row=month_header, max_row=m_end), titles_from_data=True)
    chart2.set_categories(Reference(ws, min_col=2, min_row=m_start, max_row=m_end))
    chart2.height, chart2.width = 8, 16
    ws.add_chart(chart2, 'I25')

    page_setup(ws, opts.company, landscape=True)
    ws.print_area = f'A1:Q{row + 1}'
    return ws


def build_method(wb, opts, rows, issues, source_name):
    ws = wb.create_sheet(S_METHOD)
    page_setup(ws, opts.company, landscape=False)
    set_widths(ws, [26, 30, 22, 18])
    ws['A1'] = 'المنهجية والاعتماد'
    ws['A1'].font = F_TITLE
    fx_used = defaultdict(lambda: defaultdict(int))
    for r in rows:
        if r['currency'] != opts.base and r['fx_rate']:
            fx_used[r['currency']][(r['fx_rate'], r['fx_source'])] += 1
    fx_text = '؛ '.join(f'{cur}: ' + '، '.join(f'{rate:,.2f} ({src}، {n} قيد)' for (rate, src), n in sorted(v.items()))
                        for cur, v in sorted(fx_used.items())) or 'لا توجد عملات أجنبية'
    counts = {s: sum(1 for x in issues if x.severity == s) for s in (ERROR, WARNING, INFO)}
    items = [
        ('الشركة', opts.company),
        ('مصدر البيانات', source_name),
        ('عدد القيود', len(rows)),
        ('العملة الأساسية', opts.base),
        ('أسعار الصرف المستخدمة', fx_text),
        ('دليل الترميز', 'مبني على هيكل الدليل المحاسبي للنظام المحاسبي الموحد (المستويان الأول والثاني)'),
        ('نتيجة الفحص', f'أخطاء: {counts[ERROR]}، تحذيرات: {counts[WARNING]}، ملاحظات: {counts[INFO]} — التفاصيل في ورقة "{S_ISSUES}"'),
        ('الفعلي', 'يشمل الضرائب والرسوم غير المستردة، ومحوّل للعملة الأساسية بسعر صرف تاريخ المعاملة (IAS 21)'),
        ('الوفر / (التجاوز)', 'الموازنة − الفعلي؛ الأقواس والأحمر = تجاوز'),
        ('نسبة الاستخدام', 'الفعلي ÷ الموازنة؛ 90% فأكثر = قريب من الحد'),
        ('الألوان', 'الخط الأزرق = مدخلات تُعدّل يدوياً؛ الخط الأسود = معادلات لا تُكتب فوقها'),
        ('إضافة قيود', 'أضف الصف داخل جدول السجل وانسخ معادلات الصف السابق، أو أعد بناء الملف من المصدر'),
        ('تاريخ ووقت الإعداد', f'{dt.datetime.now():%Y-%m-%d %H:%M}'),
    ]
    for i, (k, v) in enumerate(items, start=3):
        a, b = ws.cell(row=i, column=1, value=k), ws.cell(row=i, column=2, value=v)
        style_body(a, font=F_BOLD, align=RIGHT)
        style_body(b, align=RIGHT)
        a.fill = FILL_SECTION
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=4)
        ws.row_dimensions[i].height = 30 if len(str(v)) > 60 else 18
    row = len(items) + 5
    ws.cell(row=row, column=1, value='الاعتماد').font = F_SECTION
    row += 1
    for col, h in enumerate(['الدور', 'الاسم', 'التوقيع', 'التاريخ'], start=1):
        ws.cell(row=row, column=col, value=h)
    style_header(ws, row, 1, 4)
    for role, name in (('أعدّه', opts.prepared_by), ('راجعه', None), ('اعتمده (المدير المالي)', None)):
        row += 1
        ws.row_dimensions[row].height = 32
        for col, v in enumerate([role, name, None, None], start=1):
            style_body(ws.cell(row=row, column=col, value=v), font=F_BOLD if col == 1 else F_INPUT, align=RIGHT)


# ---------- التحقق بعد البناء ----------

def soffice():
    return shutil.which('soffice') or shutil.which('libreoffice')


def recalc_check(path):
    """يعيد الحساب بـ LibreOffice ويرجع قائمة الخلايا التي فيها أخطاء معادلات."""
    exe = soffice()
    if not exe:
        return None
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run([exe, '--headless', '--calc', '--convert-to', 'xlsx', '--outdir', tmp, str(path)],
                       check=True, capture_output=True, timeout=180)
        wb = load_workbook(Path(tmp) / Path(path).name, data_only=True)
        errors = [f'{ws.title}!{cell.coordinate}: {cell.value}'
                  for ws in wb.worksheets for row in ws.iter_rows() for cell in row
                  if cell.data_type == 'e' or (isinstance(cell.value, str) and cell.value.startswith('Err:'))]
        return errors, wb[S_SUMMARY]


def export_pdf(path):
    exe = soffice()
    if not exe:
        return None
    out_dir = Path(path).parent
    subprocess.run([exe, '--headless', '--convert-to', 'pdf', '--outdir', str(out_dir), str(path)],
                   check=True, capture_output=True, timeout=180)
    return out_dir / (Path(path).stem + '.pdf')


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    common_args(parser)
    parser.add_argument('--company', required=True, help='اسم الشركة كما يظهر بالتقرير')
    parser.add_argument('--out', required=True, help='مسار ملف xlsx الناتج')
    parser.add_argument('--budget', help='ملف الموازنة CSV')
    parser.add_argument('--prepared-by', help='اسم معدّ التقرير')
    parser.add_argument('--check', action='store_true', help='إعادة الحساب بـ LibreOffice وفحص أخطاء المعادلات')
    parser.add_argument('--pdf', action='store_true', help='تصدير نسخة PDF بجانب ملف xlsx')
    opts = parser.parse_args()
    opts.base = opts.base.upper()

    rows, codes, issues = run_checks(opts)
    dates = [r['date'] for r in rows if r['date']]
    if not dates:
        raise SystemExit('لا توجد قيود بتواريخ صالحة.')
    start = opts.period_start or min(dates)
    end = opts.period_end or max(dates)
    months = month_list(start, end)
    budget = load_budget(opts.budget, months) if opts.budget else {}

    row_issues = defaultdict(list)
    for x in issues:
        if x.severity in (ERROR, WARNING):
            row_issues[x.row].append(f'{x.severity}: {x.message}')
    centers = sorted({r['cost_center'] for r in rows if r['cost_center']})
    currencies = sorted({opts.base} | {r['currency'] for r in rows if r['currency']})

    wb = Workbook()
    wb.remove(wb.active)
    n_codes_last, n_centers_last, n_cur_last = build_lists(wb, codes, centers or ['—'], currencies)
    build_register(wb, rows, row_issues, opts.base, n_codes_last, n_centers_last, n_cur_last)
    bud_tot_col, bud_total_row = build_budget(wb, codes, budget, months, opts.base)
    build_variance(wb, codes, rows, budget, opts.base, bud_tot_col)
    build_issues(wb, issues)
    build_summary(wb, codes, rows, budget, centers, months, opts, bud_tot_col, bud_total_row)
    build_method(wb, opts, rows, issues, Path(opts.costs).name)
    # ترتيب الأوراق: الملخص، السجل، الموازنة، الانحراف، الملاحظات، القوائم، المنهجية
    order = [S_SUMMARY, S_REG, S_BUD, S_VAR, S_ISSUES, S_LISTS, S_METHOD]
    wb._sheets = [wb[name] for name in order]
    wb.active = 0
    for name in (S_REG, S_BUD, S_VAR, S_ISSUES, S_LISTS):
        page_setup(wb[name], opts.company, '1:1')
    wb.properties.title = f'تقرير التكاليف — {opts.company}'
    wb.properties.creator = opts.prepared_by or opts.company
    wb.properties.subject = 'Cost report: actual vs budget'

    out = Path(opts.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    counts = {s: sum(1 for x in issues if x.severity == s) for s in (ERROR, WARNING, INFO)}
    print(f'تم إنشاء {out} — القيود: {len(rows)} | أخطاء: {counts[ERROR]} | تحذيرات: {counts[WARNING]} | ملاحظات: {counts[INFO]}')

    if opts.check:
        result = recalc_check(out)
        if result is None:
            print('تنبيه: LibreOffice غير متوفر — لم يُفحص الحساب.')
        else:
            errors, summary = result
            print('فحص المعادلات: ' + ('لا توجد أخطاء ✓' if not errors else f'{len(errors)} خطأ'))
            for e in errors[:30]:
                print('  ', e)
            for col in range(2, 8):
                label, value = summary.cell(row=5, column=col).value, summary.cell(row=6, column=col).value
                print(f'   {label}: {value:,.2f}' if isinstance(value, (int, float)) else f'   {label}: {value}')
            if errors:
                sys.exit(2)
    if opts.pdf:
        pdf = export_pdf(out)
        print(f'PDF: {pdf}' if pdf else 'تنبيه: LibreOffice غير متوفر — لم يُصدَّر PDF.')


if __name__ == '__main__':
    main()
